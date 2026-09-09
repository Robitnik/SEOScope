"""SEO-стратегія (ручний Google-Sheets воркбук, завантажений як .xlsx) ->
ContentDoc.

**task-90 §4.2:** табличне ядро (мапінг колонок, блоки-заголовки, побудова
`ContentItem`) винесено в `readers/tabular.py` — воно спільне з Google
Sheets-ридером (`readers/gsheet.py`), бо обидва джерела зводяться до
`list[tuple[object, ...]]` на аркуш. Цей модуль лишає лише те, що дійсно
специфічне для openpyxl: розкрити байти `.xlsx` у рядки клітинок. Поведінка
не змінилась — `_scan_workbook` тепер лише готує вхід для
`tabular.scan_sheets`, решта викликів тонкі.

Заголовок шукається по **всьому** аркушу, не лише перших рядках: живий файл
(`GOREPORT_SEO_STRATEGY_07272026.xlsx`, аркуш "Structure Features
Description") містить дві незалежні таблиці різної форми одна під одною —
"Top Performing Pages" (Page/URL/Keyword/Traffic/Title/Description/H1) і
нижче, після банера-розділювача, "URL Structure Review & Update"
(PageStatus/Intent/Pages/URL/RecommendedURL/Breadcrumbs/WordCount/Comments,
без title/description). Пошук заголовка лише в перших 5 рядках зафіксував би
мапінг першої таблиці назавжди — тоді колонка URL другої таблиці читалась би
як title, а Word count як description (саме так і сталось при першому
прогоні на живих даних: 120 "meta-title-length" замість очікуваних ~8, бо
довгі URL і числа підмінили собою копірайт). Кожен рядок із двома і більше
впізнаними заголовками починає новий блок — навіть якщо в ньому немає
title/description: такий блок просто не дає жодного item, а не поглинається
попереднім.

Аркуш, де жоден блок не має ні title, ні description, -> `skipped_sheets`, а
не тиха «чистота» (PLAN §12: не перевірили ≠ порушень немає).

**task-87 §1.3 (вибір аркуша):** `sheets=` фільтрує, які аркуші йдуть у
`ContentDoc.items`; аркуші, яких оператор не обрав, усе одно скануються
(легко — лише підрахунок рядків, без побудови items) і потрапляють у
`ContentDoc.sheet_states` зі станом `"not_selected"`, щоб блок Scope у звіті
міг сказати «Q2 2026 Strategy (154 rows) — not selected», а не мовчати про
нього. `inspect_strategy_xlsx()` ділить той самий скан-код для прев'ю кроку
1 (`POST …/strategy-check/inspect`) — жодного другого парсера.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.contentqa.model import ContentDoc
from app.contentqa.readers.tabular import (
    SheetPreview,
    SheetScan,
    UnknownSheetError,
    WorkbookInspection,
    build_content_doc,
    build_inspection,
    scan_sheets,
)

__all__ = [
    "SheetPreview",
    "UnknownSheetError",
    "WorkbookInspection",
    "inspect_strategy_xlsx",
    "read_strategy_xlsx",
]


def _scan_workbook(content: bytes) -> list[SheetScan]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    raw_sheets = [
        (sheet_name, [tuple(row) for row in wb[sheet_name].iter_rows(values_only=True)])
        for sheet_name in wb.sheetnames
    ]
    return scan_sheets(raw_sheets)


def read_strategy_xlsx(
    content: bytes, *, source: str, sheets: list[str] | None = None
) -> ContentDoc:
    """`sheets=None` (дефолт) перевіряє всі несучі аркуші — те саме
    поводження, що до task-87. `sheets=[...]` обмежує items лише цими
    аркушами; решта аркушів workbook усе одно потрапляє в
    `ContentDoc.sheet_states` зі станом `not_selected`, а не зникає мовчки."""
    scans = _scan_workbook(content)
    return build_content_doc(scans, source=source, doc_type="strategy_xlsx", sheets=sheets)


def inspect_strategy_xlsx(content: bytes) -> WorkbookInspection:
    """Крок 1 вибору аркуша (task-87 §1.3): прев'ю без побудови items і без
    жодного побічного ефекту — той самий скан, що `read_strategy_xlsx`."""
    return build_inspection(_scan_workbook(content))
