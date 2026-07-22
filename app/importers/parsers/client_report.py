"""Client report XLSX parser (task-07): only the `Historical Data` sheet.

`client_report_<CODE>_<date>.xlsx` has three sheets. `Current Month Data`/
`Import Data` and the `.tsv` twin all key off a relative "This Month"
column whose meaning changed mid-dataset (confirmed via git log on the
AISeoTools source: a fix landed 2026-06-19 that changed "This Month" from
the current partial month to the last complete month, and nothing in the
file itself says which era a given file is from) — guessing the era would
be exactly the kind of invented period this importer avoids everywhere
else (see `app.importers.parsers.gsc` for the same reasoning applied to GSC
Family A), so those sheets are never parsed here.

`Historical Data` instead has one row per explicit calendar month label
("Jun 2026") — unambiguous, so its `Clicks`/`Impressions`/`Organic
Conversions` columns are safe to import as exact-month range metrics.
`CTR` is always `None` in real samples and is skipped regardless: derived
CTR is computed once, centrally, by `app.metrics.normalize.normalize_result`
from clicks/impressions, never hand-set by a parser.

One file can produce several snapshots: one per `(provider, month)` pair
(`gsc` clicks/impressions, `reporting_ninja` organic conversions) — never
one artifact bundling a whole time series, since `metric_values` uniques on
`(snapshot_id, metric_key, period_role)` and a single snapshot can only
hold one `gsc_clicks` value, not fourteen.
"""

from __future__ import annotations

import calendar
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.importers.parsers import ParsedArtifact, ParseSkip
from app.schemas.adapter import MetricValueResult
from app.schemas.enums import MetricUnit, PeriodRole, Provider, Quality, SourceMode

IMPORTER_VERSION = "client_report_historical.v1"

_SHEET_NAME = "Historical Data"
_MONTH_FORMATS = ("%b %Y", "%B %Y")
_REQUIRED_COLUMNS = ("date", "clicks", "impressions", "organic conversions")
_MAX_HEADER_SEARCH_ROWS = 5


def _parse_month(label: Any) -> tuple[date, date] | None:
    if isinstance(label, datetime):
        start = label.date().replace(day=1)
    elif isinstance(label, str):
        parsed = None
        for fmt in _MONTH_FORMATS:
            try:
                parsed = datetime.strptime(label.strip(), fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            return None
        start = parsed.date().replace(day=1)
    else:
        return None
    last_day = calendar.monthrange(start.year, start.month)[1]
    return start, start.replace(day=last_day)


def _range_metric(
    key: str, value: Any, start: date, end: date, unit: MetricUnit
) -> MetricValueResult | None:
    if value is None:
        return None
    try:
        numeric = Decimal(str(value))
    except InvalidOperation:
        return None
    return MetricValueResult(
        metric_key=key,
        period_role=PeriodRole.current,
        period_start=start,
        period_end=end,
        value_numeric=numeric,
        unit=unit,
        quality=Quality.complete,
    )


def parse(files: dict[str, Path], *, fallback_date: date) -> list[ParsedArtifact] | ParseSkip:
    path = files.get("client_report.xlsx")
    if path is None:
        return ParseSkip(reason="client_report xlsx not found", level="info")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several exception types for corrupt files
        return ParseSkip(reason=f"unreadable/invalid XLSX: {exc}")

    try:
        if _SHEET_NAME not in workbook.sheetnames:
            return ParseSkip(reason=f"unknown schema: no '{_SHEET_NAME}' sheet")

        sheet = workbook[_SHEET_NAME]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        return ParseSkip(reason=f"'{_SHEET_NAME}' sheet is empty (no header row)", level="info")

    # Real samples put a title row ("PJC — Monthly History") before the
    # actual header — search the first few rows rather than assuming row 0
    # is the header.
    header: list[str] | None = None
    header_idx = -1
    for i, row in enumerate(rows[:_MAX_HEADER_SEARCH_ROWS]):
        candidate = [str(c).strip().lower() if c is not None else "" for c in row]
        if set(_REQUIRED_COLUMNS).issubset(candidate):
            header, header_idx = candidate, i
            break
    if header is None:
        first_row = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        return ParseSkip(
            reason=f"unknown schema: no recognizable '{_SHEET_NAME}' header {first_row!r}"
        )

    data_rows = rows[header_idx + 1 :]
    if not data_rows:
        # Confirmed in real samples: this sheet is sometimes header-only
        # (e.g. a client with no history yet) — an expected state, not a
        # warning-worthy anomaly.
        return ParseSkip(
            reason=f"'{_SHEET_NAME}' sheet has no data rows below the header", level="info"
        )

    idx_date = header.index("date")
    idx_clicks = header.index("clicks")
    idx_impressions = header.index("impressions")
    idx_conversions = header.index("organic conversions")

    # `metric_values` uniques on `(snapshot_id, metric_key, period_role)` —
    # one snapshot can hold exactly one `gsc_clicks` value, not fourteen (one
    # per historical month). Each month is therefore its own snapshot: group
    # by `(provider, period_start, period_end)` and emit one ParsedArtifact
    # per group, not one artifact bundling the whole time series.
    by_period: dict[tuple[Provider, date, date], list[MetricValueResult]] = {}
    skipped_rows = 0

    for row in data_rows:
        if row is None or all(c is None for c in row):
            continue
        month = _parse_month(row[idx_date]) if idx_date < len(row) else None
        if month is None:
            skipped_rows += 1
            continue
        start, end = month

        if idx_clicks < len(row):
            metric = _range_metric("gsc_clicks", row[idx_clicks], start, end, MetricUnit.count)
            if metric is not None:
                by_period.setdefault((Provider.gsc, start, end), []).append(metric)
        if idx_impressions < len(row):
            metric = _range_metric(
                "gsc_impressions", row[idx_impressions], start, end, MetricUnit.count
            )
            if metric is not None:
                by_period.setdefault((Provider.gsc, start, end), []).append(metric)
        if idx_conversions < len(row):
            metric = _range_metric(
                "rn_organic_conversions", row[idx_conversions], start, end, MetricUnit.count
            )
            if metric is not None:
                by_period.setdefault((Provider.reporting_ninja, start, end), []).append(metric)

    if not by_period:
        return ParseSkip(reason=f"no parseable rows in '{_SHEET_NAME}'")

    warnings = (
        [f"skipped {skipped_rows} row(s) with an unparseable month label"] if skipped_rows else []
    )
    warnings.append(
        "Current Month Data / Import Data sheets and the .tsv twin were not parsed: their "
        "'This Month' period semantics changed mid-dataset with no in-file marker to tell "
        "which era a file is from"
    )

    artifacts: list[ParsedArtifact] = []
    for (provider, _start, end), metrics in sorted(by_period.items(), key=lambda kv: kv[0][1:]):
        artifacts.append(
            ParsedArtifact(
                provider=provider,
                source_mode=SourceMode.pdf,
                quality=Quality.complete,
                observed_at=datetime.combine(end, datetime.min.time(), tzinfo=UTC),
                data_through=end,
                metrics=metrics,
                warnings=warnings,
            )
        )
    return artifacts
