"""Excel-експорт Indexing Planner (task-100 §2.5/§2.6).

Книга будується з ТИХ САМИХ DTO, які читає екран (правило
`app/pageinventory/xlsx.py`): інакше експорт із часом стає другою,
розбіжною реалізацією тих самих правил.

Два аркуші:

- `Queue` — черга з причиною, доказами і станом, у порядку пріоритету;
- `Analytics` — воронка, лаг і **контрольна група**.

Порожня вибірка в аркуші аналітики друкується як `—`, не як `0`: нуль днів
читався б як «заходить того ж дня», тобто найкращий можливий результат.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.indexing import IndexingAnalyticsDTO, IndexingQueueDTO

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MISSING = "—"


def _dash(value: object) -> object:
    return _MISSING if value is None else value


def _num(value: Decimal | None, digits: int = 1) -> object:
    if value is None:
        return _MISSING
    return float(round(value, digits))


def _write_header(sheet, headers: list[str]) -> None:  # type: ignore[no-untyped-def]
    sheet.append(headers)
    for index in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _autosize(sheet, widths: list[int]) -> None:  # type: ignore[no-untyped-def]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_workbook(queue: IndexingQueueDTO, analytics: IndexingAnalyticsDTO) -> bytes:
    book = Workbook()

    sheet = book.active
    sheet.title = "Queue"
    _write_header(
        sheet,
        [
            "Client",
            "URL",
            "Title",
            "Reason",
            "Action",
            "Score",
            "Coverage state",
            "Last crawl",
            "Last submitted",
            "Submissions",
            "Cooldown left",
        ],
    )
    for row in queue.urls:
        sheet.append(
            [
                row.client_name,
                row.url,
                _dash(row.title),
                _dash(row.reason_code),
                _dash(row.action),
                _num(row.priority_score, 0),
                _dash(row.coverage_state),
                _dash(row.last_crawl_time.date().isoformat() if row.last_crawl_time else None),
                _dash(
                    row.last_submitted_at.date().isoformat() if row.last_submitted_at else None
                ),
                row.submissions_count,
                _dash(row.cooldown_days_left),
            ]
        )
    _autosize(sheet, [14, 60, 40, 24, 20, 8, 34, 13, 15, 13, 14])

    stats = book.create_sheet("Analytics")
    stats.append(["Funnel", "Count"])
    for label, value in (
        ("Planned", analytics.funnel.planned),
        ("Sent", analytics.funnel.sent),
        ("Quota blocked", analytics.funnel.quota_blocked),
        ("Indexed", analytics.funnel.indexed),
        ("Still in flight", analytics.funnel.in_flight),
    ):
        stats.append([label, value])

    stats.append([])
    stats.append(["Time to index", "Sample", "Median days", "p25", "p75"])
    stats.append(
        [
            "Submitted by us",
            analytics.control_group.submitted.sample,
            _num(analytics.control_group.submitted.median_days),
            _num(analytics.control_group.submitted.p25_days),
            _num(analytics.control_group.submitted.p75_days),
        ]
    )
    stats.append(
        [
            "Control group (never submitted)",
            analytics.control_group.organic.sample,
            _num(analytics.control_group.organic.median_days),
            _num(analytics.control_group.organic.p25_days),
            _num(analytics.control_group.organic.p75_days),
        ]
    )
    stats.append(
        [
            "Difference (control − submitted)",
            "",
            _num(analytics.control_group.difference_days),
        ]
    )

    stats.append([])
    stats.append(["Reason", "Sent", "Indexed", "In flight", "Median days", "Conversion"])
    for reason in analytics.by_reason:
        stats.append(
            [
                reason.reason_code,
                reason.sent,
                reason.indexed,
                reason.in_flight,
                _num(reason.median_days),
                _num(reason.conversion, 2),
            ]
        )

    stats.append([])
    stats.append(["Client", "URLs", "Coverage", "Queue", "Sent 30d", "Median days", "Cycle holds"])
    for client in analytics.by_client:
        stats.append(
            [
                client.client_name,
                _dash(client.urls_total),
                _num(client.coverage_ratio, 3),
                _dash(client.queue),
                client.sent_30d,
                _num(client.median_days),
                _dash(client.cycle_holds),
            ]
        )
    _autosize(stats, [34, 12, 14, 12, 14, 14, 14])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
