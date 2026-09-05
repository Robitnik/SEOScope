"""XLSX export for Page Inventory (task-57 §6): two sheets, `Inventory` and
`SEO QA`, built from the same `PageInventoryRowDTO` list the JSON API and
the on-screen table read — never a second, independently-derived copy of
the numbers (same rule `app/reports/monthly_xlsx.py` follows for task-54).
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.page_inventory import PageInventoryRowDTO

_HEADER_BG = "434343"
_HEADER_FG = "FFFFFF"
_ROW_LIVE = "D9EAD3"
_ROW_BROKEN = "FCE5CD"
_ROW_NOINDEX = "FFF2CC"
_ROW_DUPLICATE = "CFE2F3"
_SEVERITY_FILL = {"error": "F4CCCC", "warn": "FFF2CC", "info": "D9D9D9"}

_STATUS_ROW_FILL = {
    "live": _ROW_LIVE,
    "broken": _ROW_BROKEN,
    "redirect": _ROW_BROKEN,
    "noindex": _ROW_NOINDEX,
    "not_checked": None,
}


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _header(ws: Worksheet, labels: list[str]) -> None:
    for col_index, label in enumerate(labels, start=1):
        cell = ws.cell(1, col_index, label)
        cell.font = Font(name="Calibri", bold=True, color=_HEADER_FG, size=11)
        cell.fill = _fill(_HEADER_BG)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}1"


def _row_fill(row: PageInventoryRowDTO) -> str | None:
    if row.duplicate_urls:
        return _ROW_DUPLICATE
    if row.canonical_missing:
        return _ROW_NOINDEX
    return _STATUS_ROW_FILL.get(row.status_label)


def _breadcrumb_text(row: PageInventoryRowDTO) -> str:
    # No `row.page_label` fallback here — a URL-derived page name belongs in
    # the "Pages" column, not disguised as a breadcrumb trail this row does
    # not actually have (task-57 live report 2026-09-04, client GOREPORT).
    return " > ".join(node.name for node in row.breadcrumbs)


def build_inventory_sheet(ws: Worksheet, rows: list[PageInventoryRowDTO]) -> None:
    labels = ["Pages", "URL", "Breadcrumbs", "Title", "Description", "Keywords", "Status"]
    _header(ws, labels)
    for row_index, row in enumerate(rows, start=2):
        fill = _row_fill(row)
        values = [
            row.page_label or row.url,
            row.url,
            _breadcrumb_text(row),
            row.title or "",
            row.description or "",
            ", ".join(row.keywords),
            row.status_label,
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col_index, value)
            if fill:
                cell.fill = _fill(fill)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
    widths = [28, 42, 34, 32, 40, 24, 12]
    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width


def _worst_field_severity(row: PageInventoryRowDTO, field: str) -> str | None:
    order = {"error": 0, "warn": 1, "info": 2}
    worst: str | None = None
    for finding in row.findings_by_field.get(field, []):
        if worst is None or order[finding.severity] < order[worst]:
            worst = finding.severity
    return worst


def _field_issue_text(row: PageInventoryRowDTO, field: str) -> str:
    return "; ".join(f.detail for f in row.findings_by_field.get(field, []))


def build_qa_sheet(ws: Worksheet, rows: list[PageInventoryRowDTO]) -> None:
    labels = [
        "Pages", "URL", "Breadcrumbs", "Title", "Description", "Keywords",
        "Title issues", "Description issues", "Keywords issues", "Other issues",
    ]
    _header(ws, labels)
    for row_index, row in enumerate(rows, start=2):
        values = [
            row.page_label or row.url,
            row.url,
            _breadcrumb_text(row),
            row.title or "",
            row.description or "",
            ", ".join(row.keywords),
            _field_issue_text(row, "title"),
            _field_issue_text(row, "description"),
            _field_issue_text(row, "keywords"),
            "; ".join(
                f.detail for f in row.findings
                if f.field not in ("title", "description", "keywords")
            ),
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col_index, value)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
        for field_key, col_index in (("title", 4), ("description", 5), ("keywords", 6)):
            severity = _worst_field_severity(row, field_key)
            if severity:
                ws.cell(row_index, col_index).fill = _fill(_SEVERITY_FILL[severity])
    widths = [28, 42, 34, 32, 40, 24, 34, 34, 24, 34]
    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.cell(len(rows) + 3, 1, "Legend:").font = Font(bold=True)
    legend = [
        ("Error", _SEVERITY_FILL["error"]),
        ("Warn", _SEVERITY_FILL["warn"]),
        ("Info", _SEVERITY_FILL["info"]),
    ]
    for offset, (label, color) in enumerate(legend):
        cell = ws.cell(len(rows) + 4 + offset, 1, label)
        cell.fill = _fill(color)


def build_page_inventory_xlsx(rows: list[PageInventoryRowDTO]) -> bytes:
    workbook = Workbook()
    inventory_ws = workbook.active
    inventory_ws.title = "Inventory"
    build_inventory_sheet(inventory_ws, rows)

    qa_ws = workbook.create_sheet("SEO QA")
    build_qa_sheet(qa_ws, rows)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
