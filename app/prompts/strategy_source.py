"""Джерело стратегії для Prompt Launcher (task-102 §2): посилання на Google
Sheets/Drive або шлях до xlsx у AISeoTools -> сирі аркуші з посиланнями.

Розпізнавання входу детерміноване й нічого не вгадує: `http(s)://` -> Drive,
рядок зі слешем або `.xlsx` -> шлях, голий Drive-id -> Drive, решта ->
`unrecognized_source`. Жодного `fetch(url)` за введеним — SSRF-класу немає за
конструкцією: з посилання витягується лише Drive-id (`extract_file_id`).

Шлях замкнений коренем AISeoTools (`aiseotools_source_path`, у контейнері
read-only): оператор дає його так, як бачить на хості
(`/home/elon/dev/seo/AISeoTools/out/…`) або відносно, а читає бек лише через
контейнерний корінь. Поза ним файлів у контейнері просто немає.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal, Protocol

from openpyxl import load_workbook

from app.core.config import Settings
from app.core.security import PathTraversalError, safe_resolve
from app.integrations.google_api import GoogleApiConfigError, GoogleApiError
from app.integrations.google_api.drive_client import (
    GOOGLE_DOC_MIME,
    GOOGLE_SHEET_MIME,
    XLSX_MIME,
    DriveClient,
    extract_file_id,
)
from app.integrations.google_api.sheets_client import SheetsClient

MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_LOCAL_LISTING = 50

# Службові артефакти інших тулів, що лежать поруч зі стратегіями в `out/`
# й мають «STRATEGY» в імені, але стратегією не є (результати перевірок,
# знімки «до правок»).
_NOT_A_STRATEGY_MARKERS = ("check", "matrix", "before", "reviewed")


@dataclass(frozen=True)
class Cell:
    text: str | None
    links: tuple[str, ...] = ()


RawRow = tuple[Cell, ...]
RawSheet = tuple[str, list[RawRow]]


@dataclass(frozen=True)
class LoadedStrategy:
    name: str
    kind: Literal["gsheet", "xlsx"]
    ref: str  # Drive-id або шлях відносно кореня AISeoTools — ключ прогресу на фронті
    modified_time: str | None
    web_view_link: str | None
    sheets: list[RawSheet]


class StrategySourceError(Exception):
    """Помилка входу/читання з готовим HTTP-статусом і кодом для `ErrorDTO`."""

    def __init__(self, status: int, code: str, message: str) -> None:
        super().__init__(message)
        self.status = status
        self.code = code
        self.message = message


class GoogleClients(Protocol):
    @property
    def drive(self) -> DriveClient: ...

    @property
    def sheets(self) -> SheetsClient: ...


def _looks_like_path(source: str) -> bool:
    return "/" in source or "\\" in source or source.lower().endswith(".xlsx")


def load_strategy(
    source: str,
    *,
    settings: Settings,
    open_clients: Callable[[], GoogleClients],
) -> LoadedStrategy:
    text = source.strip()
    if not text:
        raise StrategySourceError(
            400, "unrecognized_source", "Paste a Google Sheets link or a path to an .xlsx file."
        )
    if text.lower().startswith(("http://", "https://")):
        return _load_drive(text, open_clients)
    if _looks_like_path(text):
        return _load_path(text, settings)
    if extract_file_id(text) is not None:
        return _load_drive(text, open_clients)
    raise StrategySourceError(
        400,
        "unrecognized_source",
        "Not a Google Sheets link, a Drive file id or a path to an .xlsx file.",
    )


# ── Drive ─────────────────────────────────────────────────────────────────────


def _load_drive(text: str, open_clients: Callable[[], GoogleClients]) -> LoadedStrategy:
    file_id = extract_file_id(text)
    if file_id is None:
        raise StrategySourceError(
            400, "invalid_drive_url", "Not a recognizable Google Drive URL or file id."
        )
    try:
        clients = open_clients()
    except GoogleApiConfigError as exc:
        raise StrategySourceError(
            503,
            "drive_not_configured",
            "Google Drive access is not configured — see task-90 §4.0a "
            "(scripts/google_oauth_bootstrap.py needs the drive.readonly scope).",
        ) from exc
    try:
        file = clients.drive.get_file(file_id)
        if file is None:
            about = clients.drive.about_user()
            hint = f" Share it with {about}." if about else ""
            raise StrategySourceError(
                404, "file_not_accessible", f"The agency account cannot see this file.{hint}"
            )
        if file.mime_type == GOOGLE_SHEET_MIME:
            grid = clients.sheets.get_grid(file.id)
            sheets: list[RawSheet] = [
                (
                    sheet.title,
                    [
                        tuple(
                            Cell(c.formatted_value, tuple(link.uri for link in c.links))
                            for c in row
                        )
                        for row in sheet.rows
                    ],
                )
                for sheet in sorted(grid, key=lambda s: s.index)
            ]
            kind: Literal["gsheet", "xlsx"] = "gsheet"
        elif file.mime_type == XLSX_MIME:
            content = clients.drive.download_media(file.id)
            if len(content) > MAX_XLSX_BYTES:
                raise StrategySourceError(
                    413,
                    "upload_too_large",
                    f"The file exceeds {MAX_XLSX_BYTES // (1024 * 1024)} MB.",
                )
            sheets = _read_xlsx(content)
            kind = "xlsx"
        elif file.mime_type == GOOGLE_DOC_MIME:
            raise StrategySourceError(
                400,
                "unsupported_source",
                "This file is a Google Doc, not a strategy spreadsheet — pick the "
                "SEO strategy sheet, not a framework.",
            )
        else:
            raise StrategySourceError(
                400, "unsupported_source", f"Unsupported file type: {file.mime_type!r}."
            )
    except GoogleApiError as exc:
        raise StrategySourceError(502, "drive_error", str(exc)) from exc
    return LoadedStrategy(
        name=file.name,
        kind=kind,
        ref=file.id,
        modified_time=file.modified_time,
        web_view_link=file.web_view_link,
        sheets=sheets,
    )


# ── Шлях ──────────────────────────────────────────────────────────────────────


def _host_root(settings: Settings) -> str:
    return settings.aiseotools_host_path.rstrip("/")


def _relative_to_root(text: str, settings: Settings) -> str:
    """Хостовий/контейнерний/відносний запис -> шлях відносно кореня AISeoTools.

    Абсолютний шлях поза обома коренями — `path_outside_aiseotools`: такого
    файлу в контейнері немає, а мовчки читати «схожий» ми не будемо.
    """
    if text.startswith("/"):
        for root in (_host_root(settings), str(settings.aiseotools_source_path).rstrip("/")):
            if text.startswith(root + "/"):
                return text[len(root) + 1 :]
        raise StrategySourceError(
            400,
            "path_outside_aiseotools",
            f"Only files inside {_host_root(settings)} can be read.",
        )
    return text[2:] if text.startswith("./") else text


def resolve_strategy_path(text: str, settings: Settings) -> tuple[Path, str]:
    """`(абсолютний контейнерний шлях, шлях відносно кореня)`; піднімає
    `StrategySourceError` на вихід за корінь, не-xlsx, відсутній файл, розмір."""
    relpath = _relative_to_root(text, settings)
    try:
        path = safe_resolve(settings.aiseotools_source_path, relpath)
    except PathTraversalError as exc:
        raise StrategySourceError(
            400,
            "path_outside_aiseotools",
            f"Only files inside {_host_root(settings)} can be read.",
        ) from exc
    host_path = f"{_host_root(settings)}/{relpath}"
    if path.suffix.lower() != ".xlsx":
        raise StrategySourceError(
            400, "unsupported_source", "Only .xlsx files can be read from a path."
        )
    if not path.is_file():
        raise StrategySourceError(404, "strategy_file_not_found", f"No file at {host_path}.")
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise StrategySourceError(
            413, "upload_too_large", f"The file exceeds {MAX_XLSX_BYTES // (1024 * 1024)} MB."
        )
    return path, relpath


def _load_path(text: str, settings: Settings) -> LoadedStrategy:
    path, relpath = resolve_strategy_path(text, settings)
    modified = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).isoformat()
    return LoadedStrategy(
        name=path.stem,
        kind="xlsx",
        ref=relpath,
        modified_time=modified,
        web_view_link=None,
        sheets=_read_xlsx(path.read_bytes()),
    )


def _read_xlsx(content: bytes) -> list[RawSheet]:
    # У read-only режимі openpyxl не віддає гіперпосилань клітинок, тож
    # `Cell.links` тут порожні: `framework_link` для xlsx — чесне `null`,
    # лише текст клітинки (task-102 §2.1).
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl кидає різні типи на биті файли
        raise StrategySourceError(
            400, "unreadable_strategy", "The file is not a readable .xlsx workbook."
        ) from exc
    sheets: list[RawSheet] = []
    for name in wb.sheetnames:
        rows: list[RawRow] = [
            tuple(Cell(None if v is None else str(v)) for v in row)
            for row in wb[name].iter_rows(values_only=True)
        ]
        sheets.append((name, rows))
    return sheets


# ── Локальні підказки ─────────────────────────────────────────────────────────


@dataclass(frozen=True)
class LocalStrategy:
    path_host: str
    name: str
    modified_at: datetime
    size: int


def list_local_strategies(settings: Settings) -> list[LocalStrategy]:
    """`out/*/*.xlsx` із «STRATEGY» в імені, найновіші першими. Живить
    `<datalist>` поля вводу — оператор не друкує 90-символьний шлях."""
    out_dir = settings.aiseotools_source_path / "out"
    if not out_dir.is_dir():
        return []
    found: list[LocalStrategy] = []
    for path in out_dir.glob("*/*.xlsx"):
        lowered = path.name.lower()
        if "strategy" not in lowered or any(m in lowered for m in _NOT_A_STRATEGY_MARKERS):
            continue
        try:
            info = path.stat()
        except OSError:
            continue
        found.append(
            LocalStrategy(
                path_host=f"{_host_root(settings)}/out/{path.parent.name}/{path.name}",
                name=path.name,
                modified_at=datetime.fromtimestamp(info.st_mtime, tz=UTC),
                size=info.st_size,
            )
        )
    found.sort(key=lambda f: f.modified_at, reverse=True)
    return found[:MAX_LOCAL_LISTING]
