"""Build the XLSX report (task-42 §5/§9 DoD, task-45 §2.3/§2.4): sheets
`Report` / `Data Sources` (PDF path only) / `Cross-check` (PDF path only) /
`Historical Data` / `Import Data`.

Ported from `AISeoTools/.../clientreport/xlsx_builder.py` and generalized from
a fixed 6-column monthly table to N columns (task-42 §2.4).

task-42 dropped the ported "Historical Data" sheet for having no equivalent
here; task-45 §1.5 brings it back as real data (`app/reports/history.py`) —
12 complete months with a live `=clicks/impressions` CTR formula, so a row
pasted into Google Sheets keeps working the way the hand-kept history did.
Both sources reach this module now: task-42 called `build_xlsx` only on the
PDF path, which is why an API run's Files section had nothing in it.

Number formats use the `[$-0409]` (US English locale) prefix so LibreOffice
renders commas/dots the same way regardless of the host's locale — the same
choice the ported module made.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.history import ReportHistory
from app.reports.table import ReportTable
from app.reports.verify import CrossCheckResult

_HEADER_BG = "1B5E47"
_HEADER_FG = "FFFFFF"
_ALT_ROW_BG = "F4F9F6"
_MISMATCH_BG = "F8D7DA"
_WITHIN_TOLERANCE_BG = "FFF3CD"

_FMT_INT = "[$-0409]#,##0"
_FMT_PCT = "[$-0409]0.00%"

# Where each metric physically sits inside the Reporting Ninja PDF — the
# "Data Sources" sheet's per-metric column (task-42 §5, PDF path only).
METRIC_SOURCES = {
    "users": "GA4 'Overall performance' -> 'Total users' row, current period",
    "organic_users": (
        "GA4 'First user default channel grouping' -> 'Organic Search' row -> 'Total users'"
    ),
    "direct_users": "GA4 'First user default channel grouping' -> 'Direct' row -> 'Total users'",
    "conversions": "GA4 'Overall performance' -> 'Key events' row, current period",
    "organic_conversions": (
        "GA4 'First user default channel grouping' -> 'Organic Search' row -> 'Key events'"
    ),
    "direct_conversions": (
        "GA4 'First user default channel grouping' -> 'Direct' row -> 'Key events'"
    ),
    "clicks": "GSC 'Overall performance' -> 'Clicks' row, current period",
    "impressions": "GSC 'Overall performance' -> 'Impressions' row, current period",
}


def _header_font() -> Font:
    return Font(name="Calibri", bold=True, color=_HEADER_FG, size=11)


def _cell_font(*, bold: bool = False) -> Font:
    return Font(name="Calibri", bold=bold, size=11)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _title_row(ws: Worksheet, text: str, total_cols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    cell = ws.cell(1, 1, text)
    cell.font = Font(name="Calibri", bold=True, size=13, color=_HEADER_FG)
    cell.fill = _fill(_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22


def _header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row, c, text)
        cell.font = _header_font()
        cell.fill = _fill(_HEADER_BG)
        cell.alignment = Alignment(
            horizontal="left" if c == 1 else "center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[row].height = 30


def _build_report_sheet(ws: Worksheet, client_name: str, table: ReportTable) -> None:
    ws.title = "Report"
    n_periods = len(table.periods)
    n_changes = len(table.change_labels)
    total_cols = 1 + n_periods + n_changes

    _title_row(ws, client_name, total_cols)
    _header_row(ws, 2, ["Metric", *(p.label for p in table.periods), *table.change_labels])

    column_of = {p.label: 2 + i for i, p in enumerate(table.periods)}
    for row_i, row in enumerate(table.rows, start=3):
        alt = row_i % 2 == 1
        bg = _fill(_ALT_ROW_BG) if alt else None

        label_cell = ws.cell(row_i, 1, row.label)
        label_cell.font = _cell_font()
        if bg:
            label_cell.fill = bg

        for col_i, value in enumerate(row.values, start=2):
            cell = ws.cell(row_i, col_i, value)
            cell.number_format = _FMT_INT
            cell.alignment = Alignment(horizontal="right")
            if bg:
                cell.fill = bg

        for j, (newer_label, older_label) in enumerate(table.change_pairs):
            col_i = 2 + n_periods + j
            newer_col, older_col = column_of[newer_label], column_of[older_label]
            formula = (
                f"=IFERROR(({get_column_letter(newer_col)}{row_i}-"
                f"{get_column_letter(older_col)}{row_i})/"
                f"{get_column_letter(older_col)}{row_i},\"\")"
            )
            cell = ws.cell(row_i, col_i, formula)
            cell.number_format = _FMT_PCT
            cell.alignment = Alignment(horizontal="right")
            if bg:
                cell.fill = bg

    ws.column_dimensions["A"].width = 24
    for c in range(2, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.freeze_panes = "B3"


def _build_sources_sheet(ws: Worksheet, client_name: str, table: ReportTable) -> None:
    ws.title = "Data Sources"
    total_cols = 2 + len(table.periods)
    wrap = Alignment(vertical="top", wrap_text=True)

    _title_row(ws, f"{client_name} — data sources (file + place in PDF)", total_cols)
    _header_row(ws, 2, ["Metric", "Where in the PDF", *(p.label for p in table.periods)])

    for row_i, row in enumerate(table.rows, start=3):
        bg = _fill(_ALT_ROW_BG) if row_i % 2 == 1 else None
        label_cell = ws.cell(row_i, 1, row.label)
        place_cell = ws.cell(row_i, 2, METRIC_SOURCES.get(row.key, ""))
        file_cells = [
            ws.cell(row_i, 3 + i, table.sources.get(p.label, "—"))
            for i, p in enumerate(table.periods)
        ]
        for cell in (label_cell, place_cell, *file_cells):
            cell.font = _cell_font()
            cell.alignment = wrap
            if bg:
                cell.fill = bg

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55
    for c in range(3, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 26
    ws.freeze_panes = "C3"


def _build_cross_check_sheet(
    ws: Worksheet, client_name: str, cross_check: CrossCheckResult
) -> None:
    ws.title = "Cross-check"
    headers = ["Period", "Metric", "PDF", "API", "Diff %", "Verdict"]
    total_cols = len(headers)

    verdict_bg = {"mismatch": _MISMATCH_BG, "within_tolerance": _WITHIN_TOLERANCE_BG}
    summary = f"{client_name} — Cross-check: {cross_check.status.upper()}"
    if cross_check.mismatches:
        summary += f" ({len(cross_check.mismatches)} of {len(cross_check.cells)} cells mismatched)"
    _title_row(ws, summary, total_cols)
    _header_row(ws, 2, headers)

    for row_i, cell_cmp in enumerate(cross_check.cells, start=3):
        bg = verdict_bg.get(cell_cmp.verdict) or (_ALT_ROW_BG if row_i % 2 == 1 else None)
        values = [
            cell_cmp.period_label,
            cell_cmp.metric_label,
            cell_cmp.pdf_value,
            cell_cmp.api_value,
            cell_cmp.diff_pct,
            cell_cmp.verdict,
        ]
        for col_i, value in enumerate(values, start=1):
            cell = ws.cell(row_i, col_i, value)
            cell.font = _cell_font()
            if col_i == 5 and value is not None:
                cell.number_format = _FMT_PCT
            if bg:
                cell.fill = _fill(bg) if isinstance(bg, str) else bg

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 20
    for letter in ("C", "D", "E", "F"):
        ws.column_dimensions[letter].width = 14
    ws.freeze_panes = "A3"


def _build_history_sheet(
    ws: Worksheet, client_name: str, history: ReportHistory | None
) -> None:
    """`Date · Clicks · Impressions · CTR · Conversions`, oldest month first.

    `CTR` is a live `=IFERROR(clicks/impressions,"")` formula, not a computed
    number: this sheet's whole purpose is to be pasted into the operator's
    Google Sheets history, where the row has to stay recalculable (task-45
    §2.4). `Conversions` is ORGANIC key events — the column header says so,
    because that is what the hand-kept history tracked.
    """
    ws.title = "Historical Data"
    headers = ["Date", "Clicks", "Impressions", "CTR", "Conversions (Organic)"]
    total_cols = len(headers)

    _title_row(ws, f"{client_name} — Historical Data (complete months)", total_cols)
    _header_row(ws, 2, headers)

    if history is None or not history.months:
        note = ws.cell(3, 1, "Historical data could not be collected for this run.")
        note.font = _cell_font()
        ws.column_dimensions["A"].width = 24
        return

    for row_i, month in enumerate(history.months, start=3):
        bg = _fill(_ALT_ROW_BG) if row_i % 2 == 1 else None
        label_cell = ws.cell(row_i, 1, month.label)
        label_cell.font = _cell_font()

        clicks_cell = ws.cell(row_i, 2, month.clicks)
        impressions_cell = ws.cell(row_i, 3, month.impressions)
        for cell in (clicks_cell, impressions_cell):
            cell.font = _cell_font()
            cell.number_format = _FMT_INT
            cell.alignment = Alignment(horizontal="right")

        ctr_cell = ws.cell(row_i, 4, f'=IFERROR(B{row_i}/C{row_i},"")')
        ctr_cell.font = _cell_font()
        ctr_cell.number_format = _FMT_PCT
        ctr_cell.alignment = Alignment(horizontal="right")

        conversions_cell = ws.cell(row_i, 5, month.organic_conversions)
        conversions_cell.font = _cell_font()
        conversions_cell.number_format = _FMT_INT
        conversions_cell.alignment = Alignment(horizontal="right")

        if bg:
            for col_i in range(1, total_cols + 1):
                ws.cell(row_i, col_i).fill = bg

    footnote_row = 3 + len(history.months) + 1
    notes = []
    if not history.has_gsc:
        notes.append("Clicks/Impressions/CTR unavailable — Google Search Console is not connected.")
    if not history.has_ga4:
        notes.append("Conversions unavailable — GA4 monthly channel data is not available.")
    for offset, text in enumerate(notes):
        cell = ws.cell(footnote_row + offset, 1, text)
        cell.font = _cell_font()

    ws.column_dimensions["A"].width = 20
    for letter in ("B", "C", "D", "E"):
        ws.column_dimensions[letter].width = 16
    ws.freeze_panes = "A3"


def _build_import_sheet(ws: Worksheet, table: ReportTable) -> None:
    ws.title = "Import Data"
    bold = Font(name="Calibri", bold=True, size=11)
    for c, text in enumerate(
        ["Metric", *(p.label for p in table.periods), *table.change_labels], start=1
    ):
        ws.cell(1, c, text).font = bold

    n_periods = len(table.periods)
    for row_i, row in enumerate(table.rows, start=2):
        ws.cell(row_i, 1, row.label)
        for col_i, value in enumerate(row.values, start=2):
            ws.cell(row_i, col_i, value)
        for j, change in enumerate(row.changes):
            ws.cell(row_i, 2 + n_periods + j, round(change, 6) if change is not None else None)

    ws.column_dimensions["A"].width = 24


_SLUG_RE = re.compile(r"[^\w-]+")


def build_xlsx(
    client_name: str,
    table: ReportTable,
    out_dir: Path,
    *,
    generated_at: date | None = None,
    cross_check: CrossCheckResult | None = None,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    slug = _SLUG_RE.sub("_", client_name.strip()).strip("_").upper() or "CLIENT"
    stamp = (generated_at or date.today()).isoformat()
    path = out_dir / f"client_report_{slug}_{stamp}.xlsx"

    wb = Workbook()
    ws_report = wb.active
    _build_report_sheet(ws_report, client_name, table)

    # Provenance/cross-check sheets are PDF-path only — the API path leaves
    # `table.sources`/`cross_check` empty/None, and a sheet of all-"—" rows
    # would be noise (mirrors the ported module's own skip-if-empty rule).
    if table.sources:
        _build_sources_sheet(wb.create_sheet(index=1), client_name, table)
    if cross_check is not None:
        _build_cross_check_sheet(wb.create_sheet(), client_name, cross_check)

    # Always present, on both paths (task-45 §2.4 DoD) — including when the
    # history query failed, where the sheet says so instead of vanishing.
    _build_history_sheet(wb.create_sheet(), client_name, table.history)
    _build_import_sheet(wb.create_sheet(), table)

    wb.save(path)
    return path
