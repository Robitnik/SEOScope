"""Читання таблиці аудиту: XLSX (з вибором вкладки) та CSV/TSV.

Чому двокрокова робота (`inspect_source` → `parse_rows`): вкладку й мапінг
колонок неможливо знати до того, як файл прочитано. Спершу віддаємо оператору
список вкладок, заголовки, зразок рядків і **автопропозицію** мапінгу; він
підтверджує або перебиває — і лише тоді парсимо все.

Legacy `.xls` (BIFF, Excel ≤2003) свідомо не підтримується: для нього потрібен
`xlrd`, якого немає в образі (нова залежність — зона task-01). Помилка
`UnsupportedSourceError` несе конкретну дію для оператора замість тихої
поломки; сучасні експорти (Google Sheets, Ahrefs, Semrush, Screaming Frog)
віддають `.xlsx`/`.csv`.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from enum import StrEnum
from typing import Any

from openpyxl import load_workbook

# Захист від «таблиці» на мільйон рядків і від клітинки-роману: перевищення
# показуємо оператору як `truncated`, а не ріжемо тихо.
MAX_ROWS = 20_000
MAX_CELL_LENGTH = 8_192
MAX_COLUMNS = 64
SAMPLE_ROWS = 5

_XLSX_MAGIC = b"PK\x03\x04"
_XLS_MAGIC = b"\xd0\xcf\x11\xe0"  # OLE2 container (legacy .xls)


class SourceKind(StrEnum):
    xlsx = "xlsx"
    csv = "csv"
    tsv = "tsv"


class SourceError(ValueError):
    """База для помилок читання джерела; `code` іде в `ErrorDTO.error_code`."""

    code = "unreadable_source"


class UnsupportedSourceError(SourceError):
    code = "unsupported_source"


class UnreadableSourceError(SourceError):
    code = "unreadable_source"


class ColumnMappingError(SourceError):
    code = "invalid_column_map"


# Поля, які інструмент вміє використати. `image_url` і `new_alt` обовʼязкові —
# без них немає ні що шукати, ні що копіювати.
REQUIRED_FIELDS = ("image_url", "new_alt")
OPTIONAL_FIELDS = ("page_url", "old_alt")
MAPPABLE_FIELDS = REQUIRED_FIELDS + OPTIONAL_FIELDS

# Синоніми заголовків, від найточнішого до найзагальнішого. Порядок — це
# приоритет: чим раніше в списку, тим вищий бал при автомапінгу.
_SYNONYMS: dict[str, tuple[str, ...]] = {
    "image_url": (
        "image url",
        "image link",
        "image src",
        "image address",
        "image path",
        "image file",
        "img url",
        "img src",
        "image",
        "img",
        "src",
        "picture",
        "photo",
        "file",
        "filename",
        "url kartynky",
        "url",
    ),
    "new_alt": (
        "new alt",
        "new alt text",
        "new alternative text",
        "suggested alt",
        "suggested alt text",
        "alt text new",
        "alt tag",
        "alt text",
        "alternative text",
        "alt",
    ),
    "page_url": (
        "page url",
        "page link",
        "page address",
        "post url",
        "source page",
        "found on",
        "used on",
        "landing page",
        "referring page",
        "location",
        "page",
    ),
    "old_alt": (
        "old alt",
        "old alt text",
        "current alt",
        "current alt text",
        "existing alt",
        "existing alt text",
    ),
}

_NON_ALNUM = re.compile(r"[^0-9a-zа-яіїєґ]+", re.IGNORECASE)


@dataclass(frozen=True)
class SheetPreview:
    """Одна вкладка (або єдиний «аркуш» CSV) очима оператора."""

    name: str
    headers: tuple[str, ...]
    sample_rows: tuple[tuple[str, ...], ...]
    header_row_index: int
    suggested_column_map: dict[str, int]
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class InspectedSource:
    kind: SourceKind
    sheets: tuple[SheetPreview, ...]
    default_sheet: str
    delimiter: str | None = None
    encoding: str | None = None


@dataclass
class ParsedRow:
    row_index: int
    image_url: str
    new_alt: str
    old_alt: str | None
    # Одна картинка живе на кількох сторінках, і аудити (Ahrefs «Used on
    # page(s)») пакують їх усі в ОДНУ клітинку через новий рядок/пробіл. Тому
    # це список, а не рядок: `page_url` (перша) лишається для сумісності, але
    # перевірка й посилання працюють з усіма.
    page_urls: list[str] = field(default_factory=list)
    extra: dict[str, str] = field(default_factory=dict)

    @property
    def page_url(self) -> str | None:
        return self.page_urls[0] if self.page_urls else None


@dataclass
class ParseResult:
    rows: list[ParsedRow]
    skipped_count: int
    missing_alt_count: int
    warnings: list[str]


def normalize_header(value: str) -> str:
    return _NON_ALNUM.sub(" ", value or "").strip().lower()


def detect_source_kind(filename: str, content: bytes) -> SourceKind:
    """Визначити формат за магією файлу, а не за розширенням (PLAN §13)."""
    lowered = (filename or "").lower()
    if content.startswith(_XLS_MAGIC) or lowered.endswith(".xls"):
        raise UnsupportedSourceError(
            "Legacy .xls (Excel 97-2003) is not supported. "
            "Re-save the sheet as .xlsx, or export it as CSV."
        )
    if content.startswith(_XLSX_MAGIC):
        if lowered.endswith((".xlsx", ".xlsm")) or not lowered:
            return SourceKind.xlsx
        # ZIP-контейнер із незнайомим розширенням: усе одно спробуємо як XLSX,
        # бо Google Sheets інколуди віддає файл без розширення.
        return SourceKind.xlsx
    if lowered.endswith(".tsv"):
        return SourceKind.tsv
    if lowered.endswith((".csv", ".txt")) or _looks_like_text(content):
        return SourceKind.csv
    raise UnsupportedSourceError(
        "Unsupported file. Upload .xlsx (Excel / Google Sheets export) or .csv / .tsv."
    )


def _looks_like_text(content: bytes) -> bool:
    head = content[:4096]
    if b"\x00" in head:
        return False
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            head.decode(encoding)
        except UnicodeDecodeError:
            continue
        return True
    return False


def _decode(content: bytes) -> tuple[str, str]:
    """Декодувати текстове джерело, не вигадуючи кодування молча."""
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1", errors="replace"), "latin-1"


def _sniff_delimiter(sample: str, kind: SourceKind) -> str:
    if kind is SourceKind.tsv:
        return "\t"
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        # Один стовпець або нестандартний файл: кома — найбезпечніший дефолт,
        # і оператор усе одно побачить прев'ю перед створенням звіту.
        return ","


def _clip(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value.strip()
    elif isinstance(value, float) and value.is_integer():
        # Google Sheets віддає числові клітинки як float — «12.0» у назві
        # файлу нікому не потрібне.
        text = str(int(value))
    else:
        text = str(value).strip()
    return text[:MAX_CELL_LENGTH]


def _matrix_from_csv(content: bytes, kind: SourceKind) -> tuple[list[list[str]], str, str]:
    text, encoding = _decode(content)
    if not text.strip():
        raise UnreadableSourceError("The file is empty.")
    delimiter = _sniff_delimiter(text[:8192], kind)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    matrix: list[list[str]] = []
    for row in reader:
        matrix.append([_clip(cell) for cell in row[:MAX_COLUMNS]])
        if len(matrix) > MAX_ROWS + 1:
            break
    return matrix, delimiter, encoding


def _matrix_from_xlsx(
    content: bytes, sheet: str | None, limit: int
) -> tuple[list[str], list[list[str]]]:
    """Прочитати вкладку в матрицю. `limit` — скільки рядків максимум."""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl кидає різне на битих архівах
        raise UnreadableSourceError(f"Could not read the workbook: {exc}") from exc
    try:
        names = list(workbook.sheetnames)
        if sheet is not None and sheet not in names:
            raise ColumnMappingError(f"Sheet {sheet!r} is not in this workbook.")
        worksheet = workbook[sheet] if sheet is not None else workbook[names[0]]
        matrix: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            matrix.append([_clip(cell) for cell in row[:MAX_COLUMNS]])
            if len(matrix) >= limit:
                break
        return names, matrix
    finally:
        workbook.close()


def _xlsx_hyperlinks(content: bytes, sheet: str, column: int, limit: int) -> dict[int, str]:
    """Витягти цілі гіперпосилань для однієї колонки.

    Потрібне лише для випадку, коли в колонці URL лежить не текст-URL, а
    посилання з підписом (типовий «красивий» аркуш із Google Sheets).
    `read_only`-режим гіперпосилання не віддає, тож тут — звичайне читання,
    і робимо його **тільки** як fallback, коли інакше даних немає.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception:  # pragma: no cover - вже прочитали read_only вище
        return {}
    try:
        if sheet not in workbook.sheetnames:
            return {}
        worksheet = workbook[sheet]
        found: dict[int, str] = {}
        for index, row in enumerate(worksheet.iter_rows(max_col=column + 1)):
            if index >= limit:
                break
            if len(row) <= column:
                continue
            link = row[column].hyperlink
            if link is not None and link.target:
                found[index] = str(link.target)[:MAX_CELL_LENGTH]
        return found
    finally:
        workbook.close()


def _xlsx_merged_fill(content: bytes, sheet: str, matrix: list[list[str]]) -> int:
    """Протягнути значення об'єднаних клітинок на всі їхні рядки.

    Експорти аудитів (зокрема Ahrefs) групують «одна картинка → багато
    сторінок», об'єднуючи клітинку URL по всій групі. У XLSX значення лежить
    лише у верхній лівій клітинці діапазону, решта — `None`, тож без цього
    кроку 90%+ рядків виглядали як «рядок без URL картинки» і чесно, але
    марно, потрапляли у `skipped`.

    Читання не-`read_only` (merged-діапазони недоступні в read-only) —
    навмисно **лише** як fallback, коли в колонці справді є пропуски.
    Повертає кількість заповнених клітинок.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception:  # pragma: no cover - read_only-прохід уже пройшов вище
        return 0
    try:
        if sheet not in workbook.sheetnames:
            return 0
        worksheet = workbook[sheet]
        filled = 0
        for merged in list(worksheet.merged_cells.ranges):
            top_row, left = merged.min_row - 1, merged.min_col - 1
            if top_row >= len(matrix) or left >= MAX_COLUMNS:
                continue
            source = matrix[top_row][left] if left < len(matrix[top_row]) else ""
            if not source:
                continue
            for row_index in range(top_row + 1, min(merged.max_row, len(matrix))):
                row = matrix[row_index]
                if left >= len(row):
                    row.extend([""] * (left + 1 - len(row)))
                if not row[left]:
                    row[left] = source
                    filled += 1
        return filled
    finally:
        workbook.close()


def _find_header_row(matrix: list[list[str]]) -> int:
    """Перший рядок, який виглядає як заголовок (≥2 непорожні клітинки).

    Експорти з Google Sheets часто мають рядок-назву звіту над таблицею.
    """
    for index, row in enumerate(matrix[:20]):
        if sum(1 for cell in row if cell) >= 2:
            return index
    return 0


def suggest_column_map(headers: tuple[str, ...]) -> dict[str, int]:
    """Автопропозиція мапінгу «поле → індекс колонки».

    Індекс, а не назва: експорти легко містять порожні або дубльовані
    заголовки, і тоді назва перестає бути ключем.
    """
    scored: list[tuple[int, str, int]] = []
    for column, header in enumerate(headers):
        normalized = normalize_header(header)
        if not normalized:
            continue
        for name, synonyms in _SYNONYMS.items():
            score = 0
            for index, synonym in enumerate(synonyms):
                if normalized == synonym:
                    score = 100 - index
                    break
                if normalized.startswith(synonym) or normalized.endswith(synonym):
                    score = max(score, 60 - index)
                elif synonym in normalized:
                    score = max(score, 40 - index)
            if score > 0:
                scored.append((score, name, column))

    scored.sort(key=lambda item: (-item[0], item[1], item[2]))
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for _score, name, column in scored:
        if name in mapping or column in used:
            continue
        mapping[name] = column
        used.add(column)
    return mapping


def inspect_source(filename: str, content: bytes) -> InspectedSource:
    """Крок 1: що це за файл, які вкладки, які заголовки, який мапінг пропонуємо."""
    kind = detect_source_kind(filename, content)
    if kind is SourceKind.xlsx:
        names, _ = _matrix_from_xlsx(content, None, limit=1)
        if not names:
            raise UnreadableSourceError("The workbook has no sheets.")
        previews = []
        for name in names:
            _, matrix = _matrix_from_xlsx(content, name, limit=SAMPLE_ROWS + 21)
            previews.append(_preview(name, matrix))
        default = _default_sheet(previews)
        return InspectedSource(kind=kind, sheets=tuple(previews), default_sheet=default)

    matrix, delimiter, encoding = _matrix_from_csv(content, kind)
    preview = _preview(kind.value.upper(), matrix[: SAMPLE_ROWS + 21])
    return InspectedSource(
        kind=kind,
        sheets=(preview,),
        default_sheet=preview.name,
        delimiter=delimiter,
        encoding=encoding,
    )


def _preview(name: str, matrix: list[list[str]]) -> SheetPreview:
    warnings: list[str] = []
    if not matrix or not any(any(cell for cell in row) for row in matrix):
        return SheetPreview(
            name=name,
            headers=(),
            sample_rows=(),
            header_row_index=0,
            suggested_column_map={},
            warnings=("empty_sheet",),
        )
    header_row_index = _find_header_row(matrix)
    headers = tuple(matrix[header_row_index])
    width = max((len(row) for row in matrix), default=0)
    headers = headers + ("",) * (width - len(headers))
    suggested = suggest_column_map(headers)
    if not suggested:
        warnings.append("header_row_not_detected")
    else:
        missing = [name_ for name_ in REQUIRED_FIELDS if name_ not in suggested]
        if missing:
            warnings.append("required_columns_not_detected")
    samples = tuple(
        tuple(row + [""] * (width - len(row)))
        for row in matrix[header_row_index + 1 : header_row_index + 1 + SAMPLE_ROWS]
    )
    return SheetPreview(
        name=name,
        headers=headers,
        sample_rows=samples,
        header_row_index=header_row_index,
        suggested_column_map=suggested,
        warnings=tuple(warnings),
    )


def _default_sheet(previews: list[SheetPreview]) -> str:
    for preview in previews:
        if all(name in preview.suggested_column_map for name in REQUIRED_FIELDS):
            return preview.name
    return previews[0].name


_URL_SPLIT = re.compile(r"[\s,;]+")
# Скільки сторінок використання лишаємо на рядок: більше — це вже не
# операторський список, а простирадло, і кожна ще й перевіряється мережею.
MAX_PAGE_URLS = 25


def split_page_urls(value: str) -> list[str]:
    """Розкласти клітинку «Used on page(s)» на окремі URL.

    Ahrefs та подібні пишуть усі сторінки в одну клітинку через новий рядок;
    без розбиття `page_url` ставав рядком на кілька URL — посилання й
    перевірка на такому просто не працювали.
    """
    if not value:
        return []
    seen: dict[str, None] = {}
    for chunk in _URL_SPLIT.split(value.strip()):
        candidate = chunk.strip().strip("<>\"'")
        if candidate.startswith(("http://", "https://", "//", "/")):
            seen.setdefault(candidate, None)
    if not seen and value.strip():
        # Не схоже на URL — лишаємо як є одним значенням, щоб оператор
        # побачив, що саме було в клітинці, замість тихої втрати.
        return [value.strip()]
    return list(seen)[:MAX_PAGE_URLS]


def _looks_like_url(value: str) -> bool:
    if not value:
        return False
    if value.startswith(("http://", "https://", "//", "/")):
        return True
    head = value.split("/", 1)[0]
    return "/" in value and "." in head


def parse_rows(
    filename: str,
    content: bytes,
    *,
    sheet: str | None,
    column_map: dict[str, int],
) -> ParseResult:
    """Крок 2: розпарсити все за підтвердженим мапінгом.

    Рядок без URL картинки — `skipped` (нічого не шукати). Рядок без нового
    alt лишається у звіті, але рахується окремо: оператор має бачити дірку в
    джерелі, а не втратити рядок.
    """
    for name in REQUIRED_FIELDS:
        if name not in column_map:
            raise ColumnMappingError(f"Column for {name!r} is not mapped.")

    kind = detect_source_kind(filename, content)
    warnings: list[str] = []
    if kind is SourceKind.xlsx:
        sheet_name = sheet
        if sheet_name is None:
            names, _ = _matrix_from_xlsx(content, None, limit=1)
            sheet_name = names[0]
        _, matrix = _matrix_from_xlsx(content, sheet_name, limit=MAX_ROWS + 21)
    else:
        matrix, _delimiter, _encoding = _matrix_from_csv(content, kind)
        sheet_name = None

    if not matrix:
        raise UnreadableSourceError("The sheet is empty.")
    header_row_index = _find_header_row(matrix)
    headers = list(matrix[header_row_index])
    body = matrix[header_row_index + 1 :]
    if len(body) > MAX_ROWS:
        body = body[:MAX_ROWS]
        warnings.append("row_cap_reached")

    url_column = column_map["image_url"]
    hyperlinks: dict[int, str] = {}
    # Об'єднані клітинки: значення лежить лише у верхньому рядку діапазону, тож
    # без цього кроку решта групи «одна картинка → багато сторінок» виглядала
    # як рядки без URL. Платимо за не-read_only прохід ЛИШЕ коли пропуски є.
    if kind is SourceKind.xlsx and sheet_name is not None:
        gaps = sum(1 for row in body if not _cell(row, url_column))
        if gaps:
            filled = _xlsx_merged_fill(content, sheet_name, matrix)
            if filled:
                warnings.append("merged_cells_expanded")

    non_empty = sum(1 for row in body if _cell(row, url_column))
    url_like = sum(1 for row in body if _looks_like_url(_cell(row, url_column)))
    if kind is SourceKind.xlsx and sheet_name is not None and non_empty and not url_like:
        # Уся колонка — підписи гіперпосилань, а не URL: дістаємо цілі.
        hyperlinks = _xlsx_hyperlinks(
            content, sheet_name, url_column, header_row_index + 1 + len(body)
        )
        if hyperlinks:
            warnings.append("image_urls_taken_from_hyperlinks")

    mapped_columns = set(column_map.values())
    rows: list[ParsedRow] = []
    skipped = 0
    missing_alt = 0
    for offset, row in enumerate(body):
        absolute_index = header_row_index + 1 + offset
        # Ціль гіперпосилання має пріоритет над текстом клітинки: `hyperlinks`
        # заповнюється лише тоді, коли в колонці лежать підписи, а не URL.
        image_url = hyperlinks.get(absolute_index) or _cell(row, url_column)
        if not image_url:
            skipped += 1
            continue
        new_alt = _cell(row, column_map.get("new_alt"))
        if not new_alt:
            missing_alt += 1
        extra = {
            _extra_label(headers, column): _cell(row, column)
            for column in range(len(row))
            if column not in mapped_columns and _cell(row, column)
        }
        rows.append(
            ParsedRow(
                # Номер рядка як його бачить оператор у таблиці (1-based).
                row_index=absolute_index + 1,
                image_url=image_url,
                new_alt=new_alt,
                old_alt=_cell(row, column_map.get("old_alt")) or None,
                page_urls=split_page_urls(_cell(row, column_map.get("page_url"))),
                extra=extra,
            )
        )

    if not rows:
        raise UnreadableSourceError(
            "No rows with an image URL were found. Check the column mapping and the sheet."
        )
    if missing_alt:
        warnings.append("rows_without_new_alt")
    return ParseResult(
        rows=rows, skipped_count=skipped, missing_alt_count=missing_alt, warnings=warnings
    )


def _extra_label(headers: list[str], column: int) -> str:
    """Назва незмапленої колонки; безіменну підписуємо як бачить її Excel-користувач."""
    if column < len(headers) and headers[column]:
        return headers[column]
    return f"Column {column + 1}"


def _cell(row: list[str], column: int | None) -> str:
    if column is None or column < 0 or column >= len(row):
        return ""
    return row[column]
