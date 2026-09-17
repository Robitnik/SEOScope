"""Excel-експорт для GA4 Events by Month (task-98 §2.6).

Книга будується з ТОГО САМОГО DTO, який читає екран (правило
`app/pageinventory/xlsx.py`): аркуш `Summary` — матриця `івент × місяць` за
`Event Count` порізу `All Users`, далі по аркушу на кожен обраний місяць у тому
ж порядку, що й вкладки в UI, і в тій самій формі, що й таблиця оператора:
три рядки на івент (All Users / Direct traffic / Organic traffic).

Два інваріанти, які тут легко зламати й дорого помітити:

- рядок `TOTAL` бере `Event count`/`Total users` із тоталів ПОРІЗУ (окремий
  запит без `eventName`), ніколи не складає колонку Users: сума перевищує
  реальний property-тотал у 3.9 раза на живих даних;
- місяць без даних або впалий теж отримує аркуш — з рядком причини. Аркуш,
  якого немає, читається як «нуль подій», а це неправда.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.ga4_events import Ga4EventMonthDTO, Ga4EventsReportDTO

_HEADER_BG = "434343"
_HEADER_FG = "FFFFFF"
_KEY_EVENT_BG = "D9EAD3"
_TOTAL_BG = "D9D9D9"
_PARTIAL_BG = "FFF2CC"
_EVENT_NAME_BG = "F3F3F3"

_FMT_INT = "#,##0"
_FMT_PCT = "0.00"
_FMT_RATIO = "0.00"

# Ті самі підписи, що в UI і в GA4 — провайдер і поріз називаються однаково
# скрізь (CLAUDE.md).
_SEGMENTS: tuple[tuple[str, str], ...] = (
    ("all", "All Users"),
    ("direct", "Direct traffic"),
    ("organic", "Organic traffic"),
)

_MONTH_HEADERS = [
    "Event Name",
    "Comparison",
    "Event Count",
    "Event Count %",
    "Total Users",
    "Total Users %",
    "Event count per user",
]
_MONTH_WIDTHS = [34, 17, 14, 14, 13, 14, 20]


def _f(value: Decimal | None) -> float | None:
    """Excel не має `Decimal`; `None` лишається `None` — порожня клітинка, не 0."""
    return None if value is None else float(value)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _header(ws: Worksheet, labels: list[str]) -> None:
    for index, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=index, value=label)
        cell.font = Font(bold=True, color=_HEADER_FG)
        cell.fill = _fill(_HEADER_BG)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _reason_sheet(ws: Worksheet, month: Ga4EventMonthDTO) -> None:
    _header(ws, _MONTH_HEADERS)
    reason = month.error or (
        "No events in this month." if month.state == "no_data" else "Unavailable."
    )
    cell = ws.cell(row=2, column=1, value=reason)
    cell.font = Font(italic=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_MONTH_HEADERS))
    _autosize(ws, _MONTH_WIDTHS)


def _write_month_sheet(ws: Worksheet, month: Ga4EventMonthDTO) -> None:
    if month.state != "ok":
        _reason_sheet(ws, month)
        return

    _header(ws, _MONTH_HEADERS)
    row_index = 2
    for row in month.rows:
        for key, label in _SEGMENTS:
            values = row.segments.get(key)
            cells = [
                # Ім'я івента стоїть лише в першому з трьох рядків — так само,
                # як в аркуші оператора: воно ідентифікує групу, а не кожен рядок.
                row.event_name if key == "all" else None,
                label,
                values.event_count if values else None,
                _f(values.event_share_pct if values else None),
                values.users if values else None,
                _f(values.user_share_pct if values else None),
                _f(values.per_user if values else None),
            ]
            formats = [None, None, _FMT_INT, _FMT_PCT, _FMT_INT, _FMT_PCT, _FMT_RATIO]
            for column_index, (value, number_format) in enumerate(
                zip(cells, formats, strict=True), start=1
            ):
                cell = ws.cell(row=row_index, column=column_index, value=value)
                if number_format and value is not None:
                    cell.number_format = number_format
                if key == "all" and column_index == 1:
                    cell.font = Font(bold=True)
                    cell.fill = _fill(_KEY_EVENT_BG if row.is_key_event else _EVENT_NAME_BG)
            row_index += 1

    totals_row = row_index
    for key, label in _SEGMENTS:
        totals = month.totals.get(key)
        cells = [
            "TOTAL (property)" if key == "all" else None,
            label,
            totals.event_count if totals else None,
            None,
            totals.users if totals else None,
            None,
            None,
        ]
        for column_index, value in enumerate(cells, start=1):
            cell = ws.cell(row=totals_row, column=column_index, value=value)
            cell.font = Font(bold=True)
            cell.fill = _fill(_TOTAL_BG)
            if column_index in (3, 5) and value is not None:
                cell.number_format = _FMT_INT
        totals_row += 1

    note = ws.cell(
        row=totals_row + 1,
        column=1,
        value=(
            "Total Users is a property-level figure per comparison — not the sum of the column "
            "above, because one user triggers many events. Percentages are shares of each "
            "comparison's own totals."
        ),
    )
    note.font = Font(italic=True, size=9)
    if month.is_partial:
        flag = ws.cell(
            row=totals_row + 2,
            column=1,
            value="Partial month — GA4 keeps revising these numbers for ~48 hours.",
        )
        flag.font = Font(italic=True, size=9)
        flag.fill = _fill(_PARTIAL_BG)
    _autosize(ws, _MONTH_WIDTHS)


def _write_summary_sheet(ws: Worksheet, report: Ga4EventsReportDTO) -> None:
    """Зведення по `All Users`: один рядок на івент, колонка на місяць.

    Порізи сюди не йдуть навмисно — три числа в клітинці зробили б таблицю
    нечитабельною; за ними оператор іде на аркуш місяця.
    """
    months = report.months_data
    header = ["Event", "Key event"] + [month.label for month in months]
    _header(ws, header)

    counts: dict[str, dict[str, int | None]] = {}
    key_events: set[str] = set()
    for month in months:
        for row in month.rows:
            segment = row.segments.get("all")
            counts.setdefault(row.event_name, {})[month.month] = (
                segment.event_count if segment else None
            )
            if row.is_key_event:
                key_events.add(row.event_name)

    newest = months[0].month if months else None
    ordered = sorted(
        counts,
        key=lambda name: (
            name not in key_events,
            -((counts[name].get(newest) if newest else None) or 0),
            name.lower(),
        ),
    )

    row_index = 2
    for name in ordered:
        is_key = name in key_events
        ws.cell(row=row_index, column=1, value=name)
        ws.cell(row=row_index, column=2, value="yes" if is_key else "")
        for column_index, month in enumerate(months, start=3):
            cell = ws.cell(
                row=row_index, column=column_index, value=counts[name].get(month.month)
            )
            cell.number_format = _FMT_INT
        if is_key:
            for column_index in range(1, len(header) + 1):
                ws.cell(row=row_index, column=column_index).fill = _fill(_KEY_EVENT_BG)
        row_index += 1

    ws.cell(row=row_index, column=1, value="TOTAL events (All Users)").font = Font(bold=True)
    for column_index, month in enumerate(months, start=3):
        totals = month.totals.get("all")
        cell = ws.cell(
            row=row_index, column=column_index, value=totals.event_count if totals else None
        )
        cell.font = Font(bold=True)
        cell.number_format = _FMT_INT
    for column_index in range(1, len(header) + 1):
        ws.cell(row=row_index, column=column_index).fill = _fill(_TOTAL_BG)

    ws.cell(
        row=row_index + 2,
        column=1,
        value=(
            f"{report.client_name or 'Client'} · {report.property_id} · "
            f"generated {report.created_at:%Y-%m-%d %H:%M UTC}"
        ),
    ).font = Font(italic=True, size=9)
    _autosize(ws, [34, 11] + [14] * len(months))


def build_workbook(report: Ga4EventsReportDTO) -> io.BytesIO:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    _write_summary_sheet(summary, report)
    for month in report.months_data:
        # Назва аркуша — ISO-місяць: сортується сама і завжди ≤31 символ.
        _write_month_sheet(workbook.create_sheet(title=month.month), month)
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
