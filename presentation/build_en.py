#!/usr/bin/env python3
"""Build the English The Yovth SEO Tools deck from the shared slide design."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation

import build_ua


TRANSLATIONS = {
    # Slide 1
    "SEO TOOLS для": "SEO TOOLS for",
    "контролю якості та": "quality control and",
    "продуктивності клієнтів": "client performance",
    "Я К І С Т Ь     ·     П Р О З О Р І С Т Ь     ·     Р Е З У Л Ь Т А Т": "Q U A L I T Y     ·     T R A N S P A R E N C Y     ·     R E S U L T S",
    # Slide 2
    "П Р О Б Л Е М А     →     Р І Ш Е Н Н Я": "P R O B L E M     →     S O L U T I O N",
    "Усі клієнти — на одній сторінці.": "All clients—on one page.",
    "Б У Л О   С К Л А Д Н О": "T H E   C H A L L E N G E",
    "Назва й домен": "Name and domain",
    "Notion → клієнт → Website → браузер → копіювання": "Notion → client → Website → browser → copy",
    "Статус і команда": "Status and team",
    "Хто активний? До якої групи належить? Треба було звіряти окремо.": "Who is active? Which group owns it? Each one had to be checked separately.",
    "ТЕПЕР": "NOW",
    "Усе потрібне — в одному місці.": "Everything you need—in one place.",
    "Г О Л О В Н А   С Т О Р І Н К А": "M A I N   P A G E",
    "КОПІЮВАННЯ В 1 КЛІК": "ONE-CLICK COPY",
    "ГРУПИ ТА СТАТУСИ": "GROUPS & STATUSES",
    "ШВИДКИЙ ПОШУК": "QUICK SEARCH",
    # Slide 3
    "С Т О Р І Н К А   К Л І Є Н Т А": "C L I E N T   P A G E",
    "Один клієнт — один робочий контекст.": "One client—one workspace.",
    "У С Е   Н А   О Д Н І Й   С Т О Р І Н Ц І": "E V E R Y T H I N G   O N   O N E   P A G E",
    "Профіль і статус": "Profile and status",
    "Домен, Active / Paused і health score": "Domain, Active / Paused, and health score",
    "Ключові KPI й перехід до історії": "Key KPIs and a direct link to history",
    "Помилки, sitemap, pages і DNS": "Issues, sitemap, pages, and DNS",
    "Контент": "Content",
    "Опубліковані матеріали та їх кількість": "Published content and item count",
    # Slide 4
    "П Р О Б Л Е М А     →     А В Т О М А Т И З А Ц І Я": "P R O B L E M     →     A U T O M A T I O N",
    "Повторювана задача більше не забирає години.": "Repetitive work no longer takes hours.",
    "В Р У Ч Н У": "M A N U A L L Y",
    "КІЛЬКА": "SEVERAL",
    "ГОДИН": "HOURS",
    "Збір і побудова історичної таблиці": "Collecting and building the historical table",
    "З А   Д О П О М О Г О Ю   Ш І": "W I T H   A I",
    "ХВИЛИН": "MINUTES",
    "Перший робочий результат": "First workable result",
    "СЕКУНДИ": "SECONDS",
    "Готовий повторюваний процес": "Ready, repeatable workflow",
    # Slide 5
    "Р І Ш Е Н Н Я     ·     M O N T H - O V E R - M O N T H": "S O L U T I O N     ·     M O N T H - O V E R - M O N T H",
    "Історичні дані — готова таблиця за 2 секунди.": "Historical data—a ready table in 2 seconds.",
    "Щ О   Г О Т О В О": "W H A T   I S   R E A D Y",
    "Історія за місяцями": "Monthly history",
    "Порівняння з попереднім місяцем": "Month-over-month comparison",
    "Динаміка показників ↑ / ↓": "Metric trends ↑ / ↓",
    "Файл .xlsx — в один клік": ".xlsx file—in one click",
    # Slide 6
    "К В А Р Т А Л Ь Н А   С Т Р А Т Е Г І Я     ·     П Р О Б Л Е М А": "Q U A R T E R L Y   S T R A T E G Y     ·     P R O B L E M",
    "200+ блогів — завеликий контекст для ручної перевірки.": "200+ blog posts—too much context for manual review.",
    "Повільно вручну": "Manual review is slow",
    "Зібрати всі теми й дати — довга ручна робота перед кожною стратегією.": "Collecting every topic and date takes a long time before each strategy.",
    "Ненадійно через ШІ": "AI is unreliable at scale",
    "На великому контексті модель може пропустити, змішати або вигадати записи.": "With a large context, the model may miss, mix up, or invent records.",
    "Ризик для нової стратегії": "Risk to the new strategy",
    "Нова тема може дублювати чи канібалізувати вже опублікований контент.": "A new topic may duplicate or cannibalize published content.",
    "S D S   ·   О Д И Н   К Л І Є Н Т": "S D S   ·   O N E   C L I E N T",
    "ОПУБЛІКОВАНИХ": "PUBLISHED",
    "ПОСТІВ": "POSTS",
    "Їх треба врахувати у новій стратегії": "They must be considered in the new strategy",
    # Slide 7
    "Р І Ш Е Н Н Я     ·     C L I E N T   B L O G S": "S O L U T I O N     ·     C L I E N T   B L O G S",
    "Усі теми блогів — приблизно за хвилину.": "Every blog topic—in about a minute.",
    "Т Е П Е Р   В И Д Н О": "N O W   V I S I B L E",
    "Тема, опис і word count": "Topic, description, and word count",
    "Дата та пряме посилання": "Date and direct link",
    "Пошук, сортування й експорт": "Search, sorting, and export",
    "Підбір топових тем без дублювання й канібалізації": "Choose top topics without duplication or cannibalization",
    # Slide 8
    "С Т Р А Т Е Г І Я     ·     П Р О Б Л Е М А": "S T R A T E G Y     ·     P R O B L E M",
    "300+ сторінок — це година ручного збору.": "300+ pages—an hour of manual collection.",
    "Sitemap — лише початок": "The sitemap is only the start",
    "Для стратегії потрібні URL, breadcrumbs і метадані кожної сторінки.": "A strategy needs the URL, breadcrumbs, and metadata for every page.",
    "Вручну — близько години": "Manually—about an hour",
    "Зібрати й структурувати сотні сторінок перед аналізом.": "Collect and structure hundreds of pages before analysis.",
    "Дорогий контекст для ШІ": "Expensive context for AI",
    "Сирі дані займають багато токенів і ускладнюють надійну перевірку.": "Raw data consumes many tokens and makes reliable validation harder.",
    "СТОРІНОК": "PAGES",
    "СТОРІНКИ": "PAGES",
    "Їх треба врахувати у стратегії": "They must be considered in the strategy",
    # Slide 9
    "Р І Ш Е Н Н Я     ·     S I T E   A R C H I T E C T U R E": "S O L U T I O N     ·     S I T E   A R C H I T E C T U R E",
    "Sitemap → готова структура сайту.": "Sitemap → a ready site structure.",
    "В   О Д Н І Й   Т А Б Л И Ц І": "I N   O N E   T A B L E",
    "URL і breadcrumbs": "URL and breadcrumbs",
    "Title, description і keywords": "Title, description, and keywords",
    "Пошук та Export to Excel": "Search and Export to Excel",
    "−1 ГОДИНА": "−1 HOUR",
    "на підготовку стратегії": "preparing the strategy",
    # Slide 10
    "Р І Ш Е Н Н Я     ·     S E O   Q A": "S O L U T I O N     ·     S E O   Q A",
    "SEO QA показує, що оновлювати першими.": "SEO QA shows what to update first.",
    "З Н А Х О Д И Т Ь": "F I N D S",
    "Дублікати метаданих": "Duplicate metadata",
    "Порожні або слабкі поля": "Empty or weak fields",
    "Calibrated / Strict перевірка": "Calibrated / Strict validation",
    "ПРІОРИТЕТИ": "PRIORITIES",
    "Сторінки, що точно потребують оновлення": "Pages that definitely need an update",
    # Slide 11
    "Т Е Х Н І Ч Н И Й   S E O     ·     П Р О Б Л Е М А": "T E C H N I C A L   S E O     ·     P R O B L E M",
    "Банальні помилки повторюються на сотнях сторінок.": "Basic issues repeat across hundreds of pages.",
    "Їх легко пропустити": "Easy to miss",
    "Alt, redirects, Open Graph, title і description — прості, але масові проблеми.": "Alt text, redirects, Open Graph, titles, and descriptions—simple but widespread issues.",
    "Одна причина → сотні findings": "One cause → hundreds of findings",
    "Шаблонна помилка автоматично повторюється на багатьох URL.": "A template issue automatically repeats across many URLs.",
    "Стан сайту постійно змінюється": "The site keeps changing",
    "Потрібно бачити проблеми, які існують прямо зараз, а не в старому звіті.": "We need to see issues that exist now, not in an old report.",
    "259 великих зображень": "259 large images",
    # Slide 12
    "Р І Ш Е Н Н Я     ·     S I T E   C R A W L   I S S U E S": "S O L U T I O N     ·     S I T E   C R A W L   I S S U E S",
    "Що зламано прямо зараз — за один crawl.": "What is broken right now—in a single crawl.",
    "Д А Є   В І Д П О В І Д Ь": "A N S W E R S",
    "Severity і масштаб проблеми": "Severity and issue scale",
    "Готова інструкція Fix": "Ready-to-use fix instructions",
    "Re-check після виправлення": "Re-check after correction",
    "Виміряно crawler’ом — без здогадок ШІ": "Measured by the crawler—no AI guesses",
    # Slide 13
    "M O N I T O R     ·     П Р О Б Л Е М А": "M O N I T O R     ·     P R O B L E M",
    "14 сайтів і 2000+ сторінок не перевірити вручну.": "14 sites and 2,000+ pages cannot be checked manually.",
    "Зміни відбуваються щодня": "Changes happen every day",
    "Різні люди публікують і редагують сторінки на всіх клієнтських сайтах.": "Different people publish and edit pages across all client sites.",
    "Критичну помилку легко не помітити": "Critical errors are easy to miss",
    "Noindex, broken URL у sitemap або redirect на живій сторінці.": "Noindex, a broken URL in the sitemap, or a redirect on a live page.",
    "Один фахівець не встигає": "One specialist cannot keep up",
    "Постійно контролювати весь портфель фізично неможливо.": "Continuously monitoring the entire portfolio is physically impossible.",
    "Щ О Д Е Н Н І   З М І Н И": "D A I L Y   C H A N G E S",
    "САЙТІВ": "SITES",
    "ручний контроль неможливий": "manual control is impossible",
    # Slide 14
    "Р І Ш Е Н Н Я     ·     2 4 / 7   M O N I T O R I N G": "S O L U T I O N     ·     2 4 / 7   M O N I T O R I N G",
    "Топові сторінки перевіряються кожні 2 години.": "Top pages are checked every 2 hours.",
    "З М І Н А   Н А   С А Й Т І": "S I T E   C H A N G E",
    "Помилка з’явилась": "An issue appears",
    "ГОДИНИ": "HOURS",
    "Автоматичний обхід": "Automated crawl",
    "К Р И Т И Ч Н О": "C R I T I C A L",
    "Monitor повідомляє": "Monitor sends an alert",
    "К О М А Н Д А": "T E A M",
    "Вчасне виправлення": "Fix it in time",
    # Slide 15
    "Robots, sitemap, redirects і pages — в одному місці.": "Robots, sitemap, redirects, and pages—in one place.",
    "М О Н І Т О Р И Т Ь": "M O N I T O R S",
    "Robots.txt і доступ Googlebot": "Robots.txt and Googlebot access",
    "Sitemap: OK / redirect / broken": "Sitemap: OK / redirect / broken",
    "Status, indexability і canonical": "Status, indexability, and canonical",
    "Поточний стан і ручний запуск перевірки": "Current state and manual check",
    # Slide 16
    "M O N I T O R     ·     І С Т О Р І Я   З М І Н": "M O N I T O R     ·     C H A N G E   H I S T O R Y",
    "Знаємо дату й приблизний час кожної зміни.": "We know when every change happened.",
    "Щ О   З Б Е Р І Г А Є Т Ь С Я": "W H A T   I S   S A V E D",
    "Точна дата фіксації": "Exact detection date",
    "Коли Monitor записав нову версію": "When Monitor recorded a new version",
    "Часове вікно появи": "Time window of appearance",
    "Між двома послідовними перевірками": "Between two consecutive checks",
    "Копії before / after": "Before / after copies",
    "Обидві версії можна відкрити й порівняти": "Open and compare both versions",
    "МАЙЖЕ ВСІ ПУБЛІЧНІ ЗМІНИ САЙТУ": "NEARLY ALL PUBLIC SITE CHANGES",
    # Slide 17
    "C L I E N T   T R E N D S     ·     П Р О Б Л Е М А": "C L I E N T   T R E N D S     ·     P R O B L E M",
    "Щотижневий контроль трафіку — ручний квест.": "Weekly traffic checks were a manual quest.",
    "Код для входу в GSC": "Code to access GSC",
    "Отримати код, підтвердити доступ — і лише тоді відкрити дані.": "Get the code, confirm access—and only then open the data.",
    "Один акаунт на всю компанію": "One account for the whole company",
    "Усі клієнти в одному GSC-акаунті: якщо він зайнятий, треба чекати доступу.": "All clients are in one GSC account: if someone is using it, everyone else waits.",
    "Періоди й цифри — вручну": "Periods and numbers—manual",
    "Змінювати діапазони та самостійно зіставляти показники.": "Change date ranges and compare metrics manually.",
    "К О Ж Н О Г О   Т И Ж Н Я": "E V E R Y   W E E K",
    "КОД": "CODE",
    "КЛІЄНТ × 14": "CLIENT × 14",
    "І ЗНОВУ ІНШИЙ ПЕРІОД": "THEN ANOTHER PERIOD",
    # Slide 18
    "Р І Ш Е Н Н Я     ·     C L I E N T   T R E N D S": "S O L U T I O N     ·     C L I E N T   T R E N D S",
    "Усі клієнти й періоди — на одному графіку.": "All clients and periods—on one chart.",
    "У С Е   Н А   О Д Н О М У   Е К Р А Н І": "O N E   S C R E E N",
    "Усі клієнти поруч": "All clients side by side",
    "Періоди від 7 до 365 днів": "Periods from 7 to 365 days",
    "Без очікування доступу до GSC": "No waiting for GSC access",
    # Slide 19
    "GSC + GA4 — одна точка контролю.": "GSC + GA4—one place to monitor.",
    "1 ПЕРІОД": "1 PERIOD",
    "УСІ КЛІЄНТИ": "ALL CLIENTS",
    "1 ГРАФІК": "1 CHART",
    "Одразу видно, хто росте, а хто потребує уваги цього тижня.": "See who is growing—and who needs attention this week.",
    # Slide 20
    "W E B   A R C H I V E     ·     П Р О Б Л Е М А": "W E B   A R C H I V E     ·     P R O B L E M",
    "Дані старого сайту були — візуальної копії не було.": "We had the old-site data—but no visual copy.",
    "Структура була — вигляду не було": "The structure survived. The visuals did not.",
    "URL і sitemap збереглися, але старе меню візуально вже не відновити.": "URLs and the sitemap survived, but we could no longer see the old menu.",
    "Контент був неповний": "Content was incomplete",
    "Метадані не показували, що саме бачив користувач на сторінці.": "Metadata did not show what users actually saw on the page.",
    "Публічний архів не допоміг": "The public archive did not help",
    "Версія за грудень 2025 була застарілою й відтворювалася з помилками.": "The December 2025 version was outdated and replayed with errors.",
    "П І С Л Я   М І Г Р А Ц І Ї": "A F T E R   M I G R A T I O N",
    "ЗБЕРЕГЛОСЯ": "PRESERVED",
    "НЕМОЖЛИВО ПОБАЧИТИ": "COULD NOT SEE",
    # Slide 21
    "Р І Ш Е Н Н Я     ·     W E B   A R C H I V E": "S O L U T I O N     ·     W E B   A R C H I V E",
    "Власний архів зберігає весь портфель.": "Our own archive preserves the entire portfolio.",
    "З Б Е Р І Г А Є": "P R E S E R V E S",
    "7 850 URL у контролі": "7,850 URLs tracked",
    "6 926 версій сторінок": "6,926 page versions",
    "25 822 унікальні файли": "25,822 unique files",
    "Два автоматичні обходи щодня": "Two automated captures every day",
    # Slide 22
    "Архівну версію можна переглядати як живий сайт.": "Open an archived version like the live site.",
    "М О Ж Н А   П Е Р Е В І Р И Т И": "Y O U   C A N   V E R I F Y",
    "Меню, контент і дизайн": "Menu, content, and design",
    "Meta й View source": "Meta and View source",
    "Будь-який URL — вручну": "Any URL—manually",
}


def translate_deck(path: Path) -> None:
    prs = Presentation(path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text in TRANSLATIONS:
                        run.text = TRANSLATIONS[run.text]

    untranslated = []
    for slide_index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if hasattr(shape, "text") and re.search(r"[А-Яа-яІіЇїЄєҐґ]", shape.text):
                untranslated.append((slide_index, shape.text))
    if untranslated:
        details = "\n".join(f"slide {number}: {text!r}" for number, text in untranslated)
        raise RuntimeError(f"Untranslated Ukrainian text remains:\n{details}")

    prs.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--dashboard-image", type=Path, default=Path("tests/e2e/frontend/screenshots/auth-viewer-dashboard-1440.png"))
    parser.add_argument("--client-overview-image", type=Path, default=Path("presentation/assets/slide-03-client-overview.png"))
    parser.add_argument("--monthly-table-image", type=Path, default=Path("presentation/assets/slide-05-monthly-table.png"))
    parser.add_argument("--blogs-table-image", type=Path, default=Path("presentation/assets/slide-07-blogs-table.png"))
    parser.add_argument("--page-inventory-image", type=Path, default=Path("presentation/assets/slide-09-page-inventory.png"))
    parser.add_argument("--seo-qa-image", type=Path, default=Path("presentation/assets/slide-10-seo-qa.png"))
    parser.add_argument("--site-crawl-image", type=Path, default=Path("presentation/assets/slide-12-site-crawl-issues.png"))
    parser.add_argument("--monitor-overview-image", type=Path, default=Path("presentation/assets/slide-15-monitor-overview.png"))
    parser.add_argument("--monitor-history-image", type=Path, default=Path("presentation/assets/slide-16-monitor-history.png"))
    parser.add_argument("--trends-image", type=Path, default=Path("presentation/assets/slide-18-client-trends.png"))
    parser.add_argument("--archive-overview-image", type=Path, default=Path("presentation/assets/slide-21-web-archive.png"))
    parser.add_argument("--archive-view-image", type=Path, default=Path("presentation/assets/slide-22-archive-view.png"))
    args = parser.parse_args()

    build_ua.build(
        args.reference,
        args.output,
        args.dashboard_image,
        args.client_overview_image,
        args.monthly_table_image,
        args.blogs_table_image,
        args.page_inventory_image,
        args.seo_qa_image,
        args.site_crawl_image,
        args.monitor_overview_image,
        args.monitor_history_image,
        args.trends_image,
        args.archive_overview_image,
        args.archive_view_image,
    )
    translate_deck(args.output)


if __name__ == "__main__":
    main()
