#!/usr/bin/env python3
"""Build the Ukrainian The Yovth SEO Tools presentation slide by slide."""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BLACK = RGBColor(0x21, 0x21, 0x21)
ORANGE = RGBColor(0xFF, 0x4A, 0x17)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
LIGHT_GRAY = RGBColor(0xDD, 0xDD, 0xDD)


def remove_slide(prs: Presentation, index: int) -> None:
    """Remove a slide and its relationship from a python-pptx presentation."""

    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def set_text(
    shape,
    lines: list[str],
    *,
    size: float,
    color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    line_spacing: float = 1.0,
    font_name: str = "Arial",
) -> None:
    """Replace text while applying deterministic typography."""

    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True

    for index, text in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = text
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_textbox(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    lines: list[str],
    *,
    size: float,
    color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    line_spacing: float = 1.0,
    margin: float = 0.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.text_frame.margin_left = Inches(margin)
    shape.text_frame.margin_right = Inches(margin)
    shape.text_frame.margin_top = Inches(margin)
    shape.text_frame.margin_bottom = Inches(margin)
    set_text(
        shape,
        lines,
        size=size,
        color=color,
        bold=bold,
        align=align,
        line_spacing=line_spacing,
    )
    return shape


def add_problem_card(slide, *, y: float, number: str, title: str, body: str) -> None:
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.60),
        Inches(y),
        Inches(3.05),
        Inches(1.08),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LIGHT_GRAY
    card.line.width = Pt(1)
    card.adjustments[0] = 0.12

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.78),
        Inches(y + 0.18),
        Inches(0.36),
        Inches(0.36),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = ORANGE
    badge.line.fill.background()
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge.text_frame.margin_left = 0
    badge.text_frame.margin_right = 0
    badge.text_frame.margin_top = 0
    badge.text_frame.margin_bottom = 0
    set_text(badge, [number], size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(
        slide,
        1.28,
        y + 0.15,
        2.15,
        0.28,
        [title],
        size=13,
        color=BLACK,
        bold=True,
    )
    add_textbox(
        slide,
        0.78,
        y + 0.58,
        2.67,
        0.36,
        [body],
        size=9.5,
        color=GRAY,
        line_spacing=1.05,
    )


def add_feature_chip(slide, *, x: float, width: float, text: str) -> None:
    chip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(4.45),
        Inches(width),
        Inches(0.48),
    )
    chip.fill.solid()
    chip.fill.fore_color.rgb = WHITE
    chip.line.color.rgb = LIGHT_GRAY
    chip.line.width = Pt(1)
    chip.adjustments[0] = 0.15
    chip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    chip.text_frame.margin_left = Inches(0.05)
    chip.text_frame.margin_right = Inches(0.05)
    chip.text_frame.margin_top = 0
    chip.text_frame.margin_bottom = 0
    set_text(chip, [text], size=8.5, color=BLACK, bold=True, align=PP_ALIGN.CENTER)


def add_footer(slide, number: int) -> None:
    add_textbox(
        slide,
        0.68,
        5.25,
        4.0,
        0.14,
        ["The Yovth  ·  SEO Tools"],
        size=6.5,
        color=RGBColor(0xA8, 0xA8, 0xA8),
    )
    add_textbox(
        slide,
        9.15,
        5.25,
        0.20,
        0.14,
        [str(number)],
        size=6.5,
        color=RGBColor(0xA8, 0xA8, 0xA8),
        align=PP_ALIGN.RIGHT,
    )


def crop_dashboard(source: Path, destination: Path) -> None:
    """Crop the dashboard to the controls and first visible client row."""

    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        cropped = image.crop((20, 300, 1420, 990))
        cropped.save(destination, optimize=True)


def add_slide_two(prs: Presentation, dashboard_image: Path, asset_dir: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])

    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["П Р О Б Л Е М А     →     Р І Ш Е Н Н Я"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.55,
        ["Усі клієнти — на одній сторінці."],
        size=27,
        color=BLACK,
        bold=True,
    )

    add_textbox(
        slide,
        0.60,
        1.40,
        3.05,
        0.18,
        ["Б У Л О   С К Л А Д Н О"],
        size=8.5,
        color=ORANGE,
        bold=True,
    )
    add_problem_card(
        slide,
        y=1.68,
        number="1",
        title="Назва й домен",
        body="Notion → клієнт → Website → браузер → копіювання",
    )
    add_problem_card(
        slide,
        y=2.90,
        number="2",
        title="Статус і команда",
        body="Хто активний? До якої групи належить? Треба було звіряти окремо.",
    )

    solution = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.60),
        Inches(4.12),
        Inches(3.05),
        Inches(0.81),
    )
    solution.fill.solid()
    solution.fill.fore_color.rgb = ORANGE
    solution.line.color.rgb = ORANGE
    solution.adjustments[0] = 0.12
    solution.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    solution.text_frame.margin_left = Inches(0.18)
    solution.text_frame.margin_right = Inches(0.18)
    solution.text_frame.margin_top = Inches(0.08)
    solution.text_frame.margin_bottom = Inches(0.08)
    set_text(
        solution,
        ["ТЕПЕР", "Усе потрібне — в одному місці."],
        size=10.5,
        color=WHITE,
        bold=True,
        line_spacing=1.05,
    )

    add_textbox(
        slide,
        3.98,
        1.40,
        5.40,
        0.18,
        ["Г О Л О В Н А   С Т О Р І Н К А"],
        size=8.5,
        color=ORANGE,
        bold=True,
    )

    crop_path = asset_dir / "slide-02-dashboard.png"
    crop_dashboard(dashboard_image, crop_path)

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3.98),
        Inches(1.68),
        Inches(5.40),
        Inches(2.66),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(crop_path),
        Inches(4.00),
        Inches(1.70),
        width=Inches(5.36),
        height=Inches(2.62),
    )

    add_feature_chip(slide, x=3.98, width=1.62, text="КОПІЮВАННЯ В 1 КЛІК")
    add_feature_chip(slide, x=5.72, width=1.80, text="ГРУПИ ТА СТАТУСИ")
    add_feature_chip(slide, x=7.64, width=1.74, text="ШВИДКИЙ ПОШУК")

    add_footer(slide, len(prs.slides))


def add_overview_item(slide, *, y: float, number: str, title: str, body: str) -> None:
    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.68),
        Inches(y),
        Inches(0.33),
        Inches(0.33),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = ORANGE
    badge.line.fill.background()
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge.text_frame.margin_left = 0
    badge.text_frame.margin_right = 0
    badge.text_frame.margin_top = 0
    badge.text_frame.margin_bottom = 0
    set_text(badge, [number], size=9.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(slide, 1.18, y - 0.01, 2.60, 0.25, [title], size=12.5, color=BLACK, bold=True)
    add_textbox(slide, 1.18, y + 0.28, 2.62, 0.30, [body], size=8.5, color=GRAY, line_spacing=1.04)


def add_slide_three(prs: Presentation, client_overview_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["С Т О Р І Н К А   К Л І Є Н Т А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Один клієнт — один робочий контекст."],
        size=26,
        color=BLACK,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        1.43,
        3.05,
        0.18,
        ["У С Е   Н А   О Д Н І Й   С Т О Р І Н Ц І"],
        size=8.5,
        color=ORANGE,
        bold=True,
    )

    add_overview_item(
        slide,
        y=1.79,
        number="1",
        title="Профіль і статус",
        body="Домен, Active / Paused і health score",
    )
    add_overview_item(
        slide,
        y=2.55,
        number="2",
        title="Monthly performance",
        body="Ключові KPI й перехід до історії",
    )
    add_overview_item(
        slide,
        y=3.31,
        number="3",
        title="Site monitor",
        body="Помилки, sitemap, pages і DNS",
    )
    add_overview_item(
        slide,
        y=4.07,
        number="4",
        title="Контент",
        body="Опубліковані матеріали та їх кількість",
    )

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3.92),
        Inches(1.48),
        Inches(0.035),
        Inches(3.64),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ORANGE
    accent.line.fill.background()

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(4.07),
        Inches(1.48),
        Inches(5.30),
        Inches(3.66),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(client_overview_image),
        Inches(4.09),
        Inches(1.50),
        width=Inches(5.26),
        height=Inches(3.62),
    )
    add_footer(slide, len(prs.slides))


def add_time_card(
    slide,
    *,
    x: float,
    label: str,
    value_lines: list[str],
    body: str,
    highlight: bool = False,
) -> None:
    background = ORANGE if highlight else WHITE
    foreground = WHITE if highlight else BLACK
    secondary = WHITE if highlight else GRAY
    line_color = ORANGE if highlight else LIGHT_GRAY

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(1.75),
        Inches(2.55),
        Inches(2.55),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = background
    card.line.color.rgb = line_color
    card.line.width = Pt(1.2)
    card.adjustments[0] = 0.10

    add_textbox(slide, x + 0.20, 2.00, 2.15, 0.20, [label], size=9, color=ORANGE if not highlight else WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        x + 0.18,
        2.35,
        2.19,
        0.93,
        value_lines,
        size=27,
        color=foreground,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=0.92,
    )
    add_textbox(
        slide,
        x + 0.25,
        3.60,
        2.05,
        0.35,
        [body],
        size=9.5,
        color=secondary,
        align=PP_ALIGN.CENTER,
        line_spacing=1.04,
    )


def add_slide_four(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["П Р О Б Л Е М А     →     А В Т О М А Т И З А Ц І Я"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Повторювана задача більше не забирає години."],
        size=25,
        color=BLACK,
        bold=True,
    )

    add_time_card(
        slide,
        x=0.62,
        label="В Р У Ч Н У",
        value_lines=["КІЛЬКА", "ГОДИН"],
        body="Збір і побудова історичної таблиці",
    )
    add_textbox(slide, 3.27, 2.65, 0.35, 0.35, ["→"], size=23, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_time_card(
        slide,
        x=3.72,
        label="З А   Д О П О М О Г О Ю   Ш І",
        value_lines=["15", "ХВИЛИН"],
        body="Перший робочий результат",
    )
    add_textbox(slide, 6.37, 2.65, 0.35, 0.35, ["→"], size=23, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_time_card(
        slide,
        x=6.82,
        label="S E O   T O O L",
        value_lines=["2", "СЕКУНДИ"],
        body="Готовий повторюваний процес",
        highlight=True,
    )

    add_footer(slide, len(prs.slides))


def add_benefit(slide, *, y: float, text: str) -> None:
    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(7.12),
        Inches(y),
        Inches(0.28),
        Inches(0.28),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = ORANGE
    badge.line.fill.background()
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge.text_frame.margin_left = 0
    badge.text_frame.margin_right = 0
    badge.text_frame.margin_top = 0
    badge.text_frame.margin_bottom = 0
    set_text(badge, ["✓"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 7.55, y - 0.01, 1.80, 0.38, [text], size=10.5, color=BLACK, bold=True, line_spacing=1.02)


def add_slide_five(prs: Presentation, monthly_table_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     M O N T H - O V E R - M O N T H"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Історичні дані — готова таблиця за 2 секунди."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(monthly_table_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["Щ О   Г О Т О В О"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Історія за місяцями")
    add_benefit(slide, y=2.65, text="Порівняння з попереднім місяцем")
    add_benefit(slide, y=3.45, text="Динаміка показників ↑ / ↓")

    export = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.18),
        Inches(2.30),
        Inches(0.90),
    )
    export.fill.solid()
    export.fill.fore_color.rgb = ORANGE
    export.line.color.rgb = ORANGE
    export.adjustments[0] = 0.10
    export.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    export.text_frame.margin_left = Inches(0.12)
    export.text_frame.margin_right = Inches(0.12)
    export.text_frame.margin_top = Inches(0.07)
    export.text_frame.margin_bottom = Inches(0.07)
    set_text(
        export,
        ["EXPORT TO EXCEL", "Файл .xlsx — в один клік"],
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.10,
    )
    add_footer(slide, len(prs.slides))


def add_blog_problem(slide, *, y: float, number: str, title: str, body: str) -> None:
    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.72),
        Inches(y + 0.07),
        Inches(0.34),
        Inches(0.34),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = ORANGE
    badge.line.fill.background()
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge.text_frame.margin_left = 0
    badge.text_frame.margin_right = 0
    badge.text_frame.margin_top = 0
    badge.text_frame.margin_bottom = 0
    set_text(badge, [number], size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(slide, 1.22, y, 4.75, 0.26, [title], size=13, color=BLACK, bold=True)
    add_textbox(slide, 1.22, y + 0.32, 4.80, 0.42, [body], size=9.5, color=GRAY, line_spacing=1.04)


def add_slide_six(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["К В А Р Т А Л Ь Н А   С Т Р А Т Е Г І Я     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["200+ блогів — завеликий контекст для ручної перевірки."],
        size=24,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Повільно вручну",
        body="Зібрати всі теми й дати — довга ручна робота перед кожною стратегією.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Ненадійно через ШІ",
        body="На великому контексті модель може пропустити, змішати або вигадати записи.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Ризик для нової стратегії",
        body="Нова тема може дублювати чи канібалізувати вже опублікований контент.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.88, 2.34, 0.20, ["S D S   ·   О Д И Н   К Л І Є Н Т"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.75, 2.27, 2.40, 0.88, ["248"], size=62, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.34, 2.34, 0.48, ["ОПУБЛІКОВАНИХ", "ПОСТІВ"], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=0.95)
    add_textbox(
        slide,
        6.82,
        4.28,
        2.26,
        0.35,
        ["Їх треба врахувати у новій стратегії"],
        size=9.5,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        line_spacing=1.04,
    )
    add_footer(slide, len(prs.slides))


def add_slide_seven(prs: Presentation, blogs_table_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     C L I E N T   B L O G S"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Усі теми блогів — приблизно за хвилину."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(blogs_table_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["Т Е П Е Р   В И Д Н О"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Тема, опис і word count")
    add_benefit(slide, y=2.68, text="Дата та пряме посилання")
    add_benefit(slide, y=3.42, text="Пошук, сортування й експорт")

    export = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    export.fill.solid()
    export.fill.fore_color.rgb = ORANGE
    export.line.color.rgb = ORANGE
    export.adjustments[0] = 0.10
    export.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    export.text_frame.margin_left = Inches(0.10)
    export.text_frame.margin_right = Inches(0.10)
    export.text_frame.margin_top = Inches(0.07)
    export.text_frame.margin_bottom = Inches(0.07)
    set_text(
        export,
        ["EXPORT CSV / XLSX", "Підбір топових тем без дублювання й канібалізації"],
        size=9.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_eight(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["С Т Р А Т Е Г І Я     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["300+ сторінок — це година ручного збору."],
        size=25,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Sitemap — лише початок",
        body="Для стратегії потрібні URL, breadcrumbs і метадані кожної сторінки.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Вручну — близько години",
        body="Зібрати й структурувати сотні сторінок перед аналізом.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Дорогий контекст для ШІ",
        body="Сирі дані займають багато токенів і ускладнюють надійну перевірку.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.88, 2.34, 0.20, ["S D S   ·   S I T E M A P"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.75, 2.27, 2.40, 0.88, ["369"], size=62, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.34, 2.34, 0.25, ["СТОРІНОК"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        6.82,
        4.22,
        2.26,
        0.40,
        ["Їх треба врахувати у стратегії"],
        size=9.5,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        line_spacing=1.04,
    )
    add_footer(slide, len(prs.slides))


def add_slide_nine(prs: Presentation, page_inventory_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     S I T E   A R C H I T E C T U R E"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Sitemap → готова структура сайту."],
        size=26,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(page_inventory_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["В   О Д Н І Й   Т А Б Л И Ц І"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="URL і breadcrumbs")
    add_benefit(slide, y=2.68, text="Title, description і keywords")
    add_benefit(slide, y=3.42, text="Пошук та Export to Excel")

    impact = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    impact.fill.solid()
    impact.fill.fore_color.rgb = ORANGE
    impact.line.color.rgb = ORANGE
    impact.adjustments[0] = 0.10
    impact.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    impact.text_frame.margin_left = Inches(0.10)
    impact.text_frame.margin_right = Inches(0.10)
    impact.text_frame.margin_top = Inches(0.07)
    impact.text_frame.margin_bottom = Inches(0.07)
    set_text(
        impact,
        ["−1 ГОДИНА", "на підготовку стратегії"],
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_ten(prs: Presentation, seo_qa_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     S E O   Q A"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["SEO QA показує, що оновлювати першими."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(seo_qa_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["З Н А Х О Д И Т Ь"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Дублікати метаданих")
    add_benefit(slide, y=2.68, text="Порожні або слабкі поля")
    add_benefit(slide, y=3.42, text="Calibrated / Strict перевірка")

    priority = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    priority.fill.solid()
    priority.fill.fore_color.rgb = ORANGE
    priority.line.color.rgb = ORANGE
    priority.adjustments[0] = 0.10
    priority.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    priority.text_frame.margin_left = Inches(0.10)
    priority.text_frame.margin_right = Inches(0.10)
    priority.text_frame.margin_top = Inches(0.07)
    priority.text_frame.margin_bottom = Inches(0.07)
    set_text(
        priority,
        ["ПРІОРИТЕТИ", "Сторінки, що точно потребують оновлення"],
        size=9.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


# Strategy Package Check (task-90/91): три слайди після SEO QA — проблема,
# порівняльна матриця Strategy ↔ Framework ↔ Site, SEO QA стратегії/фреймів.
# Функції нижче названі за змістом, а не за номером: після вставки нумерація
# наступних слайдів зсунулась, а футер рахує номер сам (`len(prs.slides)`).


def add_slide_strategy_package_problem(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["S T R A T E G Y   P A C K A G E   C H E C K     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Звірка стратегії з сайтом — години вручну."],
        size=24,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Три джерела, які мають збігатися",
        body="Стратегія, фреймворк кожної сторінки та живий сайт живуть окремо й розходяться непомітно.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Кожну сторінку — вручну",
        body="Відкрити стратегію, кожен фреймворк, кожну сторінку сайту; звірити поле за полем; скласти звіт.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Години, а на великому сайті — день",
        body="Що більше сторінок, то довше і то більше шансів пропустити розбіжність.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.88, 2.34, 0.20, ["S A P H N Y   ·   Q 4   2 0 2 6"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.75, 2.27, 2.40, 0.88, ["15"], size=62, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.34, 2.34, 0.25, ["ФРЕЙМВОРКІВ"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        6.82,
        4.18,
        2.26,
        0.48,
        ["× 6 полів × 3 джерела", "= 270 ручних звірок"],
        size=9.5,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_strategy_package_comparison(prs: Presentation, comparison_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     S T R A T E G Y   P A C K A G E   C H E C K"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Стратегія ↔ фреймворк ↔ сайт — за 30 секунд."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(comparison_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["П О К А З У Є"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Стратегія, фрейм і сайт поруч")
    add_benefit(slide, y=2.68, text="✓ / ✗ по кожному полю")
    add_benefit(slide, y=3.42, text="Фільтр «Mismatches only»")

    export = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    export.fill.solid()
    export.fill.fore_color.rgb = ORANGE
    export.line.color.rgb = ORANGE
    export.adjustments[0] = 0.10
    export.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    export.text_frame.margin_left = Inches(0.10)
    export.text_frame.margin_right = Inches(0.10)
    export.text_frame.margin_top = Inches(0.07)
    export.text_frame.margin_bottom = Inches(0.07)
    set_text(
        export,
        ["EXPORT TO EXCEL", "Звіт одразу в роботу"],
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_strategy_package_qa(prs: Presentation, qa_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["S T R A T E G Y   P A C K A G E   C H E C K     ·     S E O   Q A"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["SEO QA стратегії та фреймів — до публікації."],
        size=24,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(qa_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["З Н А Х О Д И Т Ь"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Дублікати й довжина meta")
    add_benefit(slide, y=2.68, text="Drift фрейму від живої сторінки")
    add_benefit(slide, y=3.42, text="Cross-checks між файлами")

    findings = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    findings.fill.solid()
    findings.fill.fore_color.rgb = ORANGE
    findings.line.color.rgb = ORANGE
    findings.adjustments[0] = 0.10
    findings.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    findings.text_frame.margin_left = Inches(0.10)
    findings.text_frame.margin_right = Inches(0.10)
    findings.text_frame.margin_top = Inches(0.07)
    findings.text_frame.margin_bottom = Inches(0.07)
    set_text(
        findings,
        ["104 ERRORS · 60 WARNINGS", "Один запуск · 15 фреймворків"],
        size=9.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_eleven(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Т Е Х Н І Ч Н И Й   S E O     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Банальні помилки повторюються на сотнях сторінок."],
        size=25,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Їх легко пропустити",
        body="Alt, redirects, Open Graph, title і description — прості, але масові проблеми.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Одна причина → сотні findings",
        body="Шаблонна помилка автоматично повторюється на багатьох URL.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Стан сайту постійно змінюється",
        body="Потрібно бачити проблеми, які існують прямо зараз, а не в старому звіті.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.88, 2.34, 0.20, ["H W E L L   ·   C R A W L"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.75, 2.27, 2.40, 0.88, ["244"], size=62, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.34, 2.34, 0.25, ["СТОРІНКИ"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        6.82,
        4.18,
        2.26,
        0.48,
        ["259 великих зображень", "244 missing alt"],
        size=9.5,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_twelve(prs: Presentation, site_crawl_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     S I T E   C R A W L   I S S U E S"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Що зламано прямо зараз — за один crawl."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(site_crawl_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["Д А Є   В І Д П О В І Д Ь"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Severity і масштаб проблеми")
    add_benefit(slide, y=2.68, text="Готова інструкція Fix")
    add_benefit(slide, y=3.42, text="Re-check після виправлення")

    action = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    action.fill.solid()
    action.fill.fore_color.rgb = ORANGE
    action.line.color.rgb = ORANGE
    action.adjustments[0] = 0.10
    action.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    action.text_frame.margin_left = Inches(0.10)
    action.text_frame.margin_right = Inches(0.10)
    action.text_frame.margin_top = Inches(0.07)
    action.text_frame.margin_bottom = Inches(0.07)
    set_text(
        action,
        ["RUN AUDIT · EXPORT XLSX", "Виміряно crawler’ом — без здогадок ШІ"],
        size=9.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_thirteen(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["M O N I T O R     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["14 сайтів і 2000+ сторінок не перевірити вручну."],
        size=25,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Зміни відбуваються щодня",
        body="Різні люди публікують і редагують сторінки на всіх клієнтських сайтах.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Критичну помилку легко не помітити",
        body="Noindex, broken URL у sitemap або redirect на живій сторінці.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Один фахівець не встигає",
        body="Постійно контролювати весь портфель фізично неможливо.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.85, 2.34, 0.20, ["Щ О Д Е Н Н І   З М І Н И"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.75, 2.20, 2.40, 0.68, ["14"], size=48, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 2.88, 2.34, 0.20, ["САЙТІВ"], size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.15), Inches(3.26), Inches(1.60), Inches(0.015))
    divider.fill.solid()
    divider.fill.fore_color.rgb = WHITE
    divider.line.fill.background()
    add_textbox(slide, 6.70, 3.49, 2.50, 0.58, ["2000+"], size=37, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 4.13, 2.34, 0.22, ["СТОРІНОК"], size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.82, 4.55, 2.26, 0.20, ["ручний контроль неможливий"], size=8.5, color=WHITE, align=PP_ALIGN.CENTER)
    add_footer(slide, len(prs.slides))


def add_process_card(
    slide,
    *,
    x: float,
    label: str,
    value_lines: list[str],
    body: str,
    highlight: bool = False,
) -> None:
    background = ORANGE if highlight else WHITE
    foreground = WHITE if highlight else BLACK
    secondary = WHITE if highlight else GRAY
    label_color = WHITE if highlight else ORANGE
    border_color = ORANGE if highlight else LIGHT_GRAY
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(1.70),
        Inches(2.00),
        Inches(2.15),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = background
    card.line.color.rgb = border_color
    card.line.width = Pt(1.1)
    card.adjustments[0] = 0.10
    add_textbox(slide, x + 0.15, 1.96, 1.70, 0.18, [label], size=8, color=label_color, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        x + 0.12,
        2.35,
        1.76,
        0.65,
        value_lines,
        size=21,
        color=foreground,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=0.94,
    )
    add_textbox(slide, x + 0.18, 3.27, 1.64, 0.30, [body], size=8.5, color=secondary, align=PP_ALIGN.CENTER, line_spacing=1.02)


def add_slide_fourteen(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     2 4 / 7   M O N I T O R I N G"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Топові сторінки перевіряються кожні 2 години."],
        size=25,
        color=BLACK,
        bold=True,
    )

    add_process_card(slide, x=0.55, label="З М І Н А   Н А   С А Й Т І", value_lines=["NOINDEX", "3XX · BROKEN"], body="Помилка з’явилась")
    add_textbox(slide, 2.58, 2.55, 0.24, 0.30, ["→"], size=19, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_process_card(slide, x=2.85, label="L I G H T   C H E C K", value_lines=["2", "ГОДИНИ"], body="Автоматичний обхід", highlight=True)
    add_textbox(slide, 4.88, 2.55, 0.24, 0.30, ["→"], size=19, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_process_card(slide, x=5.15, label="К Р И Т И Ч Н О", value_lines=["ALERT"], body="Monitor повідомляє")
    add_textbox(slide, 7.18, 2.55, 0.24, 0.30, ["→"], size=19, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_process_card(slide, x=7.45, label="К О М А Н Д А", value_lines=["FIX"], body="Вчасне виправлення")

    for x, width, label in (
        (0.72, 1.90, "TOP PAGES"),
        (2.92, 1.90, "ROBOTS.TXT"),
        (5.12, 1.90, "SITEMAP"),
        (7.32, 1.90, "REDIRECTS"),
    ):
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.20), Inches(width), Inches(0.46))
        chip.fill.solid()
        chip.fill.fore_color.rgb = WHITE
        chip.line.color.rgb = LIGHT_GRAY
        chip.adjustments[0] = 0.15
        chip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        chip.text_frame.margin_left = 0
        chip.text_frame.margin_right = 0
        chip.text_frame.margin_top = 0
        chip.text_frame.margin_bottom = 0
        set_text(chip, [label], size=9, color=BLACK, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, len(prs.slides))


def add_slide_fifteen(prs: Presentation, monitor_overview_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(slide, 0.68, 0.36, 8.65, 0.18, ["S I T E   M O N I T O R     ·     C U R R E N T   S T A T E"], size=10, color=ORANGE, bold=True)
    add_textbox(slide, 0.68, 0.78, 8.65, 0.58, ["Robots, sitemap, redirects і pages — в одному місці."], size=24, color=BLACK, bold=True)

    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.60), Inches(1.50), Inches(6.25), Inches(3.60))
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(str(monitor_overview_image), Inches(0.62), Inches(1.52), width=Inches(6.21), height=Inches(3.56))

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["М О Н І Т О Р И Т Ь"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Robots.txt і доступ Googlebot")
    add_benefit(slide, y=2.68, text="Sitemap: OK / redirect / broken")
    add_benefit(slide, y=3.42, text="Status, indexability і canonical")

    action = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(4.15), Inches(2.30), Inches(0.93))
    action.fill.solid()
    action.fill.fore_color.rgb = ORANGE
    action.line.color.rgb = ORANGE
    action.adjustments[0] = 0.10
    action.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    action.text_frame.margin_left = Inches(0.10)
    action.text_frame.margin_right = Inches(0.10)
    action.text_frame.margin_top = Inches(0.07)
    action.text_frame.margin_bottom = Inches(0.07)
    set_text(action, ["CHECK NOW · HISTORY", "Поточний стан і ручний запуск перевірки"], size=9.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.08)
    add_footer(slide, len(prs.slides))


def add_slide_sixteen(prs: Presentation, monitor_history_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(slide, 0.68, 0.36, 8.65, 0.18, ["M O N I T O R     ·     І С Т О Р І Я   З М І Н"], size=10, color=ORANGE, bold=True)
    add_textbox(slide, 0.68, 0.78, 8.65, 0.58, ["Знаємо дату й приблизний час кожної зміни."], size=25, color=BLACK, bold=True)
    add_textbox(slide, 0.68, 1.43, 3.10, 0.18, ["Щ О   З Б Е Р І Г А Є Т Ь С Я"], size=8.5, color=ORANGE, bold=True)

    add_overview_item(slide, y=1.82, number="1", title="Точна дата фіксації", body="Коли Monitor записав нову версію")
    add_overview_item(slide, y=2.68, number="2", title="Часове вікно появи", body="Між двома послідовними перевірками")
    add_overview_item(slide, y=3.54, number="3", title="Копії before / after", body="Обидві версії можна відкрити й порівняти")

    history = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.68), Inches(4.42), Inches(3.10), Inches(0.62))
    history.fill.solid()
    history.fill.fore_color.rgb = ORANGE
    history.line.color.rgb = ORANGE
    history.adjustments[0] = 0.12
    history.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    history.text_frame.margin_left = Inches(0.12)
    history.text_frame.margin_right = Inches(0.12)
    history.text_frame.margin_top = 0
    history.text_frame.margin_bottom = 0
    set_text(history, ["МАЙЖЕ ВСІ ПУБЛІЧНІ ЗМІНИ САЙТУ"], size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.07), Inches(1.48), Inches(5.30), Inches(3.79))
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(str(monitor_history_image), Inches(4.09), Inches(1.50), width=Inches(5.26), height=Inches(3.75))
    add_footer(slide, len(prs.slides))


def add_slide_seventeen(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["C L I E N T   T R E N D S     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Щотижневий контроль трафіку — ручний квест."],
        size=25,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Код для входу в GSC",
        body="Отримати код, підтвердити доступ — і лише тоді відкрити дані.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Один акаунт на всю компанію",
        body="Усі клієнти в одному GSC-акаунті: якщо він зайнятий, треба чекати доступу.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Періоди й цифри — вручну",
        body="Змінювати діапазони та самостійно зіставляти показники.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.86, 2.34, 0.20, ["К О Ж Н О Г О   Т И Ж Н Я"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 2.25, 2.34, 0.30, ["КОД"], size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 2.60, 2.34, 0.23, ["↓"], size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 2.92, 2.34, 0.30, ["GSC"], size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.27, 2.34, 0.23, ["↓"], size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.57, 2.34, 0.34, ["КЛІЄНТ × 14"], size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.15), Inches(4.12), Inches(1.60), Inches(0.015))
    divider.fill.solid()
    divider.fill.fore_color.rgb = WHITE
    divider.line.fill.background()
    add_textbox(slide, 6.80, 4.42, 2.30, 0.22, ["І ЗНОВУ ІНШИЙ ПЕРІОД"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, len(prs.slides))


def add_slide_eighteen(prs: Presentation, trends_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     C L I E N T   T R E N D S"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Усі клієнти й періоди — на одному графіку."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(trends_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["У С Е   Н А   О Д Н О М У   Е К Р А Н І"], size=8.2, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Усі клієнти поруч")
    add_benefit(slide, y=2.68, text="Day / Week / Month")
    add_benefit(slide, y=3.42, text="Періоди від 7 до 365 днів")

    action = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    action.fill.solid()
    action.fill.fore_color.rgb = ORANGE
    action.line.color.rgb = ORANGE
    action.adjustments[0] = 0.10
    action.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    action.text_frame.margin_left = Inches(0.10)
    action.text_frame.margin_right = Inches(0.10)
    action.text_frame.margin_top = Inches(0.07)
    action.text_frame.margin_bottom = Inches(0.07)
    set_text(
        action,
        ["COMPARE · INDEXED · ALL", "Без очікування доступу до GSC"],
        size=9.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_metric_source_card(
    slide,
    *,
    x: float,
    eyebrow: str,
    source: str,
    metrics: list[str],
) -> None:
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(1.68),
        Inches(2.48),
        Inches(2.27),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LIGHT_GRAY
    card.line.width = Pt(1.1)
    card.adjustments[0] = 0.10
    add_textbox(slide, x + 0.20, 1.96, 2.08, 0.18, [eyebrow], size=8, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.20, 2.28, 2.08, 0.48, [source], size=27, color=BLACK, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.24, 3.10, 2.00, 0.53, metrics, size=9, color=GRAY, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.12)


def add_slide_nineteen(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["C L I E N T   T R E N D S     ·     G S C   +   G A 4"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["GSC + GA4 — одна точка контролю."],
        size=26,
        color=BLACK,
        bold=True,
    )

    add_metric_source_card(
        slide,
        x=0.60,
        eyebrow="SEARCH CONSOLE",
        source="GSC",
        metrics=["CLICKS", "IMPRESSIONS"],
    )
    add_textbox(slide, 3.13, 2.57, 0.34, 0.32, ["→"], size=22, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

    center = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(3.55),
        Inches(1.55),
        Inches(2.90),
        Inches(2.53),
    )
    center.fill.solid()
    center.fill.fore_color.rgb = ORANGE
    center.line.color.rgb = ORANGE
    center.adjustments[0] = 0.10
    add_textbox(slide, 3.80, 1.87, 2.40, 0.18, ["C L I E N T   T R E N D S"], size=8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 3.80, 2.27, 2.40, 0.36, ["1 ПЕРІОД"], size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 3.80, 2.77, 2.40, 0.36, ["УСІ КЛІЄНТИ"], size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 3.80, 3.27, 2.40, 0.36, ["1 ГРАФІК"], size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(slide, 6.53, 2.57, 0.34, 0.32, ["←"], size=22, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_metric_source_card(
        slide,
        x=6.92,
        eyebrow="GOOGLE ANALYTICS",
        source="GA4",
        metrics=["SESSIONS · USERS", "TOTAL · ORGANIC · DIRECT CONVERSIONS"],
    )

    statement = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.60),
        Inches(4.35),
        Inches(8.80),
        Inches(0.70),
    )
    statement.fill.solid()
    statement.fill.fore_color.rgb = WHITE
    statement.line.color.rgb = LIGHT_GRAY
    statement.line.width = Pt(1)
    statement.adjustments[0] = 0.10
    statement.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    statement.text_frame.margin_left = Inches(0.20)
    statement.text_frame.margin_right = Inches(0.20)
    statement.text_frame.margin_top = 0
    statement.text_frame.margin_bottom = 0
    set_text(
        statement,
        ["Одразу видно, хто росте, а хто потребує уваги цього тижня."],
        size=13,
        color=BLACK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, len(prs.slides))


def add_slide_twenty(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["W E B   A R C H I V E     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Дані старого сайту були — візуальної копії не було."],
        size=24,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Структура була — вигляду не було",
        body="URL і sitemap збереглися, але старе меню візуально вже не відновити.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="Контент був неповний",
        body="Метадані не показували, що саме бачив користувач на сторінці.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Публічний архів не допоміг",
        body="Версія за грудень 2025 була застарілою й відтворювалася з помилками.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.80, 1.86, 2.30, 0.18, ["П І С Л Я   М І Г Р А Ц І Ї"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.80, 2.25, 2.30, 0.18, ["ЗБЕРЕГЛОСЯ"], size=8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.76, 2.57, 2.38, 0.38, ["URL · META", "STRUCTURE"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.02)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.15), Inches(3.17), Inches(1.60), Inches(0.015))
    divider.fill.solid()
    divider.fill.fore_color.rgb = WHITE
    divider.line.fill.background()
    add_textbox(slide, 6.80, 3.48, 2.30, 0.18, ["НЕМОЖЛИВО ПОБАЧИТИ"], size=8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.76, 3.80, 2.38, 0.38, ["MENU · CONTENT", "VISUAL"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.02)
    add_textbox(slide, 6.80, 4.51, 2.30, 0.20, ["DEC 2025 · BROKEN"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, len(prs.slides))


def add_slide_twenty_one(prs: Presentation, archive_overview_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     W E B   A R C H I V E"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Власний архів зберігає весь портфель."],
        size=26,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(archive_overview_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["З Б Е Р І Г А Є"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="7 850 URL у контролі")
    add_benefit(slide, y=2.68, text="6 926 версій сторінок")
    add_benefit(slide, y=3.42, text="25 822 унікальні файли")

    schedule = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    schedule.fill.solid()
    schedule.fill.fore_color.rgb = ORANGE
    schedule.line.color.rgb = ORANGE
    schedule.adjustments[0] = 0.10
    schedule.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    schedule.text_frame.margin_left = Inches(0.10)
    schedule.text_frame.margin_right = Inches(0.10)
    schedule.text_frame.margin_top = Inches(0.07)
    schedule.text_frame.margin_bottom = Inches(0.07)
    set_text(
        schedule,
        ["13:30 · 19:30", "Два автоматичні обходи щодня"],
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_twenty_two(prs: Presentation, archive_view_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["W E B   A R C H I V E     ·     R E P L A Y"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Архівну версію можна переглядати як живий сайт."],
        size=24,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(archive_view_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["М О Ж Н А   П Е Р Е В І Р И Т И"], size=8.2, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Меню, контент і дизайн")
    add_benefit(slide, y=2.68, text="Meta й View source")
    add_benefit(slide, y=3.42, text="Previous / Next / Compare")

    manual = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    manual.fill.solid()
    manual.fill.fore_color.rgb = ORANGE
    manual.line.color.rgb = ORANGE
    manual.adjustments[0] = 0.10
    manual.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    manual.text_frame.margin_left = Inches(0.10)
    manual.text_frame.margin_right = Inches(0.10)
    manual.text_frame.margin_top = Inches(0.07)
    manual.text_frame.margin_bottom = Inches(0.07)
    set_text(
        manual,
        ["ARCHIVE NOW", "Будь-який URL — вручну"],
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


# URL Inspector (task-61/88): фінальна трійка — проблема, аналітика + стан у
# Google, технічний стан сторінки (аудит, скріншоти, архів, зміни).


def add_slide_url_inspector_problem(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["U R L   I N S P E C T O R     ·     П Р О Б Л Е М А"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.60,
        ["Усе про одну сторінку — у шести різних місцях."],
        size=24,
        color=BLACK,
        bold=True,
    )

    add_blog_problem(
        slide,
        y=1.67,
        number="1",
        title="Шість вкладок на одне питання",
        body="GSC, GA4, аудит, скріншоти, архів, монітор — окремі інструменти й окремі акаунти.",
    )
    add_blog_problem(
        slide,
        y=2.72,
        number="2",
        title="«Що з цією сторінкою?» — 20 хвилин",
        body="Зібрати цифри, перевірити індексацію, глянути meta — щоразу з нуля.",
    )
    add_blog_problem(
        slide,
        y=3.77,
        number="3",
        title="Історії немає",
        body="Що змінилось на сторінці і коли — ніхто не памʼятає й ніде не записано.",
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.55),
        Inches(1.55),
        Inches(2.80),
        Inches(3.42),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = ORANGE
    card.line.color.rgb = ORANGE
    card.adjustments[0] = 0.10
    add_textbox(slide, 6.78, 1.88, 2.34, 0.20, ["О Д Н А   С Т О Р І Н К А"], size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.75, 2.27, 2.40, 0.88, ["6"], size=62, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.78, 3.34, 2.34, 0.25, ["ДЖЕРЕЛ ДАНИХ"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        6.82,
        4.18,
        2.26,
        0.48,
        ["GSC · GA4 · INDEX · AUDIT", "SCREENS · ARCHIVE"],
        size=9.5,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_url_inspector_analytics(prs: Presentation, analytics_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["Р І Ш Е Н Н Я     ·     U R L   I N S P E C T O R"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Вставив URL — отримав усе про сторінку."],
        size=25,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(analytics_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["А Н А Л І Т И К А"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="GSC + GA4: кліки, сесії, конверсії")
    add_benefit(slide, y=2.68, text="Запити й ключі сторінки")
    add_benefit(slide, y=3.42, text="Стан у Google: індекс, canonical")

    action = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    action.fill.solid()
    action.fill.fore_color.rgb = ORANGE
    action.line.color.rgb = ORANGE
    action.adjustments[0] = 0.10
    action.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    action.text_frame.margin_left = Inches(0.10)
    action.text_frame.margin_right = Inches(0.10)
    action.text_frame.margin_top = Inches(0.07)
    action.text_frame.margin_bottom = Inches(0.07)
    set_text(
        action,
        ["PASTE A PAGE URL", "Без входу в GSC чи GA4"],
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def add_slide_url_inspector_page_state(prs: Presentation, page_state_image: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_textbox(
        slide,
        0.68,
        0.36,
        8.65,
        0.18,
        ["U R L   I N S P E C T O R     ·     С Т А Н   С Т О Р І Н К И"],
        size=10,
        color=ORANGE,
        bold=True,
    )
    add_textbox(
        slide,
        0.68,
        0.78,
        8.65,
        0.58,
        ["Аудит, скріншоти, архів і зміни — на одному екрані."],
        size=24,
        color=BLACK,
        bold=True,
    )

    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(1.50),
        Inches(6.25),
        Inches(3.60),
    )
    border.fill.solid()
    border.fill.fore_color.rgb = BLACK
    border.line.color.rgb = BLACK
    slide.shapes.add_picture(
        str(page_state_image),
        Inches(0.62),
        Inches(1.52),
        width=Inches(6.21),
        height=Inches(3.56),
    )

    add_textbox(slide, 7.12, 1.54, 2.20, 0.18, ["Щ Е   Н А   С Т О Р І Н Ц І"], size=8.5, color=ORANGE, bold=True)
    add_benefit(slide, y=1.95, text="Технічний аудит: meta, H1, помилки")
    add_benefit(slide, y=2.68, text="Desktop · Tablet · Mobile")
    add_benefit(slide, y=3.42, text="Архівні версії й журнал змін")

    action = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.05),
        Inches(4.15),
        Inches(2.30),
        Inches(0.93),
    )
    action.fill.solid()
    action.fill.fore_color.rgb = ORANGE
    action.line.color.rgb = ORANGE
    action.adjustments[0] = 0.10
    action.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    action.text_frame.margin_left = Inches(0.10)
    action.text_frame.margin_right = Inches(0.10)
    action.text_frame.margin_top = Inches(0.07)
    action.text_frame.margin_bottom = Inches(0.07)
    set_text(
        action,
        ["LIVE AUDIT · RECHECK", "Жива сторінка — за секунди"],
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        line_spacing=1.08,
    )
    add_footer(slide, len(prs.slides))


def build(
    reference: Path,
    output: Path,
    dashboard_image: Path,
    client_overview_image: Path,
    monthly_table_image: Path,
    blogs_table_image: Path,
    page_inventory_image: Path,
    seo_qa_image: Path,
    strategy_package_comparison_image: Path,
    strategy_package_qa_image: Path,
    site_crawl_image: Path,
    monitor_overview_image: Path,
    monitor_history_image: Path,
    trends_image: Path,
    archive_overview_image: Path,
    archive_view_image: Path,
    url_inspector_analytics_image: Path,
    url_inspector_page_state_image: Path,
) -> None:
    prs = Presentation(reference)

    # The opening slide in the reference already contains the design grammar
    # we need. Keep it and remove the remaining reference content.
    for index in range(len(prs.slides) - 1, 0, -1):
        remove_slide(prs, index)

    slide = prs.slides[0]
    shapes = slide.shapes

    # Top-right wordmark.
    set_text(shapes[0], ["THE YOVTH"], size=13, color=BLACK, bold=False, align=PP_ALIGN.RIGHT)

    # Orange identity card.
    set_text(shapes[2], ["THE YOVTH"], size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    set_text(shapes[3], ["SEO"], size=82, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    set_text(shapes[4], ["TOOLS"], size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    # Eyebrow and title.
    set_text(
        shapes[5],
        ["T H E   Y O V T H     ·     S E O   T O O L S"],
        size=10,
        color=ORANGE,
        bold=True,
    )

    title = shapes[6]
    title.left = Inches(0.60)
    title.top = Inches(1.43)
    title.width = Inches(5.82)
    title.height = Inches(2.55)
    title.text_frame.margin_left = Inches(0.10)
    title.text_frame.margin_right = Inches(0.05)
    title.text_frame.margin_top = Inches(0.02)
    title.text_frame.margin_bottom = Inches(0.02)
    title.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    set_text(
        title,
        ["SEO TOOLS для", "контролю якості та", "продуктивності клієнтів"],
        size=37,
        color=BLACK,
        bold=True,
        line_spacing=0.98,
    )

    # This is intentionally a minimal opening slide.
    set_text(shapes[7], [""], size=15, color=BLACK)
    set_text(
        shapes[8],
        ["Я К І С Т Ь     ·     П Р О З О Р І С Т Ь     ·     Р Е З У Л Ь Т А Т"],
        size=9,
        color=ORANGE,
        bold=True,
    )

    add_slide_two(prs, dashboard_image, output.parent / "assets")
    add_slide_three(prs, client_overview_image)
    add_slide_four(prs)
    add_slide_five(prs, monthly_table_image)
    add_slide_six(prs)
    add_slide_seven(prs, blogs_table_image)
    add_slide_eight(prs)
    add_slide_nine(prs, page_inventory_image)
    add_slide_ten(prs, seo_qa_image)
    add_slide_strategy_package_problem(prs)
    add_slide_strategy_package_comparison(prs, strategy_package_comparison_image)
    add_slide_strategy_package_qa(prs, strategy_package_qa_image)
    add_slide_eleven(prs)
    add_slide_twelve(prs, site_crawl_image)
    add_slide_thirteen(prs)
    add_slide_fourteen(prs)
    add_slide_fifteen(prs, monitor_overview_image)
    add_slide_sixteen(prs, monitor_history_image)
    add_slide_seventeen(prs)
    add_slide_eighteen(prs, trends_image)
    add_slide_nineteen(prs)
    add_slide_twenty(prs)
    add_slide_twenty_one(prs, archive_overview_image)
    add_slide_twenty_two(prs, archive_view_image)
    add_slide_url_inspector_problem(prs)
    add_slide_url_inspector_analytics(prs, url_inspector_analytics_image)
    add_slide_url_inspector_page_state(prs, url_inspector_page_state_image)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--dashboard-image",
        type=Path,
        default=Path("tests/e2e/frontend/screenshots/auth-viewer-dashboard-1440.png"),
    )
    parser.add_argument(
        "--client-overview-image",
        type=Path,
        default=Path("presentation/assets/slide-03-client-overview.png"),
    )
    parser.add_argument(
        "--monthly-table-image",
        type=Path,
        default=Path("presentation/assets/slide-05-monthly-table.png"),
    )
    parser.add_argument(
        "--blogs-table-image",
        type=Path,
        default=Path("presentation/assets/slide-07-blogs-table.png"),
    )
    parser.add_argument(
        "--page-inventory-image",
        type=Path,
        default=Path("presentation/assets/slide-09-page-inventory.png"),
    )
    parser.add_argument(
        "--seo-qa-image",
        type=Path,
        default=Path("presentation/assets/slide-10-seo-qa.png"),
    )
    parser.add_argument(
        "--strategy-package-comparison-image",
        type=Path,
        default=Path("presentation/assets/strategy-package-comparison.png"),
    )
    parser.add_argument(
        "--strategy-package-qa-image",
        type=Path,
        default=Path("presentation/assets/strategy-package-qa.png"),
    )
    parser.add_argument(
        "--site-crawl-image",
        type=Path,
        default=Path("presentation/assets/slide-12-site-crawl-issues.png"),
    )
    parser.add_argument(
        "--monitor-overview-image",
        type=Path,
        default=Path("presentation/assets/slide-15-monitor-overview.png"),
    )
    parser.add_argument(
        "--monitor-history-image",
        type=Path,
        default=Path("presentation/assets/slide-16-monitor-history.png"),
    )
    parser.add_argument(
        "--trends-image",
        type=Path,
        default=Path("presentation/assets/slide-18-client-trends.png"),
    )
    parser.add_argument(
        "--archive-overview-image",
        type=Path,
        default=Path("presentation/assets/slide-21-web-archive.png"),
    )
    parser.add_argument(
        "--archive-view-image",
        type=Path,
        default=Path("presentation/assets/slide-22-archive-view.png"),
    )
    parser.add_argument(
        "--url-inspector-analytics-image",
        type=Path,
        default=Path("presentation/assets/url-inspector-analytics.png"),
    )
    parser.add_argument(
        "--url-inspector-page-state-image",
        type=Path,
        default=Path("presentation/assets/url-inspector-page-state.png"),
    )
    args = parser.parse_args()
    build(
        args.reference,
        args.output,
        args.dashboard_image,
        args.client_overview_image,
        args.monthly_table_image,
        args.blogs_table_image,
        args.page_inventory_image,
        args.seo_qa_image,
        args.strategy_package_comparison_image,
        args.strategy_package_qa_image,
        args.site_crawl_image,
        args.monitor_overview_image,
        args.monitor_history_image,
        args.trends_image,
        args.archive_overview_image,
        args.archive_view_image,
        args.url_inspector_analytics_image,
        args.url_inspector_page_state_image,
    )


if __name__ == "__main__":
    main()
