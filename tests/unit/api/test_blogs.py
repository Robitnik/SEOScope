"""Client Blogs endpoints (task-26) — read side, run side, export.

Everything here is machine data: the tests seed `blog_summary` / `blog_posts`
dataset rows exactly as `BlogsAdapter` writes them and assert the API reads
them the way the UI needs.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime

from openpyxl import load_workbook

from app.db.repositories.analytics_datasets import (
    AnalyticsDatasetRepository,
    AnalyticsDatasetRowInput,
)
from app.db.repositories.clients import ClientRepository
from app.db.repositories.jobs import JobRepository
from app.db.repositories.snapshots import SnapshotRepository
from app.integrations.blogs import post_key
from app.schemas.enums import ComparisonMode, PeriodKind, Provider, Quality, SourceMode

URL_1 = "https://blog.example.com/blog/post-1"
URL_2 = "https://blog.example.com/blog/post-2"


def _client(db, *, slug: str = "acme-blogs"):
    return ClientRepository(db).create(
        name="Acme Blogs",
        slug=slug,
        domain="blog.example.com",
        timezone="America/New_York",
    )


def _snapshot(db, created, *, observed: datetime, warnings=None):
    job = JobRepository(db).create_run(
        client_ids=[created.id],
        providers=[Provider.blogs],
        period_kind=PeriodKind.month,
        comparison_mode=ComparisonMode.closed,
        anchor_date=date(2026, 7, 23),
        mode="blogs",
    )[0]
    return SnapshotRepository(db).create_immutable(
        job_id=job.id,
        client_id=created.id,
        provider=Provider.blogs,
        source_mode=SourceMode.auto,
        observed_at=observed,
        quality=Quality.complete,
        warnings=warnings or [],
        provenance={"analytics_datasets": {"attempted": True}},
    )


def _insert(db, created, snapshot, dataset_key: str, rows: list[AnalyticsDatasetRowInput]):
    AnalyticsDatasetRepository(db).bulk_insert(
        snapshot_id=snapshot.id,
        client_id=created.id,
        provider=Provider.blogs,
        period_start=date(2026, 7, 23),
        period_end=date(2026, 7, 23),
        rows=rows,
    )


def _seed_run(db, *, slug: str = "acme-blogs"):
    created = _client(db, slug=slug)
    snapshot = _snapshot(
        db, created, observed=datetime(2026, 7, 23, 12, tzinfo=UTC),
        warnings=["crawl incomplete: 1 URL(s) failed to fetch"],
    )
    _insert(
        db, created, snapshot, "blog_summary",
        [
            AnalyticsDatasetRowInput(
                dataset_key="blog_summary", dimension_key="summary", dimension_value="blog",
                metrics={
                    "post_count": 2,
                    "blog_url": "https://blog.example.com/blog/",
                    "crawl_completeness": 0.9,
                    "warnings": ["crawl incomplete: 1 URL(s) failed to fetch"],
                },
                quality=Quality.complete, rank=0,
            ),
        ],
    )
    _insert(
        db, created, snapshot, "blog_posts",
        [
            AnalyticsDatasetRowInput(
                dataset_key="blog_posts", dimension_key="post", dimension_value=post_key(URL_1),
                metrics={
                    "url": URL_1, "title": "Post 1", "description": "Desc 1",
                    "word_count": 500, "published_date": "2026-07-01", "modified_date": None,
                },
                quality=Quality.complete, rank=1,
            ),
            AnalyticsDatasetRowInput(
                dataset_key="blog_posts", dimension_key="post", dimension_value=post_key(URL_2),
                metrics={
                    "url": URL_2, "title": "Another Post", "description": "Desc 2",
                    "word_count": 900, "published_date": "2026-07-10", "modified_date": None,
                },
                quality=Quality.complete, rank=2,
            ),
        ],
    )
    db.commit()
    return created


def test_latest_reports_no_run_for_a_fresh_client(client, db):
    created = _client(db, slug="fresh-blogs")
    db.commit()

    body = client.get(f"/api/v1/clients/{created.id}/blogs/latest").json()

    assert body["has_run"] is False
    assert body["post_count"] is None  # never a fabricated 0


def test_latest_returns_post_count_and_blog_url(client, db):
    created = _seed_run(db)

    body = client.get(f"/api/v1/clients/{created.id}/blogs/latest").json()

    assert body["has_run"] is True
    assert body["post_count"] == 2
    assert body["blog_url"] == "https://blog.example.com/blog/"
    assert body["warnings"] == ["crawl incomplete: 1 URL(s) failed to fetch"]


def test_latest_404_for_unknown_client(client):
    response = client.get("/api/v1/clients/does-not-exist/blogs/latest")
    assert response.status_code == 404


def test_posts_returns_rows(client, db):
    created = _seed_run(db)

    body = client.get(f"/api/v1/clients/{created.id}/blogs/posts").json()

    assert body["total"] == 2
    urls = {row["url"] for row in body["rows"]}
    assert urls == {URL_1, URL_2}


def test_posts_search_matches_title(client, db):
    created = _seed_run(db)

    body = client.get(
        f"/api/v1/clients/{created.id}/blogs/posts", params={"q": "Another"}
    ).json()

    assert body["total"] == 1
    assert body["rows"][0]["url"] == URL_2


def test_posts_sort_by_words_desc(client, db):
    created = _seed_run(db)

    body = client.get(
        f"/api/v1/clients/{created.id}/blogs/posts", params={"sort": "words_desc"}
    ).json()

    assert [row["url"] for row in body["rows"]] == [URL_2, URL_1]


def test_posts_sort_defaults_for_unknown_value(client, db):
    created = _seed_run(db)

    body = client.get(
        f"/api/v1/clients/{created.id}/blogs/posts", params={"sort": "not-a-real-sort"}
    ).json()

    assert body["sort"] == "date_desc"


def test_posts_null_word_count_is_not_zero(client, db):
    created = _client(db, slug="null-words")
    snapshot = _snapshot(db, created, observed=datetime(2026, 7, 23, tzinfo=UTC))
    _insert(
        db, created, snapshot, "blog_posts",
        [
            AnalyticsDatasetRowInput(
                dataset_key="blog_posts", dimension_key="post", dimension_value=post_key(URL_1),
                metrics={"url": URL_1, "title": "No word count", "description": "",
                         "word_count": None, "published_date": None, "modified_date": None},
                quality=Quality.complete, rank=1,
            ),
        ],
    )
    db.commit()

    body = client.get(f"/api/v1/clients/{created.id}/blogs/posts").json()
    assert body["rows"][0]["word_count"] is None
    assert body["rows"][0]["published_date"] is None


def test_start_run_queues_a_blogs_job(client, db):
    created = _client(db, slug="run-blogs")
    db.commit()

    response = client.post(f"/api/v1/clients/{created.id}/blogs/runs")

    assert response.status_code == 202
    body = response.json()
    assert body["events_url"] == f"/api/v1/sync-jobs/{body['job_id']}/events"
    job = JobRepository(db).get(body["job_id"])
    assert job is not None
    assert job.mode == "blogs"
    assert job.requested_providers_json == ["blogs"]


def test_start_run_404_for_unknown_client(client):
    response = client.post("/api/v1/clients/does-not-exist/blogs/runs")
    assert response.status_code == 404


def test_export_csv_without_a_run_is_404(client, db):
    created = _client(db, slug="noexport-blogs-csv")
    db.commit()
    assert client.get(f"/api/v1/clients/{created.id}/blogs/export.csv").status_code == 404


def test_export_xlsx_without_a_run_is_404(client, db):
    created = _client(db, slug="noexport-blogs-xlsx")
    db.commit()
    assert client.get(f"/api/v1/clients/{created.id}/blogs/export.xlsx").status_code == 404


def test_export_csv_has_correct_columns_and_rows(client, db):
    created = _seed_run(db)

    response = client.get(f"/api/v1/clients/{created.id}/blogs/export.csv")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    reader = csv.reader(io.StringIO(response.text))
    rows = list(reader)
    assert rows[0] == ["Title", "Description", "Word count", "URL", "Published date"]
    assert len(rows) == 3  # header + 2 posts
    urls = {row[3] for row in rows[1:]}
    assert urls == {URL_1, URL_2}


def test_export_xlsx_has_correct_columns_and_rows(client, db):
    created = _seed_run(db)

    response = client.get(f"/api/v1/clients/{created.id}/blogs/export.xlsx")

    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb["Blog posts"]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["Title", "Description", "Word count", "URL", "Published date"]
    assert ws.max_row == 3  # header + 2 posts
