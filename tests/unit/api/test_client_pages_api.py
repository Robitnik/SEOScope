"""Page Inventory REST API (task-57): 404s, `monitor_disabled`/`no_pages_yet`
409s, the `?profile=` toggle, and the XLSX export shape."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.db.repositories.clients import ClientRepository
from app.db.repositories.monitor import MonitorRepository


def _client(db, *, slug: str = "acme") -> str:
    client = ClientRepository(db).create(name="Acme", slug=slug, domain=f"{slug}.com")
    db.commit()
    return client.id


def _seed_pages(db, client_id: str) -> None:
    repo = MonitorRepository(db)
    for path, title, desc in (
        ("/", "Acme Home", "Shop everything at Acme, your one-stop shop for stuff."),
        ("/about/", "About Acme", "Learn about Acme, a company that makes things."),
    ):
        page = repo.get_or_create_page(client_id, f"https://acme.com{path}")
        repo.upsert_page_state(
            page, status=200, final_url=f"https://acme.com{path}", hop_count=0,
            indexable=True, meta_robots=None, x_robots_tag=None,
            canonical=f"https://acme.com{path}", title=title, description=desc, h1=None,
            word_count=None, analytics_ids=[], signals_sha256="s", text_sha256="t",
            raw_sha256="r", in_sitemap=True,
        )
        repo.create_page_version(
            page, http_status=200, signals={}, signals_sha256="s", text_sha256="t",
            raw_sha256="r", body_blob_sha256=None, changed_fields=[], prev_version_id=None,
        )
    db.commit()


def test_get_404_for_unknown_client(client):
    assert client.get("/api/v1/clients/nope/pages").status_code == 404


def test_get_before_any_rebuild_is_never_built_not_500(client, db):
    client_id = _client(db)
    _seed_pages(db, client_id)
    response = client.get(f"/api/v1/clients/{client_id}/pages")
    assert response.status_code == 404
    assert response.json()["detail"]["error_code"] == "never_built"


def test_rebuild_404_for_unknown_client(client):
    assert client.post("/api/v1/clients/nope/pages/rebuild").status_code == 404


def test_rebuild_409_when_no_pages_collected_yet(client, db):
    client_id = _client(db)
    response = client.post(f"/api/v1/clients/{client_id}/pages/rebuild")
    assert response.status_code == 409
    assert response.json()["detail"]["error_code"] == "no_pages_yet"


def test_rebuild_409_when_monitor_disabled(client, db):
    client_id = _client(db)
    _seed_pages(db, client_id)
    MonitorRepository(db).update_settings(client_id, enabled=False)
    db.commit()
    response = client.post(f"/api/v1/clients/{client_id}/pages/rebuild")
    assert response.status_code == 409
    assert response.json()["detail"]["error_code"] == "monitor_disabled"


def test_rebuild_then_get_returns_the_built_table(client, db):
    client_id = _client(db)
    _seed_pages(db, client_id)

    rebuild_response = client.post(f"/api/v1/clients/{client_id}/pages/rebuild")
    assert rebuild_response.status_code == 200
    table = rebuild_response.json()["table"]
    assert table["summary"]["pages_unique"] == 2
    assert table["total"] == 2

    get_response = client.get(f"/api/v1/clients/{client_id}/pages")
    assert get_response.status_code == 200
    assert get_response.json()["total"] == 2


def test_profile_toggle_changes_severity_but_not_cluster_count(client, db):
    client_id = _client(db)
    repo = MonitorRepository(db)
    shared_desc = "Buy the best widgets online today, shop now for less."
    # Deliberately DISSIMILAR titles — isolates the cluster count to the
    # single shared description, so this test measures exactly one thing.
    distinct_titles = [
        "Vintage Wooden Rocking Chair for Living Rooms",
        "Handmade Ceramic Coffee Mug Set of Four",
        "Waterproof Hiking Backpack with Rain Cover",
    ]
    for i in range(3):
        page = repo.get_or_create_page(client_id, f"https://acme.com/p{i}/")
        repo.upsert_page_state(
            page, status=200, final_url=f"https://acme.com/p{i}/", hop_count=0,
            indexable=True, meta_robots=None, x_robots_tag=None,
            canonical=f"https://acme.com/p{i}/", title=distinct_titles[i],
            description=shared_desc, h1=None, word_count=None, analytics_ids=[],
            signals_sha256="s", text_sha256="t", raw_sha256="r", in_sitemap=True,
        )
        repo.create_page_version(
            page, http_status=200, signals={}, signals_sha256="s", text_sha256="t",
            raw_sha256="r", body_blob_sha256=None, changed_fields=[], prev_version_id=None,
        )
    db.commit()
    client.post(f"/api/v1/clients/{client_id}/pages/rebuild")

    calibrated = client.get(f"/api/v1/clients/{client_id}/pages?profile=calibrated").json()
    strict = client.get(f"/api/v1/clients/{client_id}/pages?profile=strict").json()

    assert calibrated["summary"]["duplicate_groups"] == strict["summary"]["duplicate_groups"] == 1
    # Strict shows every pairwise finding (more of them); calibrated collapses
    # each row to one synthetic cluster finding per check (task-57 §2.3).
    assert strict["summary"]["severity_counts"]["error"] > calibrated["summary"]["severity_counts"]["error"]


def test_xlsx_export_has_two_sheets_with_expected_headers(client, db):
    client_id = _client(db)
    _seed_pages(db, client_id)
    client.post(f"/api/v1/clients/{client_id}/pages/rebuild")

    response = client.get(f"/api/v1/clients/{client_id}/pages/xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Inventory", "SEO QA"]
    inventory_headers = [cell.value for cell in workbook["Inventory"][1]]
    assert inventory_headers == [
        "Pages", "URL", "Breadcrumbs", "Title", "Description", "Keywords", "Status",
    ]
    qa_headers = [cell.value for cell in workbook["SEO QA"][1]]
    assert "Title issues" in qa_headers
    assert "Description issues" in qa_headers


def test_xlsx_export_404_for_client_never_built(client, db):
    client_id = _client(db)
    response = client.get(f"/api/v1/clients/{client_id}/pages/xlsx")
    assert response.status_code == 404
