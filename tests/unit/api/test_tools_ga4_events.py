"""API-тести GA4 Events by Month (task-98).

Мережі немає: `run_events_report`/`available_months` підмінені (той самий
підхід, що в `test_tools_schema_markup.py`), а персистенція, стани й експорт
перевіряються через реальну тестову БД.
"""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.api import tools_ga4_events as api_mod
from app.db.models import ClientIntegration
from app.db.repositories.clients import ClientRepository
from app.schemas.enums import Provider
from app.tools.ga4_events.service import (
    EventRow,
    Ga4EventsConfigError,
    MonthResult,
    ReportResult,
    SegmentTotals,
    SegmentValues,
)


def _create_client(db, *, slug: str = "shb", with_ga4: bool = True) -> str:
    client = ClientRepository(db).create(
        name="SHB", slug=slug, domain=f"{slug}.com", timezone="America/New_York"
    )
    if with_ga4:
        db.add(
            ClientIntegration(
                client_id=client.id,
                provider=Provider.reporting_ninja,
                config_json={"account_id": "properties/402129482"},
            )
        )
    db.commit()
    return client.id


def _result(months: list[MonthResult]) -> ReportResult:
    return ReportResult(
        property_id="properties/402129482",
        state="ok",
        months=months,
        generated_at=datetime(2026, 9, 17, tzinfo=UTC),
    )


def _month(month: str, label: str) -> MonthResult:
    return MonthResult(
        month=month,
        label=label,
        state="ok",
        is_partial=False,
        rows=[
            EventRow(
                event_name="Phone",
                is_key_event=True,
                key_events=120,
                segments={
                    "all": SegmentValues(
                        event_count=120,
                        users=100,
                        per_user=Decimal("1.20"),
                        event_share_pct=Decimal("60.00"),
                        user_share_pct=Decimal("66.67"),
                    ),
                    "direct": SegmentValues(
                        event_count=40,
                        users=35,
                        per_user=Decimal("1.14"),
                        event_share_pct=Decimal("40.00"),
                        user_share_pct=Decimal("50.00"),
                    ),
                },
            )
        ],
        totals={
            "all": SegmentTotals(event_count=200, users=150, key_events=120),
            "direct": SegmentTotals(event_count=100, users=70, key_events=40),
        },
    )


def test_clients_list_marks_who_has_a_ga4_property(client, db) -> None:
    configured = _create_client(db)
    plain = _create_client(db, slug="noga4", with_ga4=False)

    body = client.get("/api/v1/tools/ga4-events/clients").json()
    by_id = {row["id"]: row for row in body["clients"]}

    assert by_id[configured]["configured"] is True
    assert by_id[configured]["property_id"] == "properties/402129482"
    # Клієнт без GA4 не зникає зі списку — він у ньому з чесним `configured=false`.
    assert by_id[plain]["configured"] is False
    assert by_id[plain]["property_id"] is None


def test_report_is_persisted_and_readable_by_its_link(client, db, monkeypatch) -> None:
    client_id = _create_client(db)
    monkeypatch.setattr(
        api_mod,
        "run_events_report",
        lambda **kwargs: _result([_month("2026-08", "Aug 2026")]),
    )

    created = client.post(
        "/api/v1/tools/ga4-events/reports",
        json={"client_id": client_id, "months": ["2026-08"]},
    )
    assert created.status_code == 200
    body = created.json()
    assert body["client_name"] == "SHB"
    assert body["months"] == ["2026-08"]
    row = body["months_data"][0]["rows"][0]
    assert row["event_name"] == "Phone"
    assert row["segments"]["all"]["event_count"] == 120
    assert row["segments"]["all"]["event_share_pct"] == "60.00"
    assert row["segments"]["direct"]["users"] == 35
    # Поріз, якого GA4 не повернув, у відповіді просто відсутній — не нуль.
    assert "organic" not in row["segments"]
    assert body["months_data"][0]["totals"]["all"]["users"] == 150

    fetched = client.get(f"/api/v1/tools/ga4-events/reports/{body['id']}")
    assert fetched.status_code == 200
    # Знімок: читання повертає ті самі числа, не перезапитує GA4.
    assert fetched.json()["months_data"] == body["months_data"]


def test_client_without_ga4_is_not_configured_not_a_failure(client, db) -> None:
    client_id = _create_client(db, slug="noga4", with_ga4=False)

    response = client.post(
        "/api/v1/tools/ga4-events/reports",
        json={"client_id": client_id, "months": ["2026-08"]},
    )

    assert response.status_code == 422
    assert response.json()["detail"]["error_code"] == "ga4_not_configured"


def test_missing_refresh_token_surfaces_as_not_configured(client, db, monkeypatch) -> None:
    client_id = _create_client(db)

    def _boom(**kwargs):
        raise Ga4EventsConfigError("google refresh token is not set")

    monkeypatch.setattr(api_mod, "run_events_report", _boom)

    response = client.post(
        "/api/v1/tools/ga4-events/reports",
        json={"client_id": client_id, "months": ["2026-08"]},
    )

    assert response.status_code == 422
    assert response.json()["detail"]["error_code"] == "ga4_not_configured"


def test_invalid_or_oversized_month_selection_is_rejected(client, db) -> None:
    client_id = _create_client(db)

    bad = client.post(
        "/api/v1/tools/ga4-events/reports",
        json={"client_id": client_id, "months": ["2026-13"]},
    )
    assert bad.status_code == 422
    assert bad.json()["detail"]["error_code"] == "invalid_months"

    too_many = client.post(
        "/api/v1/tools/ga4-events/reports",
        json={
            "client_id": client_id,
            "months": [f"2025-{index:02d}" for index in range(1, 13)]
            + [f"2026-{index:02d}" for index in range(1, 13)]
            + ["2024-12"],
        },
    )
    assert too_many.status_code == 422
    assert too_many.json()["detail"]["error_code"] == "invalid_months"


def test_months_endpoint_returns_what_ga4_actually_has(client, db, monkeypatch) -> None:
    client_id = _create_client(db)
    monkeypatch.setattr(
        api_mod, "available_months", lambda **kwargs: ["2026-08", "2026-07"]
    )

    body = client.get(f"/api/v1/tools/ga4-events/clients/{client_id}/months").json()

    assert body["months"] == ["2026-08", "2026-07"]


def test_export_returns_a_workbook_with_one_sheet_per_month(client, db, monkeypatch) -> None:
    client_id = _create_client(db)
    monkeypatch.setattr(
        api_mod,
        "run_events_report",
        lambda **kwargs: _result(
            [_month("2026-08", "Aug 2026"), _month("2026-07", "Jul 2026")]
        ),
    )
    report_id = client.post(
        "/api/v1/tools/ga4-events/reports",
        json={"client_id": client_id, "months": ["2026-08", "2026-07"]},
    ).json()["id"]

    response = client.get(f"/api/v1/tools/ga4-events/reports/{report_id}/export.xlsx")

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    workbook = load_workbook(__import__("io").BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "2026-08", "2026-07"]


def test_unknown_report_is_404(client) -> None:
    assert client.get("/api/v1/tools/ga4-events/reports/nope").status_code == 404
