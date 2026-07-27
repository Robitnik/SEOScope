"""Наскрізні API-тести інструмента Redirect Checker (task-29)."""

from __future__ import annotations

import io
from typing import Any

import pytest
from openpyxl import load_workbook

from app.tools.redirects import trace as trace_module
from app.tools.redirects.trace import RawResponse

CHECKS = "/api/v1/tools/redirects/checks"


def _fake_network(monkeypatch: pytest.MonkeyPatch, routes: dict[str, RawResponse]) -> list[str]:
    """Підмінити реальний fetcher: жодного сокета в юніт-тестах."""
    requested: list[str] = []

    def fake_fetch(url: str, *, timeout: float = 10.0) -> RawResponse:
        requested.append(url)
        return routes.get(
            url, RawResponse(status=404, headers=(("content-type", "text/html"),), body="<h1>404")
        )

    monkeypatch.setattr(trace_module, "fetch_no_redirect", fake_fetch)
    return requested


def _redirect(location: str, status: int = 301) -> RawResponse:
    return RawResponse(status=status, headers=(("location", location),))


def _page(body: str = "<title>Landing</title>", status: int = 200) -> RawResponse:
    return RawResponse(status=status, headers=(("content-type", "text/html"),), body=body)


CHAIN = {
    "http://acme.example.com/old": _redirect("https://acme.example.com/old"),
    "https://acme.example.com/old": _redirect("https://acme.example.com/new/", 302),
    "https://acme.example.com/new/": _page(),
}


# --- каталог -------------------------------------------------------------------


def test_catalog_lists_both_tools(api_client: Any) -> None:
    response = api_client.get("/api/v1/tools")
    assert response.status_code == 200
    tools = {tool["key"]: tool for tool in response.json()["tools"]}
    assert {"alt-text", "redirects"} <= set(tools)
    assert tools["redirects"]["route"] == "#/tools/redirects"
    # Без історії — «не користувалися», а не «0 checks».
    assert tools["redirects"]["stat_value"] is None


def test_catalog_stat_counts_real_checks(api_client: Any) -> None:
    api_client.post(CHECKS, json={"urls": ["https://acme.example.com/a"]})
    tools = {tool["key"]: tool for tool in api_client.get("/api/v1/tools").json()["tools"]}
    assert tools["redirects"]["stat_value"] == "1"
    assert tools["redirects"]["stat_label"] == "check"


# --- швидкий трейс -------------------------------------------------------------


def test_trace_returns_the_whole_chain_without_persisting(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_network(monkeypatch, CHAIN)
    response = api_client.post(
        "/api/v1/tools/redirects/trace", json={"url": "http://acme.example.com/old"}
    )
    assert response.status_code == 200
    body = response.json()
    assert body["verdict"] == "chain"
    assert body["hop_count"] == 2
    assert [hop["status"] for hop in body["hops"]] == [301, 302, 200]
    assert body["hops"][0]["changes"] == ["scheme_http_to_https"]
    assert body["hops"][-1]["location"] is None
    assert body["final_url"] == "https://acme.example.com/new/"
    assert body["final_status"] == 200
    assert "temporary_redirect" in body["flags"]
    assert body["final_page"]["title"] == "Landing"
    assert body["checked_at"] is not None
    # Нічого не створили.
    assert api_client.get(CHECKS).json() == []


def test_trace_normalizes_a_pasted_link(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    requested = _fake_network(monkeypatch, {"https://acme.example.com/x": _page()})
    body = api_client.post(
        "/api/v1/tools/redirects/trace", json={"url": "acme.example.com/x#top"}
    ).json()
    assert requested == ["https://acme.example.com/x"]
    assert body["normalized_url"] == "https://acme.example.com/x"
    assert "normalized_input" in body["flags"]


def test_trace_rejects_a_non_http_link_without_touching_the_network(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    requested = _fake_network(monkeypatch, {})
    body = api_client.post(
        "/api/v1/tools/redirects/trace", json={"url": "mailto:hi@acme.example.com"}
    ).json()
    assert body["verdict"] == "unsupported_scheme"
    assert requested == []


def test_trace_respects_options(api_client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    routes = {
        "https://acme.example.com/1": _redirect("https://acme.example.com/2"),
        "https://acme.example.com/2": _redirect("https://acme.example.com/3"),
        "https://acme.example.com/3": _page(),
    }
    _fake_network(monkeypatch, routes)
    body = api_client.post(
        "/api/v1/tools/redirects/trace",
        json={"url": "https://acme.example.com/1", "options": {"max_hops": 1}},
    ).json()
    assert body["verdict"] == "too_many_hops"


def test_trace_validates_the_payload(api_client: Any) -> None:
    assert api_client.post("/api/v1/tools/redirects/trace", json={"url": ""}).status_code == 422
    assert (
        api_client.post(
            "/api/v1/tools/redirects/trace",
            json={"url": "https://a.example/", "options": {"max_hops": 500}},
        ).status_code
        == 422
    )


# --- створення прогону ---------------------------------------------------------


def test_create_check_normalizes_dedupes_and_queues(api_client: Any) -> None:
    response = api_client.post(
        CHECKS,
        json={
            "urls": [
                "acme.example.com/a\nhttps://acme.example.com/a\n"
                "https://acme.example.com/b\nmailto:hi@acme.example.com\nnot a url"
            ]
        },
    )
    assert response.status_code == 201
    body = response.json()
    assert body["target_count"] == 2
    assert body["input_count"] == 5
    assert body["duplicate_count"] == 1
    assert {row["error_code"] for row in body["rejected"]} == {
        "unsupported_scheme",
        "invalid_url",
    }
    assert body["run_status_counts"] == {"queued": 2, "running": 0, "done": 0, "failed": 0}
    assert body["verdict_counts"] == {}
    assert body["title"].startswith("Redirects · ")

    rows = api_client.get(f"{CHECKS}/{body['id']}/targets").json()["rows"]
    assert [row["run_status"] for row in rows] == ["queued", "queued"]
    # Неперевірений рядок: жодного вердикту й жодного часу перевірки.
    assert rows[0]["verdict"] is None
    assert rows[0]["checked_at"] is None
    assert rows[0]["input_notes"] == ["scheme_added"]


def test_create_check_can_be_attached_to_a_client(api_client: Any, seed_client: Any) -> None:
    client_id = seed_client()
    body = api_client.post(
        CHECKS, json={"urls": ["https://acme.example.com/a"], "client_id": client_id}
    ).json()
    assert body["client_id"] == client_id
    assert body["client_name"] == "Acme"
    assert body["title"].startswith("Acme · Redirects")


def test_create_check_rejects_an_unknown_client(api_client: Any) -> None:
    response = api_client.post(
        CHECKS, json={"urls": ["https://acme.example.com/a"], "client_id": "nope"}
    )
    assert response.status_code == 404
    assert response.json()["detail"]["error_code"] == "client_not_found"


def test_create_check_without_a_single_valid_url_is_422(api_client: Any) -> None:
    response = api_client.post(CHECKS, json={"urls": ["mailto:hi@a.example", "garbage"]})
    assert response.status_code == 422
    assert response.json()["detail"]["error_code"] == "no_valid_urls"


def test_create_check_caps_the_pasted_list(api_client: Any) -> None:
    body = api_client.post(
        CHECKS,
        json={"urls": ["\n".join(f"https://acme.example.com/{index}" for index in range(250))]},
    ).json()
    assert body["target_count"] == 200
    assert body["truncated"] is True
    assert body["input_count"] == 250


def test_create_check_refuses_an_absurd_paste(api_client: Any) -> None:
    response = api_client.post(
        CHECKS,
        json={"urls": ["\n".join(f"https://acme.example.com/{index}" for index in range(1100))]},
    )
    assert response.status_code == 413
    assert response.json()["detail"]["error_code"] == "too_many_urls"


def test_create_check_stores_the_options_used(api_client: Any) -> None:
    body = api_client.post(
        CHECKS,
        json={
            "urls": ["https://acme.example.com/a"],
            "options": {"max_hops": 5, "follow_meta_refresh": False, "cache_buster": True},
        },
    ).json()
    assert body["options"] == {
        "max_hops": 5,
        "follow_meta_refresh": False,
        "cache_buster": True,
    }


# --- прогін рядків -------------------------------------------------------------


def _seeded_check(api_client: Any, urls: list[str] | None = None) -> tuple[str, list[Any]]:
    payload = {"urls": urls or ["http://acme.example.com/old"]}
    check = api_client.post(CHECKS, json=payload).json()
    rows = api_client.get(f"{CHECKS}/{check['id']}/targets").json()["rows"]
    return check["id"], rows


def test_run_target_traces_and_persists(api_client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_network(monkeypatch, CHAIN)
    check_id, rows = _seeded_check(api_client)

    response = api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run")
    assert response.status_code == 200
    body = response.json()
    assert body["run_status"] == "done"
    assert body["verdict"] == "chain"
    assert body["hop_count"] == 2
    assert body["checked_at"] is not None
    assert len(body["hops"]) == 3

    # Результат переживає перечитування — він у БД, не в памʼяті процесу.
    reread = api_client.get(f"{CHECKS}/{check_id}/targets").json()["rows"][0]
    assert reread["verdict"] == "chain"
    assert reread["hops"][1]["status"] == 302
    check = api_client.get(f"{CHECKS}/{check_id}").json()
    assert check["run_status_counts"]["done"] == 1
    assert check["verdict_counts"] == {"chain": 1}


def test_run_target_is_repeatable(api_client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_network(monkeypatch, CHAIN)
    check_id, rows = _seeded_check(api_client)
    first = api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run").json()
    second = api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run").json()
    assert first["verdict"] == second["verdict"] == "chain"


def test_run_target_rejects_a_parallel_run(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Подвійний клік не має дати два запити до чужого сайту."""
    check_id, rows = _seeded_check(api_client)
    target_id = rows[0]["id"]
    calls: list[str] = []

    def reentrant_fetch(url: str, *, timeout: float = 10.0) -> RawResponse:
        calls.append(url)
        nested = api_client.post(f"{CHECKS}/{check_id}/targets/{target_id}/run")
        assert nested.status_code == 409
        assert nested.json()["detail"]["error_code"] == "target_running"
        return _page()

    monkeypatch.setattr(trace_module, "fetch_no_redirect", reentrant_fetch)
    assert api_client.post(f"{CHECKS}/{check_id}/targets/{target_id}/run").status_code == 200
    assert len(calls) == 1


def test_run_target_404s_for_a_foreign_row(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_network(monkeypatch, CHAIN)
    first_id, _rows = _seeded_check(api_client)
    _second_id, other_rows = _seeded_check(api_client, ["https://acme.example.com/other"])
    response = api_client.post(f"{CHECKS}/{first_id}/targets/{other_rows[0]['id']}/run")
    assert response.status_code == 404
    assert response.json()["detail"]["error_code"] == "target_not_found"


def test_run_target_404s_for_an_unknown_check(api_client: Any) -> None:
    assert api_client.post(f"{CHECKS}/nope/targets/nope/run").status_code == 404


@pytest.mark.parametrize(
    ("routes", "verdict"),
    [
        ({"https://acme.example.com/x": _page("<h1>Gone", status=404)}, "broken"),
        ({"https://acme.example.com/x": _page("<h1>Denied", status=403)}, "blocked"),
        ({"https://acme.example.com/x": _redirect("https://acme.example.com/x")}, "loop"),
        ({"https://acme.example.com/x": RawResponse(status=301)}, "broken_chain"),
        ({"https://acme.example.com/x": _page()}, "ok_direct"),
    ],
)
def test_verdicts_are_distinct_states(
    api_client: Any,
    monkeypatch: pytest.MonkeyPatch,
    routes: dict[str, RawResponse],
    verdict: str,
) -> None:
    _fake_network(monkeypatch, routes)
    check_id, rows = _seeded_check(api_client, ["https://acme.example.com/x"])
    body = api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run").json()
    assert body["verdict"] == verdict


def test_fetch_failure_is_stored_with_a_reason_and_no_leak(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    import urllib.error

    def failing(url: str, *, timeout: float = 10.0) -> RawResponse:
        raise urllib.error.URLError("Name or service not known")

    monkeypatch.setattr(trace_module, "fetch_no_redirect", failing)
    check_id, rows = _seeded_check(api_client, ["https://acme.example.com/x?token=secret"])
    body = api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run").json()
    assert body["verdict"] == "fetch_failed"
    assert body["run_status"] == "done"
    assert "secret" not in body["error_detail"]


def test_unexpected_failure_does_not_leave_the_row_running(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    def exploding(url: str, *, timeout: float = 10.0) -> RawResponse:
        raise MemoryError("boom")

    monkeypatch.setattr(trace_module, "fetch_no_redirect", exploding)
    check_id, rows = _seeded_check(api_client)
    response = api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run")
    assert response.status_code == 500
    row = api_client.get(f"{CHECKS}/{check_id}/targets").json()["rows"][0]
    assert row["run_status"] == "failed"
    assert row["error_code"] == "run_failed"


# --- читання, фільтри, видалення -----------------------------------------------


def test_list_and_get_check(api_client: Any) -> None:
    check_id, _rows = _seeded_check(api_client)
    listed = api_client.get(CHECKS).json()
    assert [check["id"] for check in listed] == [check_id]
    assert api_client.get(f"{CHECKS}/{check_id}").json()["id"] == check_id
    assert api_client.get(f"{CHECKS}/nope").status_code == 404
    assert api_client.get(f"{CHECKS}/nope/targets").status_code == 404


def test_target_filters(api_client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_network(
        monkeypatch,
        {
            "https://acme.example.com/a": _page(),
            "https://acme.example.com/b": _redirect("https://acme.example.com/c", 302),
            "https://acme.example.com/c": _page(),
        },
    )
    check_id, rows = _seeded_check(
        api_client, ["https://acme.example.com/a\nhttps://acme.example.com/b"]
    )
    for row in rows:
        api_client.post(f"{CHECKS}/{check_id}/targets/{row['id']}/run")

    by_verdict = api_client.get(f"{CHECKS}/{check_id}/targets?verdict=ok_redirect").json()
    assert by_verdict["total"] == 1
    by_flag = api_client.get(f"{CHECKS}/{check_id}/targets?flag=temporary_redirect").json()
    assert by_flag["total"] == 1
    by_query = api_client.get(f"{CHECKS}/{check_id}/targets?q=example.com/a").json()
    assert by_query["total"] == 1
    paged = api_client.get(f"{CHECKS}/{check_id}/targets?limit=1&offset=1").json()
    assert paged["total"] == 2
    assert len(paged["rows"]) == 1
    assert paged["offset"] == 1


def test_delete_check_cascades(api_client: Any) -> None:
    check_id, _rows = _seeded_check(api_client)
    assert api_client.delete(f"{CHECKS}/{check_id}").status_code == 204
    assert api_client.get(f"{CHECKS}/{check_id}").status_code == 404
    assert api_client.delete(f"{CHECKS}/{check_id}").status_code == 404


# --- експорт -------------------------------------------------------------------


def test_csv_export_contains_the_chain_as_text(
    api_client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_network(monkeypatch, CHAIN)
    check_id, rows = _seeded_check(api_client)
    api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run")

    response = api_client.get(f"{CHECKS}/{check_id}/export.csv")
    assert response.status_code == 200
    text = response.text
    assert "301 → 302 → 200" in text
    assert "https://acme.example.com/new/" in text
    assert "chain" in text


def test_csv_export_escapes_formulas(api_client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_network(
        monkeypatch,
        {"https://acme.example.com/x": _page('<title>=SUM(A1:A9)</title>')},
    )
    check_id, rows = _seeded_check(api_client, ["https://acme.example.com/x"])
    api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run")
    assert "'=SUM(A1:A9)" in api_client.get(f"{CHECKS}/{check_id}/export.csv").text


def test_export_survives_a_non_ascii_title(api_client: Any) -> None:
    check = api_client.post(
        CHECKS, json={"urls": ["https://acme.example.com/a"], "title": "Перевірка — липень"}
    ).json()
    for suffix in ("csv", "xlsx"):
        response = api_client.get(f"{CHECKS}/{check['id']}/export.{suffix}")
        assert response.status_code == 200
        assert "filename*=UTF-8''" in response.headers["content-disposition"]


def test_xlsx_export_has_a_hop_sheet(api_client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_network(monkeypatch, CHAIN)
    check_id, rows = _seeded_check(api_client)
    api_client.post(f"{CHECKS}/{check_id}/targets/{rows[0]['id']}/run")

    response = api_client.get(f"{CHECKS}/{check_id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Redirects", "Hops"]
    hop_rows = list(workbook["Hops"].iter_rows(values_only=True))
    assert hop_rows[0][0] == "#"
    assert len(hop_rows) == 4  # шапка + три хопи


def test_export_404s_for_an_unknown_check(api_client: Any) -> None:
    assert api_client.get(f"{CHECKS}/nope/export.csv").status_code == 404
    assert api_client.get(f"{CHECKS}/nope/export.xlsx").status_code == 404


def test_unchecked_rows_export_empty_not_zero(api_client: Any) -> None:
    check_id, _rows = _seeded_check(api_client)
    text = api_client.get(f"{CHECKS}/{check_id}/export.csv").text
    body = text.splitlines()[1].split(",")
    # Колонка «Redirects» лишається порожньою: «не перевіряли» ≠ «0 редіректів».
    assert body[4] == ""
