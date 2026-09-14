from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.contentqa.readers.tabular import sheet_role
from app.contentqa.readers.xlsx import (
    UnknownSheetError,
    inspect_strategy_xlsx,
    read_strategy_xlsx,
)


def _build_workbook() -> bytes:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Q3 Strategy"
    ws1.append(["Q3 2026 Strategy"])
    ws1.append([
        "Page Status", "Intent", "Page Name", "URL", "New URL",
        "Primary Key", "Secondary Key", "Longtail Key", "H1", "Title", "Description",
    ])
    ws1.append([
        "OPTIMIZE PAGE", "Navigational", "Homepage",
        "https://example.com/", "https://example.com/",
        "building surveying software", "survey reporting software",
        "RICS compliant survey software",
        "Building Surveying Software for Digital Capture, Reporting and Analysis",
        "Building Surveying Software UK | Digital Reports and RICS Templates | GoReport",
        "GoReport is UK building surveying software for mobile capture and reporting.",
    ])
    ws1.append([
        "OPTIMIZE PAGE", "Commercial", "Pricing",
        "https://example.com/pricing/", "https://example.com/pricing/",
        "survey software pricing", "", "",
        "Survey Software Pricing", "Survey Software Pricing | GoReport",
        "See plans and pricing for our survey software.",
    ])

    ws2 = wb.create_sheet("Structure v 2.0")
    ws2.append(["Pages", "URL", "Recommended URL (redirect to...)", "Breadcrumbs", "Comments"])
    ws2.append(["Home", "https://example.com/", "https://example.com/", "Home", "Q3 #1"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sheet_without_title_or_description_columns_is_skipped() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    assert doc.skipped_sheets == ("Structure v 2.0",)


def test_recognized_sheet_yields_two_rows() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    assert len(doc.items) == 2


def test_column_mapping_by_header_alias() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    home = doc.items[0]
    assert home.page_name == "Homepage"
    assert home.url == "https://example.com/"
    assert home.intent == "Navigational"
    expected_title = (
        "Building Surveying Software UK | Digital Reports and RICS Templates | GoReport"
    )
    assert home.title == expected_title
    assert home.h1 == "Building Surveying Software for Digital Capture, Reporting and Analysis"
    assert home.keywords == (
        "building surveying software", "survey reporting software",
        "RICS compliant survey software",
    )


def test_ref_carries_sheet_name_and_row_number() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    assert doc.items[0].ref == "Q3 Strategy!row 3"
    assert doc.items[1].ref == "Q3 Strategy!row 4"


def test_missing_secondary_longtail_keys_do_not_pad_with_empty_strings() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    pricing = doc.items[1]
    assert pricing.keywords == ("survey software pricing",)


def test_item_carries_its_source_sheet_name() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    assert doc.items[0].sheet == "Q3 Strategy"


def test_dash_placeholder_url_is_treated_as_absent_not_a_real_url() -> None:
    # Живий баг (3XE_SEO_STRATEGY_08032026): оператор заповнює порожню
    # клітинку URL буквальним "-" замість пустої. Якщо тримати це як
    # справжнє значення, усі рядки з "-" схлопуються в ОДИН ключ
    # зіставлення в `crosscheck.py:_item_key()` — кожен фреймворк тоді
    # "зіставляється" з останнім таким рядком стратегії замість свого.
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 Strategy"
    ws.append(["Page Name", "URL", "Title", "Description"])
    ws.append(["Page One", "-", "Title One", "Desc One"])
    ws.append(["Page Two", "-", "Title Two", "Desc Two"])
    buf = io.BytesIO()
    wb.save(buf)

    doc = read_strategy_xlsx(buf.getvalue(), source="test.xlsx")
    assert [item.url for item in doc.items] == [None, None]
    assert [item.page_name for item in doc.items] == ["Page One", "Page Two"]


# ── task-87 §1.3: вибір аркуша ────────────────────────────────────────────────


def test_default_sheets_none_checks_everything_like_before_task_87() -> None:
    doc = read_strategy_xlsx(_build_workbook(), source="test.xlsx")
    states = {s.name: s.state for s in doc.sheet_states}
    assert states == {"Q3 Strategy": "checked", "Structure v 2.0": "no_recognized_columns"}


def test_selecting_a_subset_of_sheets_limits_items() -> None:
    wb = _build_workbook()
    doc = read_strategy_xlsx(wb, source="test.xlsx", sheets=["Q3 Strategy"])
    assert len(doc.items) == 2  # той самий workbook, дефолт теж дає 2 — тут це єдиний несучий аркуш


def test_unselected_sheet_is_not_selected_not_skipped() -> None:
    # Другий несучий аркуш у тому самому workbook, щоб перевірити, що
    # "не обраний" (not_selected) — не те саме, що "нема title/description"
    # (no_recognized_columns).
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Q2"
    ws1.append(["Title", "Description"])
    ws1.append(["T1", "D1"])
    ws2 = wb.create_sheet("Q3")
    ws2.append(["Title", "Description"])
    ws2.append(["T2", "D2"])
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    doc = read_strategy_xlsx(content, source="test.xlsx", sheets=["Q3"])
    states = {s.name: s.state for s in doc.sheet_states}
    assert states == {"Q2": "not_selected", "Q3": "checked"}
    assert doc.skipped_sheets == ()  # not_selected НЕ потрапляє в skipped_sheets
    assert [i.ref for i in doc.items] == ["Q3!row 2"]


def test_unknown_sheet_raises_not_silently_checks_everything() -> None:
    with pytest.raises(UnknownSheetError):
        read_strategy_xlsx(_build_workbook(), source="test.xlsx", sheets=["Does Not Exist"])


def test_inspect_reports_load_bearing_and_row_counts_without_building_items() -> None:
    inspection = inspect_strategy_xlsx(_build_workbook())
    by_name = {s.name: s for s in inspection.sheets}
    assert by_name["Q3 Strategy"].load_bearing is True
    assert by_name["Q3 Strategy"].row_count == 2
    assert by_name["Structure v 2.0"].load_bearing is False
    assert inspection.suggested_sheets == ("Q3 Strategy",)


# ── task-97: роль аркуша (page/blog/structure/other) ───────────────────────


@pytest.mark.parametrize(
    ("name", "columns", "expected"),
    [
        ("Blog", (), "blog"),
        ("Blogs", (), "blog"),
        ("Blogs Q3-Q4", (), "blog"),
        ("Structure / Features / Description", (), "structure"),
        ("Q3 2026 Strategy", (), "page"),
        ("Q4 2026", (), "page"),
        ("Keywords", (), "other"),
        # Назва нічого не каже — фолбек на колонки (§1.4/§1.5 плану).
        ("Sheet1", ("url", "intent", "title"), "page"),
        ("Sheet1", ("keywords_csv", "title"), "blog"),
        ("Sheet1", (), "other"),
    ],
)
def test_sheet_role_by_name_with_column_fallback(name, columns, expected) -> None:
    assert sheet_role(name, columns) == expected


def test_inspect_prefers_page_sheets_over_all_load_bearing_sheets() -> None:
    """Дефолт форми (task-97 §2.1) — несучі аркуші з роллю `page`, а не
    "усі несучі" (task-90 дефолт): `Blog` тут теж несучий (title+description),
    але не входить у `suggested_sheets`, бо тул задуманий для сторінок."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Q3 2026 Strategy"
    ws1.append(["Title", "Description"])
    ws1.append(["Page title", "Page description"])
    ws2 = wb.create_sheet("Blog")
    ws2.append(["Title", "Description"])
    ws2.append(["Blog title", "Blog description"])
    buf = io.BytesIO()
    wb.save(buf)

    inspection = inspect_strategy_xlsx(buf.getvalue())
    assert inspection.suggested_sheets == ("Q3 2026 Strategy",)


def test_inspect_falls_back_to_all_load_bearing_when_no_page_sheet_found() -> None:
    """Коли жоден несучий аркуш не розпізнався як `page` — старе
    поводження (усі несучі), не порожній дефолт (§2.1 плану: широкий
    дефолт кращий за порожній)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Blog"
    ws.append(["Title", "Description"])
    ws.append(["Blog title", "Blog description"])
    buf = io.BytesIO()
    wb.save(buf)

    inspection = inspect_strategy_xlsx(buf.getvalue())
    assert inspection.suggested_sheets == ("Blog",)
