from __future__ import annotations

from dataclasses import replace
from datetime import date

import pytest
from openpyxl import load_workbook

from app.reports.batch_xlsx import NoClientDataError, build_portfolio_xlsx
from app.reports.history import HistoryMonth, ReportHistory
from app.reports.models import PeriodMetrics
from app.reports.periods import ReportPeriod
from app.reports.table import build_table


def _table(client_users: int) -> object:
    periods = [
        ReportPeriod("This Month", date(2026, 7, 1), date(2026, 7, 31)),
        ReportPeriod("Last Month", date(2026, 6, 1), date(2026, 6, 30)),
    ]
    metrics = {p.label: PeriodMetrics(users=client_users, clicks=10) for p in periods}
    return build_table("custom_months", periods, metrics)


def test_one_sheet_per_client_with_both_tables_stacked(tmp_path):
    swm = replace(
        _table(1027),
        history=ReportHistory(months=[HistoryMonth("2026-07", "July 2026", 113, 4432, None, 111)]),
    )
    grp = _table(1634)

    path = build_portfolio_xlsx(
        [("SWM", swm), ("GRP", grp)], tmp_path / "portfolio.xlsx"
    )

    wb = load_workbook(path)
    assert wb.sheetnames == ["SWM", "GRP"]

    ws = wb["SWM"]
    # The metrics table starts at row 1 (title) — same layout `xlsx.py`'s
    # own `Report` sheet uses for a single client. 8 METRICS rows (3..10)
    # means the next table starts at row 12 (row 11 is the first free row
    # `write_metrics_table` returns, plus one blank row).
    assert ws["A1"].value == "SWM"
    assert ws["A2"].value == "Metric"
    assert ws["B3"].value == 1027
    # The historical table is stacked BELOW it on the same sheet, not on a
    # separate tab — that's the whole point of the portfolio layout.
    assert ws["A12"].value == "SWM — Historical Data (complete months)"
    assert ws["A13"].value == "Date"
    assert ws["A14"].value == "July 2026"
    assert ws["B14"].value == 113


def test_a_client_with_no_history_still_gets_a_labelled_placeholder(tmp_path):
    path = build_portfolio_xlsx([("GRP", _table(1634))], tmp_path / "portfolio.xlsx")
    ws = load_workbook(path)["GRP"]
    # No history object at all -> `write_historical_table`'s own "could not
    # be collected" placeholder, same as the single-client XLSX.
    history_title_row = next(
        r for r in range(1, 20)
        if ws.cell(r, 1).value and "Historical Data" in str(ws.cell(r, 1).value)
    )
    assert "could not be collected" in str(ws.cell(history_title_row + 2, 1).value)


def test_duplicate_or_overlong_client_names_get_unique_sheet_names(tmp_path):
    long_name = "A" * 40
    path = build_portfolio_xlsx(
        [(long_name, _table(1)), (long_name, _table(2))], tmp_path / "portfolio.xlsx"
    )
    wb = load_workbook(path)
    assert len(wb.sheetnames) == 2
    assert len(set(wb.sheetnames)) == 2
    assert all(len(name) <= 31 for name in wb.sheetnames)


def test_empty_rows_raise_instead_of_producing_a_sheetless_workbook(tmp_path):
    with pytest.raises(NoClientDataError):
        build_portfolio_xlsx([], tmp_path / "portfolio.xlsx")
