"""XLSX export (task-54 §7): the exact header colors from the operator's own
sheet, column order, per-column number formats, and MoM direction as colored
text — no arrow glyph, since the source sheet has none.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.reports.monthly_table import MonthlyRow
from app.reports.monthly_xlsx import build_monthly_xlsx

_HEADER_GREEN = "0093C47D"
_HEADER_YELLOW = "00FFFF00"
_HEADER_BLUE = "006D9EEB"
_HEADER_RED = "00E06666"
_UP = "001A7A3C"
_DOWN = "00CC0000"


def _row(**overrides: object) -> MonthlyRow:
    base = dict(
        month="2026-07",
        label="July 2026",
        days_in_month=31,
        organic_users=721,
        organic_conversions=109,
        key_event_count=70,
        gsc_clicks=717,
        gsc_impressions=69176,
        gsc_days_with_data=31,
        conversion_rate=0.1512,
        ctr=0.0104,
        is_partial_period=False,
        ga4_state="ok",
        gsc_state="ok",
        truncated=False,
        directions={},
    )
    base.update(overrides)
    return MonthlyRow(**base)  # type: ignore[arg-type]


def _load(rows: list[MonthlyRow], *, key_event_label="Book Now", has_key_event_column=True):
    data = build_monthly_xlsx(
        rows, key_event_label=key_event_label, has_key_event_column=has_key_event_column
    )
    return load_workbook(io.BytesIO(data))


def test_header_colors_match_the_operators_sheet_exactly():
    wb = _load([_row()])
    ws = wb.active
    fills = [ws.cell(1, c).fill.fgColor.rgb for c in range(2, 9)]
    assert fills == [
        _HEADER_GREEN,
        _HEADER_YELLOW,
        _HEADER_BLUE,
        _HEADER_RED,
        _HEADER_GREEN,
        _HEADER_GREEN,
        _HEADER_BLUE,
    ]


def test_column_order_and_labels():
    wb = _load([_row()])
    ws = wb.active
    labels = [ws.cell(1, c).value for c in range(2, 9)]
    assert labels == [
        "Total organic Users",
        "Impressions",
        "Clicks",
        "Organic Conversions",
        "Conversion rate",
        'Organic "Book Now"',
        "CTR",
    ]


def test_key_event_column_hidden_when_not_configured():
    wb = _load([_row()], has_key_event_column=False)
    ws = wb.active
    assert ws.max_column == 7  # month + 6 metric columns, the named-event column dropped
    labels = [ws.cell(1, c).value for c in range(2, 8)]
    assert not any("Book Now" in str(label) for label in labels)


def test_named_event_label_falls_back_when_missing():
    wb = _load([_row()], key_event_label=None)
    ws = wb.active
    assert ws.cell(1, 7).value == 'Organic "Event"'


def test_number_formats_no_grouping_except_impressions():
    wb = _load([_row()])
    ws = wb.active
    assert ws.cell(2, 2).number_format == "0"  # Total organic Users
    assert ws.cell(2, 3).number_format == "#,##0"  # Impressions
    assert ws.cell(2, 4).number_format == "0"  # Clicks
    assert ws.cell(2, 5).number_format == "0"  # Organic Conversions
    assert ws.cell(2, 6).number_format == "0.00%"  # Conversion rate
    assert ws.cell(2, 8).number_format == "0.00%"  # CTR


def test_missing_value_is_em_dash_not_zero():
    row = _row(gsc_clicks=None, gsc_impressions=None, ctr=None)
    wb = _load([row])
    ws = wb.active
    assert ws.cell(2, 3).value == "—"
    assert ws.cell(2, 4).value == "—"
    assert ws.cell(2, 8).value == "—"
    assert ws.cell(2, 2).value == 721  # a present value is untouched


def test_direction_colors_up_and_down_no_arrow_glyph():
    row = _row(directions={"organic_users": "up", "organic_conversions": "down"})
    wb = _load([row])
    ws = wb.active
    users_cell = ws.cell(2, 2)
    conv_cell = ws.cell(2, 5)
    assert users_cell.font.color.rgb == _UP
    assert conv_cell.font.color.rgb == _DOWN
    assert "▲" not in str(users_cell.value)
    assert "▼" not in str(conv_cell.value)
    assert users_cell.value == 721


def test_flat_or_missing_direction_uses_default_font_color():
    row = _row(directions={"organic_users": "flat", "organic_conversions": None})
    wb = _load([row])
    ws = wb.active
    assert ws.cell(2, 2).font.color is None
    assert ws.cell(2, 5).font.color is None
