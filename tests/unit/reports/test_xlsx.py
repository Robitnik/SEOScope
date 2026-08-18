from __future__ import annotations

from dataclasses import replace
from datetime import date

import pytest
from openpyxl import load_workbook

from app.reports.history import HistoryMonth, ReportHistory
from app.reports.models import PeriodMetrics
from app.reports.periods import ReportPeriod
from app.reports.table import build_table
from app.reports.verify import cross_check
from app.reports.xlsx import build_xlsx


def _history() -> ReportHistory:
    return ReportHistory(
        months=[
            HistoryMonth("2026-06", "June 2026", 2000, 30000, 2000 / 30000, 600),
            HistoryMonth("2026-07", "July 2026", 2528, 31564, 2528 / 31564, 654),
        ]
    )


def _monthly_table() -> object:
    periods = [
        ReportPeriod("This Month", date(2026, 5, 1), date(2026, 5, 31)),
        ReportPeriod("Last Month", date(2026, 4, 1), date(2026, 4, 30)),
        ReportPeriod("Month YOY", date(2025, 5, 1), date(2025, 5, 31)),
        ReportPeriod("This Quarter", date(2026, 3, 1), date(2026, 5, 31)),
        ReportPeriod("Last Quarter", date(2025, 12, 1), date(2026, 2, 28)),
        ReportPeriod("Quarterly YOY", date(2025, 3, 1), date(2025, 5, 31)),
    ]
    metrics = {
        p.label: PeriodMetrics(
            users=6251, organic_users=3061, direct_users=2858,
            conversions=1099, organic_conversions=654, direct_conversions=391,
            clicks=2528, impressions=31564,
        )
        for p in periods
    }
    return build_table("monthly", periods, metrics)


def test_api_path_xlsx_has_report_history_and_import_sheets_only(tmp_path):
    table = _monthly_table()
    path = build_xlsx("EFFYS", table, tmp_path, generated_at=date(2026, 8, 17))

    assert path.exists()
    wb = load_workbook(path)
    # Provenance/cross-check stay PDF-only; Historical Data is on both paths
    # (task-45 §2.3/§2.4).
    assert wb.sheetnames == ["Report", "Historical Data", "Import Data"]

    ws = wb["Report"]
    assert ws["A1"].value == "EFFYS"
    assert ws["A2"].value == "Metric"
    assert ws["B2"].value == "This Month"
    assert ws["H2"].value == "% M/M"
    assert ws["A3"].value == "Users"
    assert ws["B3"].value == 6251
    assert str(ws["H3"].value).startswith("=IFERROR(")


def test_pdf_path_xlsx_adds_sources_and_cross_check_sheets(tmp_path):
    table = build_table(
        "monthly",
        _monthly_table().periods,
        {p.label: PeriodMetrics(users=6251, clicks=2528) for p in _monthly_table().periods},
        sources={"This Month": "EFFYS__this_month__2026-07-01_2026-07-31.pdf"},
    )
    pdf_metrics = {p.label: PeriodMetrics(users=6251, clicks=2528) for p in table.periods}
    api_metrics = {p.label: PeriodMetrics(users=6300, clicks=2528) for p in table.periods}
    result = cross_check(table.periods, pdf_metrics, api_metrics, tolerance_pct=0.5)

    path = build_xlsx(
        "EFFYS", table, tmp_path, generated_at=date(2026, 8, 17), cross_check=result
    )
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Report", "Data Sources", "Cross-check", "Historical Data", "Import Data"
    ]

    sources_ws = wb["Data Sources"]
    assert sources_ws["C3"].value == "EFFYS__this_month__2026-07-01_2026-07-31.pdf"

    cc_ws = wb["Cross-check"]
    assert "MISMATCH" in str(cc_ws["A1"].value) or "WITHIN_TOLERANCE" in str(cc_ws["A1"].value)
    assert cc_ws["A2"].value == "Period"


def test_import_data_sheet_has_plain_numbers_not_formulas(tmp_path):
    periods = [
        ReportPeriod("This Month", date(2026, 5, 1), date(2026, 5, 31)),
        ReportPeriod("Last Month", date(2026, 4, 1), date(2026, 4, 30)),
    ]
    table = build_table(
        "custom_months",
        periods,
        {"This Month": PeriodMetrics(users=110), "Last Month": PeriodMetrics(users=100)},
    )
    path = build_xlsx("EFFYS", table, tmp_path, generated_at=date(2026, 8, 17))
    wb = load_workbook(path)
    ws = wb["Import Data"]
    assert ws["A1"].value == "Metric"
    assert ws["B2"].value == 110
    assert ws["D2"].value == pytest.approx(0.1)
    assert "=" not in str(ws["D2"].value)


def test_history_sheet_writes_a_live_ctr_formula_not_a_number(tmp_path):
    table = replace(_monthly_table(), history=_history())
    path = build_xlsx("EFFYS", table, tmp_path, generated_at=date(2026, 8, 17))
    ws = load_workbook(path)["Historical Data"]

    assert ws["A2"].value == "Date"
    assert ws["E2"].value == "Conversions (Organic)"
    # Oldest month first, exactly as the hand-kept Google Sheets history read.
    assert ws["A3"].value == "June 2026"
    assert ws["A4"].value == "July 2026"
    assert ws["B4"].value == 2528
    assert ws["C4"].value == 31564
    assert ws["E4"].value == 654
    # A formula, so a row pasted into Google Sheets stays live (task-45 §2.4).
    assert ws["D4"].value == '=IFERROR(B4/C4,"")'


def test_history_sheet_leaves_a_month_without_data_blank(tmp_path):
    history = ReportHistory(
        months=[
            HistoryMonth("2025-01", "January 2025", None, None, None, None),
            HistoryMonth("2025-02", "February 2025", 10, 100, 0.1, 3),
        ]
    )
    table = replace(_monthly_table(), history=history)
    path = build_xlsx("3XE", table, tmp_path, generated_at=date(2026, 8, 17))
    ws = load_workbook(path)["Historical Data"]

    assert ws["B3"].value is None
    assert ws["C3"].value is None
    assert ws["E3"].value is None
    # The CTR formula still sits there and evaluates to "" — never 0.00%.
    assert ws["D3"].value == '=IFERROR(B3/C3,"")'


def test_history_sheet_exists_even_when_the_history_query_failed(tmp_path):
    path = build_xlsx("EFFYS", _monthly_table(), tmp_path, generated_at=date(2026, 8, 17))
    ws = load_workbook(path)["Historical Data"]
    assert "could not be collected" in str(ws["A3"].value)


def test_history_sheet_says_why_clicks_are_missing_without_gsc(tmp_path):
    history = ReportHistory(
        months=[HistoryMonth("2026-07", "July 2026", None, None, None, 654)], has_gsc=False
    )
    table = replace(_monthly_table(), history=history)
    path = build_xlsx("LAXTON", table, tmp_path, generated_at=date(2026, 8, 17))
    ws = load_workbook(path)["Historical Data"]
    note = " ".join(str(ws.cell(row, 1).value or "") for row in range(3, 8))
    assert "Google Search Console is not connected" in note


def test_missing_metric_renders_as_blank_not_zero(tmp_path):
    periods = [ReportPeriod("This Month", date(2026, 5, 1), date(2026, 5, 31))]
    table = build_table("custom_months", periods, {})
    path = build_xlsx("LAXTON", table, tmp_path, generated_at=date(2026, 8, 17))
    wb = load_workbook(path)
    ws = wb["Report"]
    assert ws["B3"].value is None
