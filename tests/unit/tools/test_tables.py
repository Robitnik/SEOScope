"""Тести читання таблиці аудиту: CSV/TSV, XLSX із вкладками, автомапінг (task-27)."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.tools.alt_text.tables import (
    MAX_CELL_LENGTH,
    MAX_PAGE_URLS,
    MAX_ROWS,
    ColumnMappingError,
    SourceKind,
    UnreadableSourceError,
    UnsupportedSourceError,
    detect_source_kind,
    inspect_source,
    parse_rows,
    split_page_urls,
    suggest_column_map,
)

UPLOADS = "https://example.com/wp-content/uploads/2024/03"


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


AUDIT_SHEET = [
    ["Image URL", "New alt text", "Page URL", "Current alt", "Notes"],
    [f"{UPLOADS}/hero-300x200.jpg", "Team at the office",
     "https://example.com/about", "", "high priority"],
    [f"{UPLOADS}/logo-scaled.png", "Yovth logo", "https://example.com/", "old logo", ""],
]


# --- визначення формату -------------------------------------------------------


def test_detect_xlsx_by_magic() -> None:
    assert detect_source_kind("report.xlsx", _xlsx({"Sheet1": AUDIT_SHEET})) is SourceKind.xlsx


def test_detect_csv_and_tsv() -> None:
    assert detect_source_kind("a.csv", b"a,b\n1,2\n") is SourceKind.csv
    assert detect_source_kind("a.tsv", b"a\tb\n1\t2\n") is SourceKind.tsv


def test_legacy_xls_is_rejected_with_an_actionable_message() -> None:
    with pytest.raises(UnsupportedSourceError) as exc:
        detect_source_kind("audit.xls", b"\xd0\xcf\x11\xe0legacy")
    assert ".xlsx" in str(exc.value)
    assert exc.value.code == "unsupported_source"


def test_binary_garbage_is_rejected() -> None:
    with pytest.raises(UnsupportedSourceError):
        detect_source_kind("photo.png", b"\x89PNG\r\n\x1a\n\x00\x00binary")


# --- автомапінг колонок -------------------------------------------------------


def test_suggest_column_map_on_a_typical_audit_header() -> None:
    mapping = suggest_column_map(("Image URL", "New alt text", "Page URL", "Current alt", "Notes"))
    assert mapping == {"image_url": 0, "new_alt": 1, "page_url": 2, "old_alt": 3}


def test_suggest_column_map_does_not_reuse_one_column_twice() -> None:
    mapping = suggest_column_map(("Alt text", "Image"))
    assert mapping["new_alt"] == 0
    assert mapping["image_url"] == 1


def test_suggest_column_map_prefers_exact_page_url_over_bare_url() -> None:
    mapping = suggest_column_map(("URL", "Alt", "Page URL"))
    assert mapping["image_url"] == 0
    assert mapping["page_url"] == 2


def test_suggest_column_map_ignores_unknown_headers() -> None:
    assert suggest_column_map(("foo", "bar", "")) == {}


# --- inspect ------------------------------------------------------------------


def test_inspect_xlsx_lists_every_sheet_and_picks_the_usable_one() -> None:
    content = _xlsx(
        {"Summary": [["Report", ""], ["generated", "2026-07-26"]], "Images": AUDIT_SHEET}
    )
    inspected = inspect_source("audit.xlsx", content)
    assert inspected.kind is SourceKind.xlsx
    assert [sheet.name for sheet in inspected.sheets] == ["Summary", "Images"]
    assert inspected.default_sheet == "Images"
    images = next(sheet for sheet in inspected.sheets if sheet.name == "Images")
    assert images.headers[:2] == ("Image URL", "New alt text")
    assert len(images.sample_rows) == 2
    assert images.suggested_column_map["new_alt"] == 1


def test_inspect_skips_a_title_row_above_the_header() -> None:
    content = _xlsx({"Images": [["Image audit — March"], *AUDIT_SHEET]})
    sheet = inspect_source("audit.xlsx", content).sheets[0]
    assert sheet.header_row_index == 1
    assert sheet.headers[0] == "Image URL"


def test_inspect_reports_an_empty_sheet_instead_of_crashing() -> None:
    content = _xlsx({"Empty": [], "Images": AUDIT_SHEET})
    empty = inspect_source("audit.xlsx", content).sheets[0]
    assert empty.warnings == ("empty_sheet",)
    assert empty.headers == ()


def test_inspect_flags_a_sheet_without_the_required_columns() -> None:
    content = _xlsx({"Other": [["foo", "bar"], ["1", "2"]]})
    sheet = inspect_source("audit.xlsx", content).sheets[0]
    assert "header_row_not_detected" in sheet.warnings


def test_inspect_csv_sniffs_semicolons_and_bom() -> None:
    raw = "﻿Image URL;New alt;Page URL\r\nhttps://a/b-300x200.jpg;An alt;https://a/page\r\n"
    inspected = inspect_source("audit.csv", raw.encode("utf-8"))
    assert inspected.delimiter == ";"
    assert inspected.encoding == "utf-8-sig"
    assert inspected.sheets[0].headers[0] == "Image URL"


def test_inspect_csv_reads_cp1251() -> None:
    raw = "Image URL,New alt\nhttps://a/b.jpg,Офіс команди\n".encode("cp1251")
    inspected = inspect_source("audit.csv", raw)
    assert inspected.encoding == "cp1251"
    assert inspected.sheets[0].sample_rows[0][1] == "Офіс команди"


def test_inspect_empty_csv_is_a_clear_error() -> None:
    with pytest.raises(UnreadableSourceError):
        inspect_source("audit.csv", b"   \n")


# --- parse --------------------------------------------------------------------


def _map() -> dict[str, int]:
    return {"image_url": 0, "new_alt": 1, "page_url": 2, "old_alt": 3}


def test_parse_xlsx_rows() -> None:
    content = _xlsx({"Images": AUDIT_SHEET})
    result = parse_rows("audit.xlsx", content, sheet="Images", column_map=_map())
    assert [row.image_url for row in result.rows] == [
        f"{UPLOADS}/hero-300x200.jpg",
        f"{UPLOADS}/logo-scaled.png",
    ]
    first, second = result.rows
    assert first.row_index == 2  # 1-based, під заголовком
    assert first.new_alt == "Team at the office"
    assert first.page_url == "https://example.com/about"
    assert first.old_alt is None  # порожня клітинка → None, не ""
    assert first.extra == {"Notes": "high priority"}
    assert second.old_alt == "old logo"
    assert result.skipped_count == 0
    assert result.missing_alt_count == 0


def test_parse_skips_rows_without_an_image_url_but_counts_them() -> None:
    content = _xlsx(
        {"Images": [*AUDIT_SHEET, ["", "orphan alt", "", "", ""], [None, None, None, None, None]]}
    )
    result = parse_rows("audit.xlsx", content, sheet="Images", column_map=_map())
    assert len(result.rows) == 2
    assert result.skipped_count == 2


def test_parse_keeps_rows_without_a_new_alt_and_flags_them() -> None:
    content = _xlsx({"Images": [*AUDIT_SHEET, [f"{UPLOADS}/x.jpg", "", "", "", ""]]})
    result = parse_rows("audit.xlsx", content, sheet="Images", column_map=_map())
    assert len(result.rows) == 3
    assert result.missing_alt_count == 1
    assert "rows_without_new_alt" in result.warnings


def test_parse_requires_the_mandatory_columns() -> None:
    content = _xlsx({"Images": AUDIT_SHEET})
    with pytest.raises(ColumnMappingError):
        parse_rows("audit.xlsx", content, sheet="Images", column_map={"image_url": 0})


def test_parse_rejects_an_unknown_sheet() -> None:
    content = _xlsx({"Images": AUDIT_SHEET})
    with pytest.raises(ColumnMappingError):
        parse_rows("audit.xlsx", content, sheet="Nope", column_map=_map())


def test_parse_errors_when_nothing_has_an_image_url() -> None:
    content = _xlsx({"Images": [["Image URL", "New alt"], ["", "alt only"]]})
    with pytest.raises(UnreadableSourceError):
        parse_rows("audit.xlsx", content, sheet="Images", column_map={"image_url": 0, "new_alt": 1})


def test_parse_csv_with_tabs() -> None:
    raw = "Image URL\tNew alt\nhttps://a/b.jpg\tAn alt\n"
    result = parse_rows("audit.tsv", raw.encode(), sheet=None,
                        column_map={"image_url": 0, "new_alt": 1})
    assert result.rows[0].new_alt == "An alt"


def test_parse_normalizes_numeric_cells() -> None:
    # Google Sheets віддає числа як float — «2024.0» у назві файлу не потрібне.
    content = _xlsx({"Images": [["Image URL", "New alt", "Width"],
                                [f"{UPLOADS}/a.jpg", "An alt", 1024]]})
    result = parse_rows("audit.xlsx", content, sheet="Images",
                        column_map={"image_url": 0, "new_alt": 1})
    assert result.rows[0].extra == {"Width": "1024"}


def test_parse_recovers_urls_from_hyperlinks() -> None:
    """Аркуш, де в колонці URL — підписи посилань, а не самі URL."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Images"
    worksheet.append(["Image URL", "New alt"])
    worksheet.append(["hero.jpg", "Team at the office"])
    worksheet["A2"].hyperlink = f"{UPLOADS}/hero-300x200.jpg"
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_rows(
        "audit.xlsx", buffer.getvalue(), sheet="Images",
        column_map={"image_url": 0, "new_alt": 1},
    )
    assert result.rows[0].image_url == f"{UPLOADS}/hero-300x200.jpg"
    assert "image_urls_taken_from_hyperlinks" in result.warnings


# --- одна картинка на кількох сторінках ----------------------------------------


def test_page_cell_with_several_urls_is_split() -> None:
    """Ahrefs «Used on page(s)» пакує всі сторінки в одну клітинку."""
    cell = "https://a.test/one\nhttps://a.test/two https://a.test/three"
    content = _xlsx({"Images": [["Image URL", "New alt", "Used on page(s)"],
                                [f"{UPLOADS}/a.jpg", "An alt", cell]]})
    result = parse_rows("audit.xlsx", content, sheet="Images",
                        column_map={"image_url": 0, "new_alt": 1, "page_url": 2})
    row = result.rows[0]
    assert row.page_urls == ["https://a.test/one", "https://a.test/two", "https://a.test/three"]
    # `page_url` лишається первинною сторінкою для простих читань.
    assert row.page_url == "https://a.test/one"


def test_split_page_urls_deduplicates_and_caps() -> None:
    assert split_page_urls("https://a/1 https://a/1 https://a/2") == ["https://a/1", "https://a/2"]
    assert split_page_urls("") == []
    # Не-URL не губимо: оператор має побачити, що було в клітинці.
    assert split_page_urls("see the brief") == ["see the brief"]
    many = " ".join(f"https://a/{index}" for index in range(MAX_PAGE_URLS + 10))
    assert len(split_page_urls(many)) == MAX_PAGE_URLS


def test_merged_image_url_cells_are_expanded_not_skipped() -> None:
    """Об'єднана клітинка URL: значення лише у верхньому рядку групи.

    Без розгортання 90%+ рядків реального Ahrefs-експорту виглядали як
    «рядок без URL картинки» й потрапляли у `skipped`.
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Images"
    worksheet.append(["Image URL", "New alt", "Used on page(s)"])
    worksheet.append([f"{UPLOADS}/grouped.jpg", "An alt", "https://a.test/one"])
    worksheet.append([None, "An alt", "https://a.test/two"])
    worksheet.append([None, "An alt", "https://a.test/three"])
    worksheet.merge_cells(start_row=2, start_column=1, end_row=4, end_column=1)
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_rows("audit.xlsx", buffer.getvalue(), sheet="Images",
                        column_map={"image_url": 0, "new_alt": 1, "page_url": 2})
    assert len(result.rows) == 3
    assert {row.image_url for row in result.rows} == {f"{UPLOADS}/grouped.jpg"}
    assert result.skipped_count == 0
    assert "merged_cells_expanded" in result.warnings


def test_parse_clips_an_enormous_cell() -> None:
    huge = "x" * (MAX_CELL_LENGTH + 500)
    content = _xlsx({"Images": [["Image URL", "New alt"], [f"{UPLOADS}/a.jpg", huge]]})
    result = parse_rows("audit.xlsx", content, sheet="Images",
                        column_map={"image_url": 0, "new_alt": 1})
    assert len(result.rows[0].new_alt) == MAX_CELL_LENGTH


def test_parse_caps_the_row_count_and_says_so() -> None:
    rows = [["Image URL", "New alt"]]
    rows.extend([[f"{UPLOADS}/img-{index}.jpg", f"alt {index}"] for index in range(MAX_ROWS + 25)])
    result = parse_rows("audit.csv", _csv(rows), sheet=None,
                        column_map={"image_url": 0, "new_alt": 1})
    assert len(result.rows) == MAX_ROWS
    assert "row_cap_reached" in result.warnings


def _csv(rows: list[list[str]]) -> bytes:
    return "\n".join(",".join(cell for cell in row) for row in rows).encode()
